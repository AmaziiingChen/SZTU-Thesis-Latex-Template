#!/usr/bin/env python3
"""Render the official SZTU midterm DOCX from schema-constrained JSON."""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image as PILImage

TEMPLATES_DIR = Path(__file__).resolve().parents[2]
if str(TEMPLATES_DIR) not in sys.path:
    sys.path.insert(0, str(TEMPLATES_DIR))

from common.python.content import (  # noqa: E402
    ContentDataError as DataError,
    display_width,
    figure_caption_runs,
    normalize_content_block,
    normalize_paragraph,
    plain_runs,
    prepare_equation_content,
    prepare_figure_content,
    require_object,
    require_text,
    runs_text,
)
from common.python.process_form import load_process_document_layout  # noqa: E402
from common.python.equation import add_numbered_omml_table, append_omml  # noqa: E402
from common.python.outline_numbering import (  # noqa: E402
    number_outline_levels,
    require_numbering_style,
    style_max_depth,
)


SCHEMA_VERSION = "0.1"
LAYOUT = load_process_document_layout(
    Path(__file__).resolve().parents[1] / "spec" / "layout.json"
)
BODY = LAYOUT["typography"]["body"]
LABEL = LAYOUT["typography"]["label"]
DATA = LAYOUT["typography"]["data"]
TEACHER_HEADER = LAYOUT["typography"]["teacher_header"]
TEACHER_BODY = LAYOUT["typography"]["teacher_body"]
SIGNATURE = LAYOUT["typography"]["signature"]
LINE_SPACING = LAYOUT["paragraphs"]["line_spacing"]
LIST_LEVEL_INDENT_PT = BODY["size_pt"] * LAYOUT["paragraphs"]["list_level_indent_em"]
LIST_HANGING_INDENT_PT = BODY["size_pt"] * LAYOUT["paragraphs"]["list_hanging_indent_em"]
LIST_MAX_DEPTH = LAYOUT["paragraphs"]["list_max_depth"]
UNORDERED_LIST_MARKERS = LAYOUT["paragraphs"]["unordered_list_markers"]
ADVISOR_TITLE_PATTERN = re.compile(r"(?:老师|教授|副教授|讲师|博士|硕士|导师)$")
METADATA_LIMITS = {
    "student_name": 20,
    "college": 50,
    "major": 50,
    "class_name": 30,
    "advisor": 30,
}


def validate_data(raw: Any) -> dict[str, Any]:
    data = require_object(raw, "root")
    expected_root = {"schema_version", "metadata", "sections"}
    unknown_root = set(data) - expected_root
    if unknown_root:
        raise DataError(f"unknown root fields: {sorted(unknown_root)}")
    if data.get("schema_version") != SCHEMA_VERSION:
        raise DataError(f"schema_version must be {SCHEMA_VERSION!r}")

    metadata = require_object(data.get("metadata"), "metadata")
    expected_metadata = set(METADATA_LIMITS) | {"check_date", "title"}
    unknown_metadata = set(metadata) - expected_metadata
    if unknown_metadata:
        raise DataError(f"unknown metadata fields: {sorted(unknown_metadata)}")
    clean_metadata = {
        key: require_text(metadata.get(key), f"metadata.{key}", limit)
        for key, limit in METADATA_LIMITS.items()
    }
    if ADVISOR_TITLE_PATTERN.search(clean_metadata["advisor"]):
        raise DataError("metadata.advisor must contain the name only, without a title")
    check_date = require_text(metadata.get("check_date"), "metadata.check_date", 10)
    try:
        parsed_date = date.fromisoformat(check_date)
    except ValueError as exc:
        raise DataError("metadata.check_date must be a valid ISO date") from exc
    title = normalize_paragraph(metadata.get("title"), "metadata.title", 90)
    if display_width(runs_text(title)) > 80:
        raise DataError("metadata.title exceeds the official two-line capacity")
    clean_metadata.update({"check_date": parsed_date.isoformat(), "title": title})

    sections = require_object(data.get("sections"), "sections")
    expected_sections = {
        "numbering_style",
        "directory",
        "main_research_content",
        "progress",
    }
    unknown_sections = set(sections) - expected_sections
    if unknown_sections:
        raise DataError(f"unknown section fields: {sorted(unknown_sections)}")

    numbering_style = sections.get("numbering_style")
    if numbering_style is not None:
        numbering_style = require_numbering_style(
            numbering_style, "sections.numbering_style"
        )
    maximum_outline_depth = (
        style_max_depth(numbering_style) if numbering_style is not None else 6
    )

    raw_directory = sections.get("directory")
    if not isinstance(raw_directory, list) or not 1 <= len(raw_directory) <= 150:
        raise DataError("sections.directory must contain 1 to 150 items")
    clean_directory = []
    previous_level = 1
    for index, raw_item in enumerate(raw_directory):
        path = f"sections.directory[{index}]"
        item = require_object(raw_item, path)
        unknown = set(item) - {"level", "number", "title"}
        if unknown:
            raise DataError(f"unknown fields in {path}: {sorted(unknown)}")
        level = item.get("level")
        if (
            isinstance(level, bool)
            or not isinstance(level, int)
            or not 1 <= level <= maximum_outline_depth
        ):
            raise DataError(
                f"{path}.level must be an integer from 1 to {maximum_outline_depth}"
            )
        if index == 0 and level != 1:
            raise DataError("the first directory item must be level 1")
        if level > previous_level + 1:
            raise DataError(f"{path}.level skips an outline level")
        previous_level = level
        number = item.get("number")
        if number is not None:
            number = require_text(number, f"{path}.number", 20)
        clean_directory.append(
            {
                "level": level,
                "number": number,
                "title": normalize_paragraph(item.get("title"), f"{path}.title", 200),
            }
        )

    if numbering_style is not None:
        canonical_numbers = number_outline_levels(
            [item["level"] for item in clean_directory], numbering_style
        )
        for item, number in zip(clean_directory, canonical_numbers, strict=True):
            item["number"] = number

    clean_sections: dict[str, Any] = {"directory": clean_directory}
    if numbering_style is not None:
        clean_sections["numbering_style"] = numbering_style
    for key in ("main_research_content", "progress"):
        raw_blocks = sections.get(key)
        if not isinstance(raw_blocks, list) or not 1 <= len(raw_blocks) <= 80:
            raise DataError(f"sections.{key} must contain 1 to 80 blocks")
        clean_sections[key] = [
            normalize_content_block(
                block,
                f"sections.{key}[{index}]",
                5000,
                max_list_depth=LIST_MAX_DEPTH,
            )
            for index, block in enumerate(raw_blocks)
        ]
    prepare_figure_content(
        [
            clean_sections["main_research_content"],
            clean_sections["progress"],
        ]
    )
    prepare_equation_content(
        [
            clean_sections["main_research_content"],
            clean_sections["progress"],
        ]
    )
    return {
        "schema_version": SCHEMA_VERSION,
        "metadata": clean_metadata,
        "sections": clean_sections,
    }


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def _clear_cell(cell) -> None:
    while len(cell.paragraphs) > 1:
        _remove_paragraph(cell.paragraphs[-1])
    cell.paragraphs[0].clear()


def _set_run_font(run, *, style: dict[str, Any], bold: bool = False) -> None:
    text = run.text or " "
    name = style["latin_word_family"] if all(ord(char) < 128 for char in text) else style["cjk_word_family"]
    run.font.name = name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), style["latin_word_family"])
    rfonts.set(qn("w:hAnsi"), style["latin_word_family"])
    rfonts.set(qn("w:cs"), style["latin_word_family"])
    rfonts.set(qn("w:eastAsia"), style["cjk_word_family"])
    run.font.size = Pt(style["size_pt"])
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _append_runs(paragraph, runs: list[dict[str, Any]], *, style: dict[str, Any]) -> None:
    for rich_run in runs:
        pieces = re.findall(r"[\x00-\x7f]+|[^\x00-\x7f]+", rich_run["text"])
        for piece in pieces:
            run = paragraph.add_run(piece)
            _set_run_font(run, style=style, bold=rich_run["bold"])
            run.italic = rich_run["italic"]
            run.font.subscript = rich_run["script"] == "sub"
            run.font.superscript = rich_run["script"] == "super"


def _format_paragraph(
    paragraph,
    *,
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    first_line_indent_pt: float | None = None,
    left_indent_pt: float | None = None,
) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = LINE_SPACING
    paragraph.paragraph_format.first_line_indent = (
        Pt(first_line_indent_pt) if first_line_indent_pt is not None else None
    )
    paragraph.paragraph_format.left_indent = (
        Pt(left_indent_pt) if left_indent_pt is not None else None
    )


def _fill_centered_cell(
    cell,
    value: str | list[dict[str, Any]],
    *,
    style: dict[str, Any],
) -> None:
    _clear_cell(cell)
    paragraph = cell.paragraphs[0]
    _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    _append_runs(paragraph, plain_runs(value) if isinstance(value, str) else value, style=style)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _fill_label_cell(cell, text: str) -> None:
    _fill_centered_cell(cell, text, style=LABEL)


def _new_paragraph(cell, *, indent: bool = False, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph = cell.add_paragraph()
    _format_paragraph(
        paragraph,
        alignment=alignment,
        first_line_indent_pt=BODY["size_pt"] * 2 if indent else None,
    )
    return paragraph


def _resolve_image(path_text: str, data_dir: Path) -> Path:
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", path_text):
        raise DataError("remote image URLs are forbidden; use a local file")
    path = Path(path_text).expanduser()
    if not path.is_absolute():
        path = data_dir / path
    path = path.resolve()
    if not path.is_file():
        raise DataError(f"image file does not exist: {path}")
    return path


def _set_picture_alt(run, alt: str) -> None:
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", alt)


def _fit_image_dimensions(
    image_path: Path,
    requested_width_mm: float,
    *,
    max_height_mm: float | None = None,
) -> tuple[float, float]:
    with PILImage.open(image_path) as image:
        pixel_width, pixel_height = image.size
    if pixel_width <= 0 or pixel_height <= 0:
        raise DataError("image dimensions must be positive")
    width_mm = float(requested_width_mm)
    height_mm = width_mm * pixel_height / pixel_width
    max_height_mm = (
        float(max_height_mm)
        if max_height_mm is not None
        else float(LAYOUT["image"]["max_height_mm"])
    )
    if height_mm > max_height_mm:
        scale = max_height_mm / height_mm
        width_mm *= scale
        height_mm = max_height_mm
    return width_mm, height_mm


def _set_table_borders(table, *, border_pt: float | None) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        if border_pt is None:
            element.set(qn("w:val"), "nil")
        else:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(round(border_pt * 8)))
            element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)


def _set_table_cell_margins(table, margin_mm: float) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblCellMar"))
    if old is not None:
        tbl_pr.remove(old)
    margins = OxmlElement("w:tblCellMar")
    value = str(round(Mm(margin_mm).twips))
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.append(margins)


def _set_table_geometry(table, widths_mm: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total_twips = round(Mm(sum(widths_mm)).twips)
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table._tbl.tblPr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    for grid_col, width_mm in zip(table._tbl.tblGrid.gridCol_lst, widths_mm, strict=True):
        grid_col.set(qn("w:w"), str(round(Mm(width_mm).twips)))
    for row in table.rows:
        for cell, width_mm in zip(row.cells, widths_mm, strict=True):
            width_twips = str(round(Mm(width_mm).twips))
            cell.width = Mm(width_mm)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), width_twips)
            tc_w.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _alignment(value: str):
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[value]


def _append_data_table(cell, block: dict[str, Any]) -> None:
    config = LAYOUT["embedded_table"]
    weights = [column["width_weight"] for column in block["columns"]]
    total_weight = sum(weights)
    widths = [float(config["width_mm"]) * weight / total_weight for weight in weights]
    table = cell.add_table(rows=1 + len(block["rows"]), cols=len(block["columns"]))
    _set_table_geometry(table, widths)
    _set_table_borders(table, border_pt=float(config["border_pt"]))
    _set_table_cell_margins(table, float(config["cell_padding_mm"]))
    _set_repeat_table_header(table.rows[0])
    for row in table.rows:
        _prevent_row_split(row)
    for column_index, column in enumerate(block["columns"]):
        target = table.rows[0].cells[column_index]
        target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = target.paragraphs[0]
        _format_paragraph(paragraph, alignment=_alignment(column["alignment"]))
        _append_runs(paragraph, column["header_runs"], style=BODY)
        for run in paragraph.runs:
            run.bold = True
    for row_index, values in enumerate(block["rows"], start=1):
        for column_index, runs in enumerate(values):
            target = table.rows[row_index].cells[column_index]
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = target.paragraphs[0]
            _format_paragraph(
                paragraph,
                alignment=_alignment(block["columns"][column_index]["alignment"]),
            )
            _append_runs(paragraph, runs, style=BODY)
    if block["caption_runs"]:
        caption = _new_paragraph(cell, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        caption.paragraph_format.keep_together = True
        caption.paragraph_format.space_before = Pt(config["caption_space_before_pt"])
        caption.paragraph_format.space_after = Pt(config["caption_space_after_pt"])
        _append_runs(caption, block["caption_runs"], style=BODY)


def _append_equation(cell, block: dict[str, Any]) -> None:
    config = LAYOUT["equation"]
    label = block.get("equation_label")
    if label is not None:
        _, paragraphs = add_numbered_omml_table(
            cell,
            block["expression"],
            label,
            width_mm=float(config["width_mm"]),
        )
        for paragraph in paragraphs:
            _format_paragraph(paragraph, alignment=paragraph.alignment)
            paragraph.paragraph_format.keep_together = True
        center = paragraphs[1]
        center.paragraph_format.space_before = Pt(config["space_before_pt"])
        center.paragraph_format.space_after = Pt(config["space_after_pt"])
        fallback = center.add_run(block["alt"])
        _set_run_font(fallback, style=BODY)
        fallback.font.hidden = True
        _set_run_font(paragraphs[2].runs[0], style=BODY)
        return
    paragraph = _new_paragraph(cell, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph.paragraph_format.space_before = Pt(config["space_before_pt"])
    paragraph.paragraph_format.space_after = Pt(config["space_after_pt"])
    paragraph.paragraph_format.keep_together = True
    fallback = paragraph.add_run(block["alt"])
    _set_run_font(fallback, style=BODY)
    fallback.font.hidden = True
    append_omml(paragraph, block["expression"])


def _append_figure_group(cell, block: dict[str, Any], *, data_dir: Path) -> None:
    config = LAYOUT["figure_group"]
    count = len(block["items"])
    total_width = float(config["width_mm"])
    gap = float(config["column_gap_mm"])
    item_width = (total_width - gap * (count - 1)) / count
    widths = [item_width] * count
    table = cell.add_table(rows=1, cols=count)
    _set_table_geometry(table, widths)
    _set_table_borders(table, border_pt=None)
    _set_table_cell_margins(table, gap / 2)
    _prevent_row_split(table.rows[0])
    for index, item in enumerate(block["items"]):
        target = table.rows[0].cells[index]
        target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = target.paragraphs[0]
        _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run()
        _set_run_font(run, style=BODY)
        image_path = _resolve_image(item["path"], data_dir)
        width_mm, height_mm = _fit_image_dimensions(
            image_path,
            item_width - gap,
            max_height_mm=float(config["max_item_height_mm"]),
        )
        run.add_picture(str(image_path), width=Mm(width_mm), height=Mm(height_mm))
        _set_picture_alt(run, item["alt"])
        subcaption = target.add_paragraph()
        _format_paragraph(subcaption, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        subcaption.paragraph_format.keep_together = True
        subcaption.paragraph_format.keep_with_next = True
        _append_runs(subcaption, plain_runs(item["subfigure_label"]), style=BODY)
        if item["caption_runs"]:
            _append_runs(subcaption, item["caption_runs"], style=BODY)
    caption = _new_paragraph(cell, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_before = Pt(
        LAYOUT["paragraphs"]["caption_space_before_pt"]
    )
    caption.paragraph_format.space_after = Pt(
        LAYOUT["paragraphs"]["caption_space_after_pt"]
    )
    _append_runs(caption, figure_caption_runs(block), style=BODY)


def _list_marker(block_type: str, item: dict[str, Any], index: int, depth: int) -> str:
    if block_type == "ordered_list":
        return item["marker"] or f"{index}、"
    return UNORDERED_LIST_MARKERS[depth - 1]


def _append_list_block(cell, block: dict[str, Any], *, depth: int = 1) -> None:
    if not 1 <= depth <= LIST_MAX_DEPTH:
        raise AssertionError(f"normalized list depth escaped bounds: {depth}")
    for index, item in enumerate(block["items"], start=1):
        paragraph = cell.add_paragraph()
        _format_paragraph(
            paragraph,
            left_indent_pt=depth * LIST_LEVEL_INDENT_PT + LIST_HANGING_INDENT_PT,
            first_line_indent_pt=-LIST_HANGING_INDENT_PT,
        )
        marker = _list_marker(block["type"], item, index, depth)
        _append_runs(paragraph, plain_runs(f"{marker} "), style=BODY)
        _append_runs(paragraph, item["runs"], style=BODY)
        if item["children"]:
            _append_list_block(cell, item["children"], depth=depth + 1)


def _append_content_blocks(cell, blocks: list[dict[str, Any]], *, data_dir: Path) -> None:
    for block in blocks:
        if block["type"] == "paragraph":
            paragraph = _new_paragraph(cell, indent=True)
            _append_runs(paragraph, block["runs"], style=BODY)
        elif block["type"] in {"ordered_list", "unordered_list"}:
            _append_list_block(cell, block)
        elif block["type"] == "image":
            paragraph = _new_paragraph(cell, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run()
            _set_run_font(run, style=BODY)
            image_path = _resolve_image(block["path"], data_dir)
            width_mm, height_mm = _fit_image_dimensions(
                image_path, block["width_mm"]
            )
            run.add_picture(
                str(image_path), width=Mm(width_mm), height=Mm(height_mm)
            )
            _set_picture_alt(run, block["alt"])
            caption = _new_paragraph(cell, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            caption.paragraph_format.keep_together = True
            caption.paragraph_format.space_before = Pt(
                LAYOUT["paragraphs"]["caption_space_before_pt"]
            )
            caption.paragraph_format.space_after = Pt(
                LAYOUT["paragraphs"]["caption_space_after_pt"]
            )
            _append_runs(caption, figure_caption_runs(block), style=BODY)
        elif block["type"] == "data_table":
            _append_data_table(cell, block)
        elif block["type"] == "figure_group":
            _append_figure_group(cell, block, data_dir=data_dir)
        elif block["type"] == "equation":
            _append_equation(cell, block)
        else:
            raise AssertionError(f"unsupported normalized block: {block['type']}")


def _fill_student_sections(table, data: dict[str, Any], *, data_dir: Path) -> None:
    sections = data["sections"]
    directory_cell = table.rows[4].cells[0]
    _clear_cell(directory_cell)
    label = directory_cell.paragraphs[0]
    _format_paragraph(label)
    _append_runs(label, plain_runs("毕业论文（设计）的目录和主要研究内容："), style=LABEL)
    directory_heading = _new_paragraph(directory_cell, indent=True)
    _append_runs(directory_heading, plain_runs("目录："), style=BODY)
    for item in sections["directory"]:
        paragraph = directory_cell.add_paragraph()
        level_indent = (
            BODY["size_pt"]
            * LAYOUT["paragraphs"]["outline_level_indent_em"]
            * item["level"]
        )
        hanging_indent = (
            BODY["size_pt"] * LAYOUT["paragraphs"]["outline_hanging_indent_em"]
            if item["number"]
            else 0
        )
        _format_paragraph(
            paragraph,
            left_indent_pt=level_indent + hanging_indent,
            first_line_indent_pt=-hanging_indent if hanging_indent else None,
        )
        if item["number"]:
            _append_runs(paragraph, plain_runs(f"{item['number']} "), style=BODY)
        _append_runs(paragraph, item["title"], style=BODY)
    research_heading = _new_paragraph(directory_cell, indent=True)
    _append_runs(research_heading, plain_runs("主要研究内容："), style=BODY)
    _append_content_blocks(
        directory_cell,
        sections["main_research_content"],
        data_dir=data_dir,
    )

    progress_cell = table.rows[5].cells[0]
    _clear_cell(progress_cell)
    progress_label = progress_cell.paragraphs[0]
    _format_paragraph(progress_label)
    _append_runs(progress_label, plain_runs("毕业论文（设计）工作进展情况（详述）："), style=LABEL)
    _append_content_blocks(progress_cell, sections["progress"], data_dir=data_dir)


def _fill_teacher_fixed_rows(table) -> None:
    fixed_rows = {
        6: ["指导教师填写栏目（在正确项后方框内划√）"],
        7: [
            "1、毕业论文（设计）进展情况：",
            "（1）提前完成□；    （2）正常进行□；    （3）延期滞后□",
        ],
        8: [
            "2、学生对毕业论文（设计）的认真程度：",
            "（1）认真□；        （2）较认真□；      （3）不认真□",
        ],
        9: [
            "3、查阅文献资料的能力：",
            "（1）强□；          （2）一般□；        （3）差□",
        ],
        10: [
            "4、已完成的毕业论文（设计）中期质量评价：",
            "（1）好□；          （2）中□；          （3）差□",
        ],
        11: ["5、毕业论文（设计）方向有无更改：（1）有□；        （2）无□"],
        12: ["6、对能否按期完成毕业论文（设计）的评估：（1）能□；    （2）否□"],
    }
    for row_index, lines in fixed_rows.items():
        cell = table.rows[row_index].cells[0]
        _clear_cell(cell)
        for line_index, text in enumerate(lines):
            paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
            _format_paragraph(paragraph)
            style = TEACHER_HEADER if row_index == 6 else TEACHER_BODY
            _append_runs(paragraph, plain_runs(text), style=style)


def _fill_signature_region(cell, *, label: str, signer: str, min_height_mm: float) -> None:
    _clear_cell(cell)
    label_paragraph = cell.paragraphs[0]
    _format_paragraph(label_paragraph)
    _append_runs(label_paragraph, plain_runs(label), style=SIGNATURE)
    for _ in range(9):
        paragraph = cell.add_paragraph()
        _format_paragraph(paragraph)
    signature = cell.add_paragraph()
    _format_paragraph(signature, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _append_runs(signature, plain_runs(signer), style=SIGNATURE)
    date_paragraph = cell.add_paragraph()
    _format_paragraph(date_paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _append_runs(date_paragraph, plain_runs("年      月      日"), style=SIGNATURE)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _set_row_min_height(row, height_mm: float) -> None:
    row.height = Mm(height_mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def render(template: Path, data_path: Path, output: Path, *, overwrite: bool) -> Path:
    if output.exists() and not overwrite:
        raise FileExistsError(f"output exists; pass --overwrite: {output}")
    data = validate_data(json.loads(data_path.read_text(encoding="utf-8")))
    document = Document(template)
    if len(document.tables) != 1:
        raise RuntimeError("official template must contain exactly one table")
    table = document.tables[0]
    if len(table.rows) != 15 or len(table.columns) != 4:
        raise RuntimeError("official template must expose a 15-row, 4-column grid")

    metadata = data["metadata"]
    labels = ((0, 0, "学生姓名"), (0, 2, "学院"), (1, 0, "专业"), (1, 2, "班级"),
              (2, 0, "指导教师"), (2, 2, "中期检查时间"), (3, 0, "论文题目"))
    for row_index, cell_index, text in labels:
        _fill_label_cell(table.rows[row_index].cells[cell_index], text)
    values = (
        (0, 1, metadata["student_name"]),
        (0, 3, metadata["college"]),
        (1, 1, metadata["major"]),
        (1, 3, metadata["class_name"]),
        (2, 1, metadata["advisor"]),
        (2, 3, f"{metadata['check_date'][0:4]} 年 {int(metadata['check_date'][5:7])} 月 {int(metadata['check_date'][8:10])} 日"),
        (3, 1, metadata["title"]),
    )
    for row_index, cell_index, value in values:
        _fill_centered_cell(table.rows[row_index].cells[cell_index], value, style=DATA)

    for row_index in range(4):
        _set_row_min_height(table.rows[row_index], LAYOUT["row_min_heights_mm"]["metadata"])
    _fill_student_sections(table, data, data_dir=data_path.resolve().parent)
    _set_row_min_height(table.rows[5], LAYOUT["section_min_heights_mm"]["progress"])
    _fill_teacher_fixed_rows(table)
    _set_row_min_height(table.rows[6], LAYOUT["row_min_heights_mm"]["teacher_header"])
    _set_row_min_height(table.rows[7], LAYOUT["row_min_heights_mm"]["teacher_option_first"])
    for row_index in range(8, 13):
        _set_row_min_height(table.rows[row_index], LAYOUT["row_min_heights_mm"]["teacher_option"])
    _fill_signature_region(
        table.rows[13].cells[0],
        label="存在的问题及后期指导工作意见：",
        signer="指导教师签名：",
        min_height_mm=LAYOUT["section_min_heights_mm"]["teacher_opinion"],
    )
    _set_row_min_height(table.rows[13], LAYOUT["section_min_heights_mm"]["teacher_opinion"])
    _fill_signature_region(
        table.rows[14].cells[0],
        label="审查小组检查意见：",
        signer="审查小组负责人签名：",
        min_height_mm=LAYOUT["section_min_heights_mm"]["review_group_opinion"],
    )
    _set_row_min_height(table.rows[14], LAYOUT["section_min_heights_mm"]["review_group_opinion"])

    document.core_properties.author = "SZTU Thesis Template"
    document.core_properties.last_modified_by = "SZTU Thesis Template"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    reopened = Document(output)
    if len(reopened.tables) != 1 or len(reopened.tables[0].rows) != 15:
        raise RuntimeError("generated DOCX failed structural reopen validation")
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=Path(__file__).with_name("official-template.docx"))
    parser.add_argument("--data", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    try:
        output = render(args.template, args.data, args.output, overwrite=args.overwrite)
    except (DataError, OSError, RuntimeError, ValueError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
