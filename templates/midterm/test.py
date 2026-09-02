#!/usr/bin/env python3
"""Run structural and rendered regression checks for the midterm template."""

from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path

import pdfplumber
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

TEMPLATES_DIR = Path(__file__).resolve().parents[1]
if str(TEMPLATES_DIR) not in sys.path:
    sys.path.insert(0, str(TEMPLATES_DIR))

from common.python.process_form import load_process_document_layout  # noqa: E402


def run(command: list[str], *, cwd: Path) -> str:
    result = subprocess.run(
        command,
        cwd=cwd,
        check=True,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
    )
    return result.stdout


def load_renderer(path: Path, module_name: str):
    spec = importlib.util.spec_from_file_location(module_name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load renderer: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def paragraph_runs(cell) -> list:
    return [run for paragraph in cell.paragraphs for run in paragraph.runs]


def assert_run_role(run, *, cjk: str, latin: str, size: float, bold: bool) -> None:
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    assert rfonts.get(qn("w:eastAsia")) == cjk
    assert rfonts.get(qn("w:ascii")) == latin
    assert run.font.size is not None and run.font.size.pt == size
    assert bool(run.bold) is bold
    assert run.font.color.rgb is not None
    assert str(run.font.color.rgb) == "000000"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-pdf", action="store_true")
    args = parser.parse_args()

    midterm_dir = Path(__file__).resolve().parent
    project_dir = midterm_dir.parents[1]
    output_root = project_dir / "tmp" / "midterm-tests"
    output_root.mkdir(parents=True, exist_ok=True)
    layout = load_process_document_layout(midterm_dir / "spec" / "layout.json")
    word_renderer = load_renderer(midterm_dir / "word" / "render.py", "midterm_word_test")
    latex_renderer = load_renderer(midterm_dir / "latex" / "render.py", "midterm_latex_test")
    latex_source = (midterm_dir / "latex" / "main.tex").read_text(encoding="utf-8")
    assert r"\input{sztu-process-form.tex}" in latex_source
    assert r"\newcommand{\TightJoin}" not in latex_source
    assert r"\usepackage{needspace}" not in latex_source
    assert r"\Needspace" not in latex_source
    assert r"\SZTUMainMinHeight" not in latex_source
    assert r"\SZTUTeacherBundleNeedspace" not in latex_source
    assert r"\begin{minipage}{\linewidth}" in latex_source
    assert r"height=\SZTUImageMaxHeight,keepaspectratio" in latex_source

    plain_run = {"text": "完成开题", "script": "normal", "italic": False, "bold": False}
    assert latex_renderer.rich_runs([plain_run]) == r"完成开\nobreak{}题"
    punctuated_run = dict(plain_run, text="完成开题。")
    assert latex_renderer.rich_runs([punctuated_run]) == r"完成开\nobreak{}题。"
    latin_ending_run = dict(plain_run, text="完成模型 CNN")
    assert latex_renderer.rich_runs([latin_ending_run]) == "完成模型 CNN"
    emphasis_runs = [
        {"text": "中文加粗", "bold": True, "italic": False, "script": "normal"},
        {"text": "中文倾斜", "bold": False, "italic": True, "script": "normal"},
        {"text": "中文加粗倾斜", "bold": True, "italic": True, "script": "normal"},
    ]
    emphasis_tex = latex_renderer.rich_runs(emphasis_runs)
    emphasis_tex = emphasis_tex.replace(r"\nobreak{}", "")
    assert r"\textbf{中文加粗}" in emphasis_tex
    assert r"\textit{中文倾斜}" in emphasis_tex
    assert latex_renderer._render_rich_run(
        emphasis_runs[2], "中文加粗倾斜"
    ) == r"\textbf{\textit{中文加粗倾斜}}"
    font_dir = output_root / "synthetic-fonts"
    font_tex = latex_renderer.fonts_tex(
        {
            "FZXiaoBiaoSong": font_dir / "form-title.ttf",
            "SimSun": font_dir / "simsun.ttf",
            "Times New Roman": font_dir / "times.ttf",
            "Times New Roman Bold": font_dir / "timesbd.ttf",
            "Times New Roman Italic": font_dir / "timesi.ttf",
            "Times New Roman Bold Italic": font_dir / "timesbi.ttf",
        }
    )
    assert "AutoFakeBold=3,AutoFakeSlant=0.2" in font_tex
    assert font_tex.count("AutoFakeSlant=0.2") >= 3

    assert layout["typography"]["title"]["size_pt"] == 18.0
    assert layout["typography"]["title"]["bold"] is False
    assert layout["typography"]["label"]["size_pt"] == 12.0
    assert layout["typography"]["data"]["size_pt"] == 10.5
    assert layout["typography"]["teacher_header"]["bold"] is True
    assert layout["table"]["border_pt"] == 0.5
    assert layout["table"]["flow_vertical_padding_mm"] == 1.5
    assert layout["table"]["flow_end_space_mm"] == 4.0
    assert layout["table"]["section_title_content_gap_mm"] == 0.0
    assert layout["paragraphs"]["list_max_depth"] == 4
    assert layout["paragraphs"]["unordered_list_markers"] == ["•", "◦", "▪", "▫"]
    assert layout["image"]["max_height_mm"] == 180.0
    assert layout["signature"]["teacher_blank_width_mm"] == 45.0
    assert layout["signature"]["review_blank_width_mm"] == 35.0
    assert layout["latex_pagination"]["teacher_opinion_height_mm"] == 84.0
    assert layout["latex_pagination"]["review_group_opinion_height_mm"] == 72.0
    assert "teacher_bundle_needspace_mm" not in layout["latex_pagination"]
    assert "directory_and_research" not in layout["section_min_heights_mm"]
    assert layout["section_min_heights_mm"]["progress"] == 55.12
    assert layout["required_font_files"] == [
        "FZXiaoBiaoSong",
        "SimSun",
        "Times New Roman",
        "Times New Roman Bold",
        "Times New Roman Bold Italic",
        "Times New Roman Italic",
    ]

    official_template = midterm_dir / "word" / "official-template.docx"
    assert sha256(official_template) == "5411d1471ba1a30def0b8c945a11674cce4616a6b3c0d3a8953e3f45b88af1a2"
    official = Document(official_template)
    assert len(official.tables) == 1
    assert len(official.tables[0].rows) == 15
    assert official.tables[0].rows[4].height is None
    assert official.tables[0].rows[4].height_rule is None
    assert round(official.tables[0].rows[5].height.mm, 2) == 55.12
    assert official.tables[0].rows[5].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
    title_runs = [run for paragraph in official.paragraphs[:2] for run in paragraph.runs]
    title_runs = [run for run in title_runs if run.text.strip()]
    assert title_runs
    for title_run in title_runs:
        assert title_run.font.size is not None and title_run.font.size.pt == 18.0
        assert not bool(title_run.bold)
        rfonts = title_run._element.get_or_add_rPr().get_or_add_rFonts()
        assert rfonts.get(qn("w:eastAsia")) == "方正小标宋简体"

    fixture_names = (
        "minimal",
        "normal",
        "layout-stress",
        "long-with-image",
        "nested-list",
        "outline-numbering",
        "structured-content",
        "extreme-pagination",
    )
    for name in fixture_names:
        fixture = midterm_dir / "fixtures" / f"{name}.json"
        raw = json.loads(fixture.read_text(encoding="utf-8"))
        normalized = word_renderer.validate_data(raw)
        assert normalized["schema_version"] == "0.1"
        assert normalized["sections"]["directory"][0]["level"] == 1
        if name == "long-with-image":
            research = normalized["sections"]["main_research_content"]
            progress = normalized["sections"]["progress"]
            figures = [
                block
                for block in [*research, *progress]
                if block["type"] == "image"
            ]
            assert [block["figure_label"] for block in figures] == [
                "图 3-1",
                "图 3-2",
            ]
            normalized_text = "".join(
                word_renderer.runs_text(block["runs"])
                for block in [*research, *progress]
                if block["type"] == "paragraph"
            )
            assert "{{fig:" not in normalized_text
            assert "图 3-1" in normalized_text
            assert "图 3-2" in normalized_text
        if name == "structured-content":
            research = normalized["sections"]["main_research_content"]
            progress = normalized["sections"]["progress"]
            data_tables = [block for block in research if block["type"] == "data_table"]
            figure_groups = [block for block in progress if block["type"] == "figure_group"]
            equations = [block for block in research if block["type"] == "equation"]
            assert len(data_tables) == 1
            assert len(data_tables[0]["columns"]) == 4
            assert len(data_tables[0]["rows"]) == 3
            assert len(figure_groups) == 1
            assert len(equations) == 1
            assert equations[0]["expression"]["type"] == "row"
            assert equations[0]["equation_label"] == "(3-1)"
            assert "{{eq:" not in research[0]["runs"][0]["text"]
            assert "(3-1)" in research[0]["runs"][0]["text"]
            assert figure_groups[0]["figure_label"] == "图 3-1"
            assert [item["subfigure_label"] for item in figure_groups[0]["items"]] == [
                "（a）",
                "（b）",
            ]
        if name == "extreme-pagination":
            research = normalized["sections"]["main_research_content"]
            progress = normalized["sections"]["progress"]
            assert len([block for block in research if block["type"] == "data_table"]) == 2
            assert len([block for block in progress if block["type"] == "data_table"]) == 1
            groups = [block for block in progress if block["type"] == "figure_group"]
            assert len(groups) == 4
            assert [block["figure_label"] for block in groups] == [
                "图 3-1",
                "图 3-2",
                "图 3-3",
                "图 3-4",
            ]

        docx_output = output_root / f"{name}.docx"
        run(
            [
                sys.executable,
                str(midterm_dir / "word" / "render.py"),
                "--data",
                str(fixture),
                "--output",
                str(docx_output),
                "--overwrite",
            ],
            cwd=project_dir,
        )
        document = Document(docx_output)
        assert len(document.sections) == 1
        section = document.sections[0]
        assert abs(section.page_width.mm - 210.0) <= 0.02
        assert abs(section.page_height.mm - 297.0) <= 0.02
        assert round(section.top_margin.mm, 2) == 25.4
        assert round(section.bottom_margin.mm, 2) == 25.4
        assert round(section.left_margin.mm, 2) == 31.75
        assert round(section.right_margin.mm, 2) == 31.75
        assert len(document.tables) == 1
        table = document.tables[0]
        assert len(table.rows) == 15 and len(table.columns) == 4
        assert table.rows[4].height is None
        assert table.rows[4].height_rule is None
        assert round(table.rows[5].height.mm, 2) == 55.12
        assert table.rows[5].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
        assert "FF0000" not in document.element.xml.upper()
        assert "导师应对论文完成进度" not in table.rows[13].cells[0].text
        assert "导师手签" not in table.rows[13].cells[0].text
        assert "存在的问题及后期指导工作意见：" in table.rows[13].cells[0].text
        assert "审查小组负责人签名：" in table.rows[14].cells[0].text
        assert "指导教师填写栏目（在正确项后方框内划√）" in table.rows[6].cells[0].text

        for row_index, cell_index in (
            (0, 0), (0, 1), (0, 2), (0, 3),
            (1, 0), (1, 1), (1, 2), (1, 3),
            (2, 0), (2, 1), (2, 2), (2, 3),
            (3, 0), (3, 1),
        ):
            cell = table.rows[row_index].cells[cell_index]
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
            assert all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in cell.paragraphs)

        for row_index, cell_index in ((0, 0), (0, 2), (1, 0), (1, 2), (2, 0), (2, 2), (3, 0)):
            for run_item in paragraph_runs(table.rows[row_index].cells[cell_index]):
                if run_item.text:
                    assert_run_role(
                        run_item,
                        cjk="宋体",
                        latin="Times New Roman",
                        size=12.0,
                        bold=False,
                    )
        for row_index, cell_index in ((0, 1), (0, 3), (1, 1), (1, 3), (2, 1), (2, 3), (3, 1)):
            for run_item in paragraph_runs(table.rows[row_index].cells[cell_index]):
                if run_item.text:
                    assert_run_role(
                        run_item,
                        cjk="宋体",
                        latin="Times New Roman",
                        size=10.5,
                        bold=False,
                    )

        directory_cell = table.rows[4].cells[0]
        research_start = next(
            index
            for index, paragraph in enumerate(directory_cell.paragraphs)
            if paragraph.text == "主要研究内容："
        )
        body_paragraphs = [
            paragraph
            for paragraph in directory_cell.paragraphs[research_start + 1 :]
            if paragraph.text.strip()
        ]
        assert body_paragraphs
        paragraph_blocks = [
            block
            for block in normalized["sections"]["main_research_content"]
            if block["type"] == "paragraph"
        ]
        assert paragraph_blocks
        assert any(
            paragraph.paragraph_format.first_line_indent is not None
            and paragraph.paragraph_format.first_line_indent.pt == 21.0
            for paragraph in body_paragraphs
        )
        first_outline = directory_cell.paragraphs[2]
        assert first_outline.paragraph_format.left_indent is not None
        assert first_outline.paragraph_format.left_indent.pt == 36.75
        assert first_outline.paragraph_format.first_line_indent is not None
        assert first_outline.paragraph_format.first_line_indent.pt == -15.75

        progress_cell = table.rows[5].cells[0]
        progress_body = [p for p in progress_cell.paragraphs[1:] if p.text.strip()]
        assert progress_body
        assert any(
            p.paragraph_format.first_line_indent is not None
            and p.paragraph_format.first_line_indent.pt == 21.0
            for p in progress_body
        )
        if any(block["type"] == "ordered_list" for block in normalized["sections"]["progress"]):
            list_paragraphs = [p for p in progress_body if re.match(r"^(?:（\d+）|\d+、)", p.text)]
            assert list_paragraphs
            assert all(p.paragraph_format.first_line_indent.pt == -21.0 for p in list_paragraphs)
            assert all(p.paragraph_format.left_indent.pt == 42.0 for p in list_paragraphs)

        if name == "normal":
            all_runs = [run_item for cell in (directory_cell, progress_cell) for run_item in paragraph_runs(cell)]
            assert any(run_item.text == "i" and run_item.font.subscript for run_item in all_runs)
        if name == "layout-stress":
            title_cell = table.rows[3].cells[1]
            assert len(title_cell.paragraphs[0].text) > 35
            assert title_cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
        if name == "long-with-image":
            xml = document.element.xml
            assert "统一数据模型驱动 Word 与 LaTeX 双路渲染流程示意图" in xml
            assert "{{fig:" not in xml
            figure_paragraphs = [
                paragraph
                for cell in (directory_cell, progress_cell)
                for paragraph in cell.paragraphs
                if paragraph._p.xpath(".//w:drawing")
            ]
            assert len(figure_paragraphs) == 2
            assert all(
                paragraph.paragraph_format.keep_with_next is True
                and paragraph.paragraph_format.keep_together is True
                for paragraph in figure_paragraphs
            )
            caption_paragraphs = [
                paragraph
                for cell in (directory_cell, progress_cell)
                for paragraph in cell.paragraphs
                if paragraph.text.startswith("图 3-")
            ]
            assert [paragraph.text for paragraph in caption_paragraphs] == [
                "图 3-1 统一数据入口与双路渲染流程（来源：SyntheticWorkflow2026）",
                "图 3-2 中期工作进展验证结果",
            ]
            assert all(
                paragraph.paragraph_format.keep_together is True
                for paragraph in caption_paragraphs
            )
            all_runs = [run_item for cell in (directory_cell, progress_cell) for run_item in paragraph_runs(cell)]
            assert any(run_item.text == "2" and run_item.font.subscript for run_item in all_runs)
            assert any(run_item.text == "−" and run_item.font.superscript for run_item in all_runs)
            assert any(run_item.text == "3" and run_item.font.superscript for run_item in all_runs)
        if name == "nested-list":
            by_prefix = {
                prefix: next(p for p in body_paragraphs if p.text.startswith(prefix))
                for prefix in ("•", "（1）", "▪", "A.", "列表结束后")
            }
            for depth, prefix in enumerate(("•", "（1）", "▪", "A."), start=1):
                paragraph = by_prefix[prefix]
                assert paragraph.paragraph_format.left_indent is not None
                assert paragraph.paragraph_format.left_indent.pt == (
                    depth * layout["typography"]["body"]["size_pt"]
                    * layout["paragraphs"]["list_level_indent_em"]
                    + layout["typography"]["body"]["size_pt"]
                    * layout["paragraphs"]["list_hanging_indent_em"]
                )
                assert paragraph.paragraph_format.first_line_indent.pt == -21.0
            trailing = by_prefix["列表结束后"]
            assert trailing.paragraph_format.left_indent is None
            assert trailing.paragraph_format.first_line_indent.pt == 21.0
            nested_runs = [run_item for paragraph in body_paragraphs for run_item in paragraph.runs]
            assert any(run_item.text == "2" and run_item.font.subscript for run_item in nested_runs)
            assert normalized["sections"]["main_research_content"][1]["type"] == "unordered_list"
        if name == "structured-content":
            assert len(directory_cell.tables) == 2
            nested_data_table = directory_cell.tables[0]
            equation_table = directory_cell.tables[1]
            assert len(nested_data_table.rows) == 4
            assert len(nested_data_table.columns) == 4
            assert nested_data_table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
            assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in nested_data_table.rows)
            assert len(progress_cell.tables) == 1
            figure_group_table = progress_cell.tables[0]
            assert len(figure_group_table.rows) == 1
            assert len(figure_group_table.columns) == 2
            assert figure_group_table.rows[0]._tr.xpath("./w:trPr/w:cantSplit")
            assert "基线方法流程示意图" in document.element.xml
            assert "改进方法流程示意图" in document.element.xml
            assert "图 3-1 两种方法的处理流程对比" in progress_cell.text
            document_xml = document.element.xml
            for math_tag in ("m:oMathPara", "m:f", "m:sSub", "m:sSup", "m:rad"):
                assert f"<{math_tag}" in document_xml
            assert "效率等于输入输出浓度差" in document_xml
            assert equation_table.rows[0].cells[2].text == "(3-1)"
            assert equation_table.rows[0]._tr.xpath("./w:trPr/w:cantSplit")
            assert all(
                border.get(qn("w:val")) == "nil"
                for border in equation_table._tbl.tblPr.find(qn("w:tblBorders"))
            )
            assert equation_table.rows[0].cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
        if name == "extreme-pagination":
            assert len(directory_cell.tables) == 2
            assert len(progress_cell.tables) == 5
            assert all(
                row._tr.xpath("./w:trPr/w:cantSplit")
                for nested in [*directory_cell.tables, *progress_cell.tables]
                for row in nested.rows
            )
            assert "图 3-4 第四组组合图" in progress_cell.text

        if args.skip_pdf:
            if name == "nested-list":
                data_tex = latex_renderer.data_tex(
                    normalized,
                    data_dir=fixture.parent,
                    assets_dir=output_root / "nested-list-assets",
                )
                for depth, marker in ((1, "•"), (2, "（1）"), (3, "▪"), (4, "A.")):
                    assert rf"\MidtermListItem{{{depth}}}{{{marker}}}" in data_tex
                assert r"\MidtermParagraph{列表结束后" in data_tex
                assert r"H\textsubscript{2}O" in data_tex
            continue

        latex_dir = output_root / f"latex-{name}"
        build_output = run(
            [
                sys.executable,
                str(midterm_dir / "latex" / "render.py"),
                "--data",
                str(fixture),
                "--output-dir",
                str(latex_dir),
                "--overwrite",
                "--compile",
            ],
            cwd=project_dir,
        )
        assert "Overfull" not in build_output
        assert "Underfull" not in build_output
        assert "Missing character" not in build_output
        fonts_tex = (latex_dir / "midterm-fonts.tex").read_text(encoding="utf-8")
        data_tex = (latex_dir / "midterm-data.tex").read_text(encoding="utf-8")
        typography_tex = (latex_dir / "midterm-typography.tex").read_text(encoding="utf-8")
        process_form_tex = (latex_dir / "sztu-process-form.tex").read_text(
            encoding="utf-8"
        )
        assert "FZXiaoBiaoSong" not in fonts_tex or "方正小标宋简体" in fonts_tex
        assert "AutoFakeBold=3" in fonts_tex
        assert r"\newcommand{\SZTUTitleSize}{\zihao{-2}}" in typography_tex
        assert r"\newcommand{\SZTUDataSize}{\zihao{5}}" in typography_tex
        assert r"\newcommand{\SZTUFormRuleWidth}{0.5pt}" in typography_tex
        assert r"\newcommand{\SZTUFlowVerticalPadding}{1.5mm}" in typography_tex
        assert r"\newcommand{\SZTUFlowEndSpace}{4mm}" in typography_tex
        assert r"\newcommand{\SZTUSectionTitleContentGap}{0mm}" in typography_tex
        assert "sztuformflow/.style" in process_form_tex
        assert r"\newcommand{\SZTUListLevelIndent}{2em}" in typography_tex
        assert r"\newcommand{\SZTUListHangingIndent}{2em}" in typography_tex
        assert r"\newcommand{\SZTUTeacherOpinionHeight}{84mm}" in typography_tex
        assert r"\newcommand{\SZTUReviewOpinionHeight}{72mm}" in typography_tex
        assert r"\newcommand{\SZTUProgressMinHeight}{55.12mm}" in typography_tex
        assert r"\newcommand{\SZTUImageMaxHeight}{180mm}" in typography_tex
        assert r"\newcommand{\SZTUMainMinHeight}" not in typography_tex
        assert r"\newcommand{\SZTUTeacherBundleNeedspace}" not in typography_tex
        assert r"\newcommand{\SZTUTeacherSignatureBlank}{45mm}" in typography_tex
        assert r"\newcommand{\SZTUReviewSignatureBlank}{35mm}" in typography_tex
        assert r"\long\def\DirectoryContent" in data_tex
        if name == "normal":
            assert r"\textit{x}\textsubscript{i}" in data_tex
            assert r"\MidtermListItem{1}{（1）}" in data_tex
        if name == "long-with-image":
            assert r"H\textsubscript{2}O" in data_tex
            assert r"10\textsuperscript{−3}" in data_tex
            assert r"R\textsuperscript{2}" in data_tex
            assert r"\MidtermFigure{assets/" in data_tex
            assert "图 3-1" in data_tex
            assert "图 3-2" in data_tex
            assert "{{fig:" not in data_tex
            assert "SyntheticWorkflow2026" in data_tex
            assert len(list((latex_dir / "assets").glob("*.png"))) == 1
        if name == "nested-list":
            for depth, marker in ((1, "•"), (2, "（1）"), (3, "▪"), (4, "A.")):
                assert rf"\MidtermListItem{{{depth}}}{{{marker}}}" in data_tex
            assert r"\MidtermParagraph{列表结束后" in data_tex
            assert r"H\textsubscript{2}O" in data_tex
        if name == "structured-content":
            assert r"\begin{tblr}{width=\linewidth" in data_tex
            assert r"\begin{minipage}[t]{0.4875\linewidth}" in data_tex
            assert "图 3-1" in data_tex
            assert "（a）" in data_tex and "（b）" in data_tex
            assert "{{fig:" not in data_tex
            assert r"\frac{" in data_tex
            assert r"_{out}" in data_tex and r"_{in}" in data_tex
            assert r"\sqrt{{x}^{2}}" in data_tex
            assert r"\input" not in data_tex
            assert "(3-1)" in data_tex and "{{eq:" not in data_tex

        pdf_path = latex_dir / "main.pdf"
        info = run(["pdfinfo", str(pdf_path)], cwd=project_dir)
        assert "Page size:       595.28 x 841.89 pts (A4)" in info
        match = re.search(r"^Pages:\s+(\d+)$", info, re.MULTILINE)
        assert match is not None
        page_count = int(match.group(1))
        if name == "extreme-pagination":
            assert 7 <= page_count <= 17
        else:
            assert 2 <= page_count <= 8
        fonts = run(["pdffonts", str(pdf_path)], cwd=project_dir)
        assert "FZXBSJW" in fonts
        assert "SimSun" in fonts
        assert "TimesNewRoman" in fonts
        assert all(
            line.split()[-5:-2] == ["yes", "yes", "yes"]
            for line in fonts.splitlines()[2:]
            if line.strip()
        )
        extracted = run(["pdftotext", str(pdf_path), "-"], cwd=project_dir)
        assert "深圳技术大学本科毕业论文（设计）" in extracted
        assert raw["metadata"]["student_name"] in extracted
        assert "存在的问题及后期指导工作意见" in extracted
        gate = run(
            [sys.executable, str(project_dir / "scripts" / "validate_cjk_render.py"), str(pdf_path)],
            cwd=project_dir,
        )
        assert "PASS" in gate and "100.0%" in gate

        with pdfplumber.open(pdf_path) as pdf:
            first_page = pdf.pages[0]
            teacher_bundle_pages = {}
            student_section_pages = {}
            for page_index, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""
                for label in (
                    "毕业论文（设计）的目录和主要研究内容：",
                    "毕业论文（设计）工作进展情况（详述）：",
                ):
                    if label in page_text:
                        student_section_pages[label] = page_index
                for label in (
                    "指导教师填写栏目",
                    "存在的问题及后期指导工作意见",
                    "审查小组检查意见",
                ):
                    if label in page_text:
                        teacher_bundle_pages[label] = page_index
            assert set(teacher_bundle_pages) == {
                "指导教师填写栏目",
                "存在的问题及后期指导工作意见",
                "审查小组检查意见",
            }
            teacher_pages = [
                teacher_bundle_pages["指导教师填写栏目"],
                teacher_bundle_pages["存在的问题及后期指导工作意见"],
                teacher_bundle_pages["审查小组检查意见"],
            ]
            assert teacher_pages == sorted(teacher_pages)
            assert teacher_pages[1] - teacher_pages[0] <= 1
            assert teacher_pages[2] - teacher_pages[1] <= 1
            if name == "minimal":
                assert student_section_pages == {
                    "毕业论文（设计）的目录和主要研究内容：": 1,
                    "毕业论文（设计）工作进展情况（详述）：": 1,
                }
                full_width_tops = sorted(
                    {
                        round(float(edge["top"]), 1)
                        for edge in first_page.edges
                        if abs(float(edge["bottom"]) - float(edge["top"])) < 0.1
                        and float(edge["x1"]) - float(edge["x0"]) > 420
                    }
                )
                boundaries = []
                for top in full_width_tops:
                    if not boundaries or top - boundaries[-1] > 1.0:
                        boundaries.append(top)
                assert len(boundaries) >= 3
                progress_height_mm = (boundaries[-1] - boundaries[-2]) * 25.4 / 72.0
                assert 54.5 <= progress_height_mm <= 56.5
            chars = [char for page in pdf.pages for char in page.chars if char.get("text", "").strip()]
            assert any("FZXBSJW" in char["fontname"] and round(float(char["size"]), 2) == 18.0 for char in chars)
            assert any("SimSun" in char["fontname"] and round(float(char["size"]), 2) == 12.0 for char in chars)
            assert any("SimSun" in char["fontname"] and round(float(char["size"]), 2) == 10.5 for char in chars)
            vertical = [
                edge
                for edge in first_page.edges
                if abs(float(edge["x1"]) - float(edge["x0"])) < 0.1
                and float(edge["top"]) < 220
                and float(edge["bottom"]) > 135
            ]
            x_values = sorted({round(float(edge["x0"]), 1) for edge in vertical})
            assert x_values[0] <= 84.8 and x_values[-1] >= 510.2
            table_top = min(float(edge["top"]) for edge in vertical)
            assert abs(table_top - 134.4) <= 1.2

            for page_index, page in enumerate(pdf.pages, start=1):
                long_h = [
                    edge
                    for edge in page.edges
                    if abs(float(edge["bottom"]) - float(edge["top"])) < 0.1
                    and float(edge["x1"]) - float(edge["x0"]) > 420
                ]
                long_v = [
                    edge
                    for edge in page.edges
                    if abs(float(edge["x1"]) - float(edge["x0"])) < 0.1
                    and float(edge["bottom"]) - float(edge["top"]) > 20
                ]
                assert long_h, f"page {page_index} lacks a full-width closing rule"
                assert long_v, f"page {page_index} lacks continuous side borders"
                if page_index > 1:
                    assert min(float(edge["top"]) for edge in long_h) <= 82.0

            words_by_text = {
                word["text"]: word
                for page in pdf.pages
                for word in page.extract_words()
                if word["text"] in {"指导教师签名：", "审查小组负责人签名："}
            }
            assert set(words_by_text) == {"指导教师签名：", "审查小组负责人签名："}
            form_right = max(
                float(edge["x0"])
                for page in pdf.pages
                for edge in page.edges
                if abs(float(edge["x1"]) - float(edge["x0"])) < 0.1
            )
            mm_to_pt = 72.0 / 25.4
            teacher_gap = form_right - float(words_by_text["指导教师签名："]["x1"])
            review_gap = form_right - float(words_by_text["审查小组负责人签名："]["x1"])
            assert teacher_gap >= (45.0 + 7.5) * mm_to_pt - 2.0
            assert review_gap >= (35.0 + 7.5) * mm_to_pt - 2.0
        if name == "long-with-image":
            with pdfplumber.open(pdf_path) as pdf:
                assert any(page.images for page in pdf.pages)

    invalid = json.loads((midterm_dir / "fixtures" / "minimal.json").read_text(encoding="utf-8"))
    numbered = copy.deepcopy(invalid)
    numbered["sections"]["numbering_style"] = "chapter_section_hierarchy"
    numbered["sections"]["directory"] = [
        {"level": 1, "number": "旧编号", "title": "一级"},
        {"level": 2, "number": "旧编号", "title": "二级"},
        {"level": 3, "number": "旧编号", "title": "三级"},
        {"level": 4, "number": "旧编号", "title": "四级"},
        {"level": 5, "number": "旧编号", "title": "五级"},
        {"level": 6, "number": "旧编号", "title": "六级"},
    ]
    normalized_numbered = word_renderer.validate_data(numbered)
    assert [
        item["number"] for item in normalized_numbered["sections"]["directory"]
    ] == ["第一章", "第一节", "一、", "（一）", "1.", "(1)"]
    numbered_tex = latex_renderer.data_tex(
        normalized_numbered,
        data_dir=midterm_dir / "fixtures",
        assets_dir=output_root / "numbered-assets",
    )
    for expected_number in ("第一章", "第一节", "一、", "（一）", "1.", "(1)"):
        assert expected_number in numbered_tex

    invalid_advisor = copy.deepcopy(invalid)
    invalid_advisor["metadata"]["advisor"] = "李华教授"
    try:
        word_renderer.validate_data(invalid_advisor)
    except ValueError:
        pass
    else:
        raise AssertionError("advisor title must be rejected")
    invalid_level = copy.deepcopy(invalid)
    invalid_level["sections"]["directory"][1]["level"] = 3
    try:
        word_renderer.validate_data(invalid_level)
    except ValueError:
        pass
    else:
        raise AssertionError("skipped outline level must be rejected")
    invalid_remote = copy.deepcopy(invalid)
    invalid_remote["sections"]["progress"] = [
        {"type": "image", "path": "https://example.com/a.png", "alt": "remote"}
    ]
    normalized_remote = word_renderer.validate_data(invalid_remote)
    try:
        word_renderer._resolve_image(
            normalized_remote["sections"]["progress"][0]["path"],
            midterm_dir,
        )
    except ValueError:
        pass
    else:
        raise AssertionError("remote image URL must be rejected")

    structured_figures = json.loads(
        (midterm_dir / "fixtures" / "long-with-image.json").read_text(
            encoding="utf-8"
        )
    )
    duplicate_figure = copy.deepcopy(structured_figures)
    duplicate_figure["sections"]["progress"][2]["id"] = "fig-workflow_0001"
    try:
        word_renderer.validate_data(duplicate_figure)
    except ValueError as error:
        assert "duplicate figure id" in str(error)
    else:
        raise AssertionError("duplicate figure ids must be rejected")

    unknown_reference = copy.deepcopy(structured_figures)
    unknown_reference["sections"]["progress"][0]["runs"][0]["text"] += (
        "未知引用{{fig:fig-missing_0001}}。"
    )
    try:
        word_renderer.validate_data(unknown_reference)
    except ValueError as error:
        assert "unknown figure" in str(error)
    else:
        raise AssertionError("unknown figure references must be rejected")

    missing_caption = copy.deepcopy(structured_figures)
    del missing_caption["sections"]["main_research_content"][2]["caption"]
    try:
        word_renderer.validate_data(missing_caption)
    except ValueError as error:
        assert "caption is required" in str(error)
    else:
        raise AssertionError("new figures without captions must be rejected")

    unified_output = output_root / "unified-minimal"
    run(
        [
            sys.executable,
            str(midterm_dir / "render.py"),
            "--data",
            str(midterm_dir / "fixtures" / "minimal.json"),
            "--output-dir",
            str(unified_output),
            "--format",
            "all",
            "--overwrite",
        ],
        cwd=project_dir,
    )
    assert (unified_output / "midterm.docx").is_file()
    assert (unified_output / "latex" / "main.tex").is_file()

    print("midterm template regression checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
