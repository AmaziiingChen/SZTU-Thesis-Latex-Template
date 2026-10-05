"""Regression for compact reviewer guidance in actual DOCX and PDF output."""
import json
import os
from pathlib import Path
import subprocess
import sys
import unittest

import pdfplumber
from docx import Document
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from common.python.assessment_word import render as render_word
from common.python.assessment_latex import render_bundle


class PromptSpacingTests(unittest.TestCase):
    def test_static_and_editable_guidance_keep_compact_spacing(self):
        root = HERE.parents[1] / 'tmp/reviewer-prompt-spacing-20261005'
        root.mkdir(parents=True, exist_ok=True)
        for mode in ('static', 'editable'):
            with self.subTest(mode=mode):
                data = json.loads((HERE / 'fixtures/prompt-spacing.json').read_text(encoding='utf-8'))
                if mode == 'static':
                    data['sections'].pop('writing_guidance')
                folder = root / mode
                folder.mkdir(exist_ok=True)
                source = folder / 'input.json'
                source.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
                docx = folder / 'reviewer.docx'
                render_word('reviewer-review', HERE / 'word/official-template.docx', source, docx, overwrite=True)
                document = Document(docx)
                prompt = next(p for row in document.tables[0].rows for cell in row.cells for p in cell.paragraphs if p.text.startswith('评阅教师评语'))
                spacing = prompt._p.pPr.find(qn('w:spacing'))
                self.assertEqual(spacing.get(qn('w:line')), '364')
                self.assertEqual(spacing.get(qn('w:lineRule')), 'exact')
                render_bundle('reviewer-review', source, folder / 'latex', overwrite=True, compile_pdf=True)
                pdf = folder / 'latex/main.pdf'
                with pdfplumber.open(pdf) as rendered:
                    self.assertEqual(len(rendered.pages), 1)
                    page = rendered.pages[0]
                    first = page.search('从选题价值和难度', regex=False)[0]
                    second = page.search('水平与规范性等方面进行评述', regex=False)[0]
                    self.assertAlmostEqual(second['top'] - first['top'], 18.2, delta=0.25)
                    self.assertTrue(page.search('签名', regex=False))
                subprocess.run([sys.executable, str(HERE.parents[1] / 'scripts/validate_cjk_render.py'), str(pdf), '--dpi', '180'], check=True)


if __name__ == '__main__':
    os.environ['PYTHONUTF8'] = '1'
    unittest.main()
