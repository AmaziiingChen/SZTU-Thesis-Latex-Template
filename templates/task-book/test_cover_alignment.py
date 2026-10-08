"""Check cover title centering, baselines and underline clearance."""
import json
import math
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import pdfplumber
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = Path(__file__).resolve().parent


def render_cover(pdf_path):
    """Use the same Poppler renderer as the Chinese glyph gate."""
    with tempfile.TemporaryDirectory(prefix="task-cover-ink-") as folder:
        output = Path(folder) / "cover"
        subprocess.run([
            "pdftocairo", "-f", "1", "-l", "1", "-singlefile", "-png", "-r", "600",
            str(pdf_path), str(output),
        ], check=True, capture_output=True)
        with Image.open(output.with_suffix(".png")) as image:
            return image.convert("L")


def title_ink_clearance_mm(page, image, chars, rule):
    """Measure visible strokes, excluding the underline itself and font-box whitespace."""
    sx, sy = image.width / page.width, image.height / page.height
    left = math.floor(min(c["x0"] for c in chars) * sx)
    right = math.ceil(max(c["x1"] for c in chars) * sx)
    top = math.floor(min(c["top"] for c in chars) * sy)
    rule_edge = rule["top"] - rule["linewidth"] / 2
    bottom = math.floor(rule_edge * sy)
    ink = image.crop((left, top, right, bottom)).point(lambda p: 255 if p < 128 else 0)
    bounds = ink.getbbox()
    if bounds is None:
        raise AssertionError("cover title has no visible strokes")
    return (rule_edge - (top + bounds[3]) / sy) * 25.4 / 72


class CoverAlignmentTests(unittest.TestCase):
    def test_cover_titles(self):
        layout = json.loads((TEMPLATE / "spec/layout.json").read_text(encoding="utf-8"))
        base = json.loads((TEMPLATE / "fixtures/minimal.json").read_text(encoding="utf-8"))
        cases = {
            "single": "合成材料性能研究",
            "wrapped": "合成材料组织结构及光学性能影响机制与实验方法研究",
            "single_sub": {"type": "paragraph", "runs": [
                {"text": "陶瓷 Al"}, {"text": "2", "script": "sub"},
                {"text": "O"}, {"text": "3", "script": "sub"}, {"text": " 性能研究"},
            ]},
            "mixed": {"type": "paragraph", "runs": [
                {"text": "陶瓷助剂 ZrO"}, {"text": "2", "script": "sub"},
                {"text": " 与温度对 Al"}, {"text": "2", "script": "sub"},
                {"text": "O"}, {"text": "3", "script": "sub"},
                {"text": " 复合材料组织结构及光学性能影响的合成验证研究"},
            ]},
        }
        for name, title in cases.items():
            with self.subTest(name=name):
                output = ROOT / "tmp/task-cover-alignment-20261005" / name
                output.mkdir(parents=True, exist_ok=True)
                data = json.loads(json.dumps(base))
                data["metadata"].update(title=title, student_name="合成测试学生", student_id="TEST001", advisor="合成教师")
                fixture = output / "input.json"
                fixture.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
                result = subprocess.run([sys.executable, str(TEMPLATE / "render.py"), "--data", str(fixture),
                    "--output-dir", str(output), "--compile", "--overwrite"], capture_output=True, text=True, encoding="utf-8")
                self.assertEqual(result.returncode, 0, result.stdout[-2000:] + result.stderr[-2000:])
                title_table = Document(output / "task-book.docx").tables[0]
                self.assertEqual(len(title_table.rows), 2 if name in {"mixed", "wrapped"} else 1)
                for index, row in enumerate(title_table.rows):
                    value_cell = row.cells[1] if index == 0 else row.cells[0]
                    self.assertEqual(value_cell.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
                raster = render_cover(output / "latex/main.pdf")
                with pdfplumber.open(output / "latex/main.pdf") as pdf:
                    page = pdf.pages[0]
                    top = layout["cover"]["title_field_top_mm"] * 72 / 25.4
                    chars = [c for c in page.chars if top - 5 < c["top"] < top + 65 and c["text"].strip()]
                    label = next(c for c in chars if c["text"] == "题")
                    first = next(c for c in chars if c["text"] == ("合" if name in {"single", "wrapped"} else "陶"))
                    self.assertAlmostEqual(label["top"], first["top"], delta=0.1)
                    previous = top - 5
                    rules = sorted([e for e in page.lines if top + 10 < e["top"] < top + 65 and e["width"] > 300], key=lambda e: e["top"])
                    self.assertEqual(len(rules), 2 if name in {"mixed", "wrapped"} else 1)
                    clearances = []
                    for rule in rules:
                        row = [c for c in chars if previous <= c["top"] < rule["top"] and c["x0"] >= rule["x0"] - 0.5]
                        text_center = (min(c["x0"] for c in row) + max(c["x1"] for c in row)) / 2
                        rule_center = (rule["x0"] + rule["x1"]) / 2
                        self.assertAlmostEqual(text_center, rule_center, delta=0.5)
                        gap = title_ink_clearance_mm(page, raster, row, rule)
                        self.assertGreaterEqual(gap, layout["cover"]["title_underline_clearance_mm"])
                        self.assertLessEqual(gap, layout["cover"]["title_underline_clearance_max_mm"])
                        clearances.append(gap)
                        previous = rule["top"]
                        print(name, "underline clearance (mm):", round(gap, 3))
                    self.assertLessEqual(max(clearances) - min(clearances), layout["cover"]["title_underline_clearance_difference_max_mm"])


if __name__ == "__main__":
    unittest.main()
