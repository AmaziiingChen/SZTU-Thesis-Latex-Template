#!/usr/bin/env python3
"""Render the SZTU undergraduate thesis task book as an editable DOCX."""

from __future__ import annotations

import argparse
import importlib.util
import json
import math
import re
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image as PILImage


TASK_BOOK_DIR = Path(__file__).resolve().parents[1]
TEMPLATES_DIR = TASK_BOOK_DIR.parent
if str(TEMPLATES_DIR) not in sys.path:
    sys.path.insert(0, str(TEMPLATES_DIR))

from common.python.content import (  # noqa: E402
    display_width,
    figure_caption_runs,
    plain_runs,
    starts_with_calendar_date,
)
from common.python.font_files import font_roots, resolve_font_files  # noqa: E402
from common.python.process_form import load_process_document_layout  # noqa: E402
from common.python.equation import add_numbered_omml_table, append_omml  # noqa: E402


def _load_model():
    path = TASK_BOOK_DIR / "python" / "model.py"
    spec = importlib.util.spec_from_file_location("task_book_model", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load task-book model: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


MODEL = _load_model()
DataError = MODEL.DataError
validate_data = MODEL.validate_data
LAYOUT = load_process_document_layout(TASK_BOOK_DIR / "spec" / "layout.json")
BODY_STYLE = LAYOUT["typography"]["body"]
LIST_FIRST_LEVEL_INDENT_PT = (
    BODY_STYLE["size_pt"] * LAYOUT["paragraphs"]["list_first_level_indent_em"]
)
LIST_LEVEL_INDENT_PT = (
    BODY_STYLE["size_pt"] * LAYOUT["paragraphs"]["list_level_indent_em"]
)
LIST_HANGING_INDENT_PT = (
    BODY_STYLE["size_pt"] * LAYOUT["paragraphs"]["list_hanging_indent_em"]
)
LIST_MAX_DEPTH = LAYOUT["paragraphs"]["list_max_depth"]
UNORDERED_LIST_MARKERS = LAYOUT["paragraphs"]["unordered_list_markers"]
CHECKBOX_SYMBOL_FONT = "Wingdings 2"
CHECKBOX_SYMBOL_CODES = {
    "outline_box": "00A3",
    "outline_box_with_tick": "0052",
}
CHECKBOX_FONT_FILENAMES = {
    "wingdings 2.ttf",
    "wingdings2.ttf",
    "wingdng2.ttf",
}


def _fixed_content() -> dict[str, Any]:
    return json.loads(
        (TASK_BOOK_DIR / "spec" / "fixed-content.json").read_text(encoding="utf-8")
    )


def _require_checkbox_symbol_font() -> Path:
    """Reject output that would silently substitute the checkbox symbol font."""

    for root in font_roots():
        for path in root.rglob("*"):
            if path.is_file() and path.name.casefold() in CHECKBOX_FONT_FILENAMES:
                return path.resolve()
    raise FileNotFoundError(
        "required checkbox font is missing; install Wingdings 2 "
        "(Wingdings 2.ttf/Wingdng2.ttf) before rendering"
    )


def _remove_all_body_content(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_page_system(document: Document) -> None:
    section = document.sections[0]
    page = LAYOUT["page"]
    section.page_width = Mm(page["width_mm"])
    section.page_height = Mm(page["height_mm"])
    section.top_margin = Mm(page["top_margin_mm"])
    section.bottom_margin = Mm(page["bottom_margin_mm"])
    section.left_margin = Mm(page["left_margin_mm"])
    section.right_margin = Mm(page["right_margin_mm"])
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)


def _set_run_font(run, *, style: dict[str, Any], bold: bool | None = None) -> None:
    text = run.text or " "
    latin = all(ord(char) < 128 for char in text)
    run.font.name = style["latin_word_family"] if latin else style["cjk_word_family"]
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), style["latin_word_family"])
    rfonts.set(qn("w:hAnsi"), style["latin_word_family"])
    rfonts.set(qn("w:cs"), style["latin_word_family"])
    rfonts.set(qn("w:eastAsia"), style["cjk_word_family"])
    run.font.size = Pt(style["size_pt"])
    run.bold = style["bold"] if bold is None else bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _append_runs(paragraph, runs: list[dict[str, Any]], *, style: dict[str, Any]) -> None:
    for rich_run in runs:
        for piece in re.findall(r"[\x00-\x7f]+|[^\x00-\x7f]+", rich_run["text"]):
            run = paragraph.add_run(piece)
            _set_run_font(
                run,
                style=style,
                bold=bool(style.get("bold", False) or rich_run.get("bold", False)),
            )
            run.italic = rich_run.get("italic", False)
            script = rich_run.get("script", "normal")
            run.font.subscript = script == "sub"
            run.font.superscript = script == "super"


def _split_runs_for_cover_title(
    runs: list[dict[str, Any]], *, line_capacity: float, max_lines: int
) -> list[list[dict[str, Any]]]:
    """Split the official cover title into at most two style-preserving lines."""

    if line_capacity <= 0:
        raise RuntimeError("cover.title_line_capacity_units must be positive")
    if max_lines <= 0:
        raise RuntimeError("cover.title_max_lines must be positive")
    lines: list[list[dict[str, Any]]] = [[]]
    line_width = 0.0
    for source_run in runs:
        for char in source_run["text"]:
            if char == "\n":
                if not lines[-1]:
                    continue
                if len(lines) == max_lines:
                    raise DataError(
                        "metadata.title exceeds the official cover line capacity"
                    )
                lines.append([])
                line_width = 0.0
                continue
            width = display_width(char)
            if line_width + width > line_capacity and lines[-1]:
                if len(lines) == max_lines:
                    raise DataError(
                        "metadata.title exceeds the official cover line capacity"
                    )
                lines.append([])
                line_width = 0.0
            if lines[-1] and all(
                lines[-1][-1].get(key) == source_run.get(key)
                for key in ("script", "italic", "bold")
            ):
                lines[-1][-1]["text"] += char
            else:
                copied = dict(source_run)
                copied["text"] = char
                lines[-1].append(copied)
            line_width += width
    return [line for line in lines if line]


def _format_paragraph(
    paragraph,
    *,
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    line_spacing: float = 1.0,
    first_line_indent_pt: float | None = None,
    left_indent_pt: float | None = None,
    space_before_pt: float = 0,
    space_after_pt: float = 0,
) -> None:
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before_pt)
    fmt.space_after = Pt(space_after_pt)
    fmt.line_spacing = line_spacing
    fmt.first_line_indent = Pt(first_line_indent_pt) if first_line_indent_pt is not None else None
    fmt.left_indent = Pt(left_indent_pt) if left_indent_pt is not None else None


def _paragraph_alignment(value: str, *, path: str) -> WD_ALIGN_PARAGRAPH:
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    try:
        return mapping[value]
    except KeyError as exc:
        raise RuntimeError(f"{path} must be one of {sorted(mapping)}") from exc


def _set_auto_line_spacing_twips(paragraph, line_twips: int) -> None:
    """Write an explicit Word/WPS auto-line-spacing value."""

    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:line"), str(line_twips))
    spacing.set(qn("w:lineRule"), "auto")


def _paragraph(
    container,
    runs: str | list[dict[str, Any]],
    *,
    style: dict[str, Any],
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    line_spacing: float = 1.0,
    first_line_indent_pt: float | None = None,
    left_indent_pt: float | None = None,
    space_before_pt: float = 0,
    space_after_pt: float = 0,
):
    paragraph = container.add_paragraph()
    _format_paragraph(
        paragraph,
        alignment=alignment,
        line_spacing=line_spacing,
        first_line_indent_pt=first_line_indent_pt,
        left_indent_pt=left_indent_pt,
        space_before_pt=space_before_pt,
        space_after_pt=space_after_pt,
    )
    _append_runs(paragraph, plain_runs(runs) if isinstance(runs, str) else runs, style=style)
    return paragraph


def _clear_cell(cell) -> None:
    for paragraph in list(cell.paragraphs)[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    cell.paragraphs[0].clear()


def _twips(mm: float) -> str:
    return str(round(mm / 25.4 * 1440))


def _set_cell_margins(cell, *, left: float, right: float, top: float, bottom: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("left", left), ("right", right), ("top", top), ("bottom", bottom)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), _twips(value))
        element.set(qn("w:type"), "dxa")


def _set_cell_border(cell, **edges: dict[str, Any]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, values in edges.items():
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        for key, value in values.items():
            element.set(qn(f"w:{key}"), str(value))


def _set_table_borders(table, *, size_eighth_pt: int) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size_eighth_pt))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def _remove_table_borders(table) -> None:
    _set_table_borders(table, size_eighth_pt=0)
    borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
    for element in borders:
        element.set(qn("w:val"), "nil")


def _set_table_geometry(
    table,
    widths_mm: list[float] | tuple[float, ...],
    *,
    alignment: WD_TABLE_ALIGNMENT = WD_TABLE_ALIGNMENT.CENTER,
) -> None:
    """Set tblW, tblGrid and every tcW so Word/WPS cannot recalculate widths."""

    if any(len(row.cells) != len(widths_mm) for row in table.rows):
        raise ValueError("table geometry width count must match every row")
    width_mm = sum(widths_mm)
    table.autofit = False
    table.alignment = alignment
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), _twips(width_mm))
    tbl_w.set(qn("w:type"), "dxa")
    old_grid = table._tbl.tblGrid
    if old_grid is not None:
        table._tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_mm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), _twips(width))
        grid.append(grid_col)
    table._tbl.insert(table._tbl.index(tbl_pr) + 1, grid)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), _twips(width))
            tc_w.set(qn("w:type"), "dxa")


def _set_table_indent(table, width_mm: float) -> None:
    tbl_pr = table._tbl.tblPr
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), _twips(width_mm))
    indent.set(qn("w:type"), "dxa")


def _set_row_height(row, height_mm: float) -> None:
    row.height = Mm(height_mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _set_row_exact_height(row, height_mm: float) -> None:
    row.height = Mm(height_mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY


def _set_row_cant_split(row, enabled: bool) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:cantSplit"))
    if enabled and existing is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    elif not enabled and existing is not None:
        tr_pr.remove(existing)


def _add_page_break(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _add_cover(document: Document, data: dict[str, Any], fixed: dict[str, Any]) -> None:
    styles = LAYOUT["typography"]
    metadata = data["metadata"]
    cover = fixed["cover"]
    _paragraph(
        document,
        f"{cover['document_number_label']}{metadata['student_id']}",
        style=styles["cover_number"],
    )
    _paragraph(
        document,
        cover["school_name"],
        style=styles["school_name"],
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before_pt=float(
            LAYOUT["cover"].get("school_name_space_before_pt", 58.0)
        ),
        space_after_pt=float(
            LAYOUT["cover"].get("school_name_space_after_pt", 8.0)
        ),
    )
    _paragraph(
        document,
        cover["document_title"],
        style=styles["document_title"],
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_after_pt=float(
            LAYOUT["cover"].get("document_title_space_after_pt", 10.0)
        ),
    )
    _paragraph(
        document,
        cover["cohort_format"].format(graduation_year=metadata["graduation_year"]),
        style=styles["cohort"],
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_after_pt=float(
            LAYOUT["cover"].get("cohort_space_after_pt", 38.0)
        ),
    )

    title_style = styles["cover_title_value"]
    title_em_mm = float(title_style["size_pt"]) * 25.4 / 72.0
    title_lines = _split_runs_for_cover_title(
        metadata["title"],
        line_capacity=float(
            LAYOUT["cover"].get(
                "title_line_capacity_units",
                float(LAYOUT["cover"]["title_underline_width_mm"]) / title_em_mm,
            )
        ),
        max_lines=int(LAYOUT["cover"]["title_max_lines"]),
    )
    if not 1 <= len(title_lines) <= int(LAYOUT["cover"]["title_max_lines"]):
        raise DataError("metadata.title exceeds the official cover line capacity")
    title_table = document.add_table(rows=len(title_lines), cols=2)
    title_underline = LAYOUT["cover"]["title_underline_width_mm"]
    title_label = LAYOUT["cover"]["content_width_mm"] - title_underline
    _set_table_geometry(title_table, (title_label, title_underline))
    _remove_table_borders(title_table)
    title_line_height = float(LAYOUT["cover"].get("title_line_height_mm", 9.0))
    title_cell_vertical_padding = float(
        LAYOUT["cover"].get("title_cell_vertical_padding_mm", 0.5)
    )
    for index, (row, line_runs) in enumerate(zip(title_table.rows, title_lines)):
        for cell in row.cells:
            _set_cell_margins(
                cell,
                left=1.5,
                right=1.5,
                top=title_cell_vertical_padding,
                bottom=title_cell_vertical_padding,
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if index == 0:
            _clear_cell(row.cells[0])
            label = row.cells[0].paragraphs[0]
            _format_paragraph(label, alignment=WD_ALIGN_PARAGRAPH.LEFT)
            _append_runs(
                label,
                plain_runs(cover["field_labels"]["title"]),
                style=styles["cover_label"],
            )
            value_cell = row.cells[1]
        else:
            value_cell = row.cells[0].merge(row.cells[1])
        _clear_cell(value_cell)
        value = value_cell.paragraphs[0]
        value_alignment = WD_ALIGN_PARAGRAPH.CENTER
        if index > 0:
            value_alignment = _paragraph_alignment(
                str(
                    LAYOUT["cover"].get(
                        "title_second_line_alignment",
                        "center",
                    )
                ),
                path="cover.title_second_line_alignment",
            )
        _format_paragraph(value, alignment=value_alignment)
        _append_runs(value, line_runs, style=title_style)
        _set_cell_border(
            value_cell,
            bottom={"val": "single", "sz": 4, "space": 0, "color": "000000"},
        )
        _set_row_height(row, title_line_height)

    spacer = document.add_paragraph()
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.paragraph_format.space_after = Pt(
        float(
            LAYOUT["cover"].get(
                "metadata_space_after_pt_one_line"
                if len(title_lines) == 1
                else "metadata_space_after_pt_two_lines",
                (4 if len(title_lines) == 1 else 3)
                * float(title_style["size_pt"]),
            )
        )
    )

    info_table = document.add_table(rows=4, cols=5)
    short_width = LAYOUT["cover"]["short_field_underline_width_mm"]
    gap_width = LAYOUT["cover"]["metadata_column_gap_mm"]
    label_width = (
        LAYOUT["cover"]["content_width_mm"] - 2 * short_width - gap_width
    ) / 2
    widths = (label_width, short_width, gap_width, label_width, short_width)
    _set_table_geometry(info_table, widths)
    _remove_table_borders(info_table)
    metadata_row_height = float(LAYOUT["cover"].get("metadata_row_height_mm", 14.0))
    for row in info_table.rows:
        _set_row_exact_height(row, metadata_row_height)
        for index, cell in enumerate(row.cells):
            side_padding = 0.2 if index in {0, 3} else 1.0
            _set_cell_margins(cell, left=side_padding, right=side_padding, top=1.0, bottom=1.0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    field_labels = cover["field_labels"]
    full_width_rows = (
        (field_labels["college"], metadata["college"]),
        (field_labels["major"], metadata["major"]),
    )
    for row, (label_text, value_text) in zip(info_table.rows[:2], full_width_rows):
        label_cell = row.cells[0]
        value_cell = row.cells[1].merge(row.cells[4])
        _clear_cell(label_cell)
        label = label_cell.paragraphs[0]
        _format_paragraph(label, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        _append_runs(label, plain_runs(label_text), style=styles["cover_label"])
        _clear_cell(value_cell)
        value = value_cell.paragraphs[0]
        _format_paragraph(value, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        _append_runs(value, plain_runs(value_text), style=styles["cover_value"])
        _set_cell_border(
            value_cell,
            bottom={"val": "single", "sz": 4, "space": 0, "color": "000000"},
        )

    paired_rows = (
        (field_labels["class_name"], metadata["class_name"], field_labels["student_id"], metadata["student_id"]),
        (field_labels["student_name"], metadata["student_name"], field_labels["advisor"], metadata["advisor"]),
    )
    for row, values in zip(info_table.rows[2:], paired_rows):
        mapped = ((0, values[0], True), (1, values[1], False), (3, values[2], True), (4, values[3], False))
        _clear_cell(row.cells[2])
        for index, value, is_label in mapped:
            _clear_cell(row.cells[index])
            paragraph = row.cells[index].paragraphs[0]
            _format_paragraph(
                paragraph,
                alignment=WD_ALIGN_PARAGRAPH.LEFT if is_label else WD_ALIGN_PARAGRAPH.CENTER,
            )
            _append_runs(
                paragraph,
                plain_runs(value),
                style=styles["cover_label"] if is_label else styles["cover_value"],
            )
            if not is_label:
                _set_cell_border(
                    row.cells[index],
                    bottom={"val": "single", "sz": 4, "space": 0, "color": "000000"},
                )


def _add_notice(document: Document, fixed: dict[str, Any]) -> None:
    styles = LAYOUT["typography"]
    notice = fixed["notice"]
    notice_layout = LAYOUT.get("notice", {})
    top_skip = float(notice_layout.get("page_top_skip_mm", 13.5))
    spacer = document.add_paragraph()
    spacer.paragraph_format.page_break_before = True
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Mm(max(0.0, top_skip))
    _paragraph(
        document,
        notice["title"],
        style=styles["notice_title"],
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=LAYOUT["paragraphs"]["notice_line_spacing"],
        space_before_pt=0,
        space_after_pt=float(notice_layout.get("title_space_after_pt", 39.25)),
    )
    notices = notice["items"]
    if len(notices) != 6:
        raise RuntimeError("fixed task-book notice must contain exactly six items")
    for item in notices:
        paragraph = document.add_paragraph()
        hanging = float(LAYOUT.get("notice", {}).get("hanging_indent_twip", 403)) / 20.0
        _format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            line_spacing=LAYOUT["paragraphs"]["notice_line_spacing"],
            left_indent_pt=hanging,
            first_line_indent_pt=-hanging,
            space_after_pt=float(notice_layout.get("item_space_after_pt", 0.0)),
        )
        notice_line_twips = round(
            float(styles["notice_body"]["size_pt"])
            * 20
            * float(LAYOUT["paragraphs"]["notice_line_spacing"])
        )
        _set_auto_line_spacing_twips(paragraph, notice_line_twips)
        indent = paragraph._p.get_or_add_pPr().find(qn("w:ind"))
        if indent is None:
            raise RuntimeError("notice hanging indent was not created")
        indent.set(
            qn("w:hangingChars"),
            str(int(notice_layout.get("hanging_chars_hundredth", 168))),
        )
        _append_runs(
            paragraph,
            plain_runs(f"{item['marker']}{item['text']}"),
            style=styles["notice_body"],
        )
    _add_page_break(document)


def _resolve_image(path_text: str, data_dir: Path) -> Path:
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", path_text):
        raise DataError("remote image URLs are forbidden; use a local file")
    path = Path(path_text).expanduser()
    if not path.is_absolute():
        path = data_dir / path
    path = path.resolve()
    if not path.is_file():
        raise DataError(f"image file does not exist: {path}")
    return path


def _set_picture_alt(run, alt: str) -> None:
    items = run._r.xpath(".//wp:docPr")
    if items:
        items[0].set("descr", alt)


def _fit_image_dimensions(
    image_path: Path,
    requested_width_mm: float,
    *,
    max_height_mm: float,
) -> tuple[float, float]:
    with PILImage.open(image_path) as image:
        pixel_width, pixel_height = image.size
    if pixel_width <= 0 or pixel_height <= 0:
        raise DataError("image dimensions must be positive")
    width_mm = float(requested_width_mm)
    height_mm = width_mm * pixel_height / pixel_width
    if height_mm > max_height_mm:
        scale = max_height_mm / height_mm
        width_mm *= scale
        height_mm = max_height_mm
    return width_mm, height_mm


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _block_alignment(value: str):
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[value]


def _append_data_table(cell, block: dict[str, Any]) -> None:
    config = LAYOUT["embedded_table"]
    weights = [column["width_weight"] for column in block["columns"]]
    total_weight = sum(weights)
    widths = [float(config["width_mm"]) * weight / total_weight for weight in weights]
    table = cell.add_table(rows=1 + len(block["rows"]), cols=len(block["columns"]))
    _set_table_geometry(table, widths)
    _set_table_borders(table, size_eighth_pt=round(float(config["border_pt"]) * 8))
    _set_repeat_table_header(table.rows[0])
    for row in table.rows:
        _set_row_cant_split(row, True)
        for target in row.cells:
            _set_cell_margins(
                target,
                left=float(config["cell_padding_mm"]),
                right=float(config["cell_padding_mm"]),
                top=float(config["cell_padding_mm"]),
                bottom=float(config["cell_padding_mm"]),
            )
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    style = LAYOUT["typography"]["body"]
    for column_index, column in enumerate(block["columns"]):
        paragraph = table.rows[0].cells[column_index].paragraphs[0]
        _format_paragraph(paragraph, alignment=_block_alignment(column["alignment"]))
        _append_runs(paragraph, column["header_runs"], style=style)
        for run in paragraph.runs:
            run.bold = True
    for row_index, values in enumerate(block["rows"], start=1):
        for column_index, runs in enumerate(values):
            paragraph = table.rows[row_index].cells[column_index].paragraphs[0]
            _format_paragraph(
                paragraph,
                alignment=_block_alignment(block["columns"][column_index]["alignment"]),
            )
            _append_runs(paragraph, runs, style=style)
    if block["caption_runs"]:
        caption = cell.add_paragraph()
        _format_paragraph(
            caption,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            space_before_pt=float(config["caption_space_before_pt"]),
            space_after_pt=float(config["caption_space_after_pt"]),
        )
        caption.paragraph_format.keep_together = True
        _append_runs(caption, block["caption_runs"], style=style)


def _append_figure_group(cell, block: dict[str, Any], *, data_dir: Path) -> None:
    config = LAYOUT["figure_group"]
    count = len(block["items"])
    total_width = float(config["width_mm"])
    gap = float(config["column_gap_mm"])
    item_width = (total_width - gap * (count - 1)) / count
    table = cell.add_table(rows=1, cols=count)
    _set_table_geometry(table, [item_width] * count)
    _remove_table_borders(table)
    _set_row_cant_split(table.rows[0], True)
    style = LAYOUT["typography"]["body"]
    for index, item in enumerate(block["items"]):
        target = table.rows[0].cells[index]
        _set_cell_margins(target, left=gap / 2, right=gap / 2, top=0, bottom=0)
        target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = target.paragraphs[0]
        _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run()
        _set_run_font(run, style=style)
        image_path = _resolve_image(item["path"], data_dir)
        width_mm, height_mm = _fit_image_dimensions(
            image_path,
            item_width - gap,
            max_height_mm=float(config["max_item_height_mm"]),
        )
        run.add_picture(str(image_path), width=Mm(width_mm), height=Mm(height_mm))
        _set_picture_alt(run, item["alt"])
        subcaption = target.add_paragraph()
        _format_paragraph(subcaption, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        subcaption.paragraph_format.keep_together = True
        _append_runs(subcaption, plain_runs(item["subfigure_label"]), style=style)
        if item["caption_runs"]:
            _append_runs(subcaption, item["caption_runs"], style=style)
    caption = cell.add_paragraph()
    _format_paragraph(
        caption,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before_pt=float(LAYOUT["paragraphs"]["caption_space_before_pt"]),
        space_after_pt=float(LAYOUT["paragraphs"]["caption_space_after_pt"]),
    )
    caption.paragraph_format.keep_together = True
    _append_runs(caption, figure_caption_runs(block), style=style)


def _list_marker(block_type: str, item: dict[str, Any], index: int, depth: int) -> str:
    if block_type == "ordered_list":
        return item["marker"] or f"（{index}）"
    return UNORDERED_LIST_MARKERS[depth - 1]


def _append_list_block(cell, block: dict[str, Any], *, depth: int = 1) -> None:
    if not 1 <= depth <= LIST_MAX_DEPTH:
        raise AssertionError(f"normalized list depth escaped bounds: {depth}")
    style = LAYOUT["typography"]["body"]
    marker_indent = LIST_FIRST_LEVEL_INDENT_PT + (depth - 1) * LIST_LEVEL_INDENT_PT
    for index, item in enumerate(block["items"], start=1):
        paragraph = cell.add_paragraph()
        _format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            left_indent_pt=marker_indent + LIST_HANGING_INDENT_PT,
            first_line_indent_pt=-LIST_HANGING_INDENT_PT,
        )
        marker = _list_marker(block["type"], item, index, depth)
        marker_text = marker if block["type"] == "ordered_list" else f"{marker} "
        _append_runs(paragraph, plain_runs(marker_text), style=style)
        _append_runs(paragraph, item["runs"], style=style)
        if item["children"]:
            _append_list_block(cell, item["children"], depth=depth + 1)


def _append_content_blocks(cell, blocks: list[dict[str, Any]], *, data_dir: Path) -> None:
    style = LAYOUT["typography"]["body"]
    indent = style["size_pt"] * LAYOUT["paragraphs"]["first_line_indent_em"]
    for block in blocks:
        if block["type"] == "paragraph":
            paragraph = cell.add_paragraph()
            _format_paragraph(
                paragraph,
                alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                first_line_indent_pt=indent,
            )
            _append_runs(paragraph, block["runs"], style=style)
        elif block["type"] in {"ordered_list", "unordered_list"}:
            _append_list_block(cell, block)
        elif block["type"] == "image":
            paragraph = cell.add_paragraph()
            _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            run = paragraph.add_run()
            _set_run_font(run, style=style)
            path = _resolve_image(block["path"], data_dir)
            run.add_picture(str(path), width=Mm(block["width_mm"]))
            _set_picture_alt(run, block["alt"])
            if block.get("caption_runs"):
                caption = cell.add_paragraph()
                _format_paragraph(caption, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before_pt=3, space_after_pt=3)
                _append_runs(caption, block["caption_runs"], style=style)
        elif block["type"] == "data_table":
            _append_data_table(cell, block)
        elif block["type"] == "figure_group":
            _append_figure_group(cell, block, data_dir=data_dir)
        elif block["type"] == "equation":
            label = block.get("equation_label")
            if label is not None:
                _, paragraphs = add_numbered_omml_table(
                    cell,
                    block["expression"],
                    label,
                    width_mm=float(LAYOUT["equation"]["width_mm"]),
                )
                for equation_paragraph in paragraphs:
                    _format_paragraph(
                        equation_paragraph,
                        alignment=equation_paragraph.alignment,
                    )
                    equation_paragraph.paragraph_format.keep_together = True
                center = paragraphs[1]
                center.paragraph_format.space_before = Pt(
                    LAYOUT["equation"]["space_before_pt"]
                )
                center.paragraph_format.space_after = Pt(
                    LAYOUT["equation"]["space_after_pt"]
                )
                fallback = center.add_run(block["alt"])
                _set_run_font(fallback, style=style)
                fallback.font.hidden = True
                _set_run_font(paragraphs[2].runs[0], style=style)
                continue
            paragraph = cell.add_paragraph()
            _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            paragraph.paragraph_format.space_before = Pt(
                LAYOUT["equation"]["space_before_pt"]
            )
            paragraph.paragraph_format.space_after = Pt(
                LAYOUT["equation"]["space_after_pt"]
            )
            paragraph.paragraph_format.keep_together = True
            fallback = paragraph.add_run(block["alt"])
            _set_run_font(fallback, style=style)
            fallback.font.hidden = True
            append_omml(paragraph, block["expression"])
        else:
            raise AssertionError(f"unsupported content block: {block['type']}")


def _heading(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    _format_paragraph(paragraph)
    _append_runs(paragraph, plain_runs(text), style=LAYOUT["typography"]["section_title"])


def _append_schedule(cell, schedule: list[dict[str, Any]]) -> None:
    style = LAYOUT["typography"]["body"]
    indent = style["size_pt"] * LAYOUT["paragraphs"]["first_line_indent_em"]
    for item in schedule:
        paragraph = cell.add_paragraph()
        period = item["period"]
        _format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line_indent_pt=0 if starts_with_calendar_date(period) else indent,
        )
        _append_runs(paragraph, plain_runs(f"{period}："), style=style)
        _append_runs(paragraph, item["content"], style=style)


def _append_references(cell, references: list[list[dict[str, Any]]]) -> None:
    style = LAYOUT["typography"]["body"]
    for runs in references:
        paragraph = cell.add_paragraph()
        _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        _append_runs(paragraph, runs, style=style)


def _append_checkbox(
    paragraph,
    *,
    selected: bool,
    fixed_topic: dict[str, Any],
    style: dict[str, Any],
) -> None:
    """Append a monochrome hollow checkbox using deterministic font glyphs.

    The current official WPS-authored forms use ``w:sym`` for checkbox glyphs.
    Wingdings 2 code 00A3 is an empty outline and 0052 is the same outline with
    a check mark, so both states retain identical metrics without emoji fallback.
    """

    token = (
        fixed_topic["checked_mark"]
        if selected
        else fixed_topic["unchecked_box"]
    )
    try:
        symbol_code = CHECKBOX_SYMBOL_CODES[token]
    except KeyError as exc:
        raise RuntimeError(f"unsupported task-book checkbox token: {token!r}") from exc
    run = paragraph.add_run()
    _set_run_font(run, style=style)
    symbol = OxmlElement("w:sym")
    symbol.set(qn("w:font"), CHECKBOX_SYMBOL_FONT)
    symbol.set(qn("w:char"), symbol_code)
    run._r.append(symbol)


def _append_checkbox_option(
    paragraph,
    label: str,
    *,
    selected: bool,
    fixed_topic: dict[str, Any],
    style: dict[str, Any],
) -> None:
    _append_runs(paragraph, plain_runs(label), style=style)
    _append_checkbox(
        paragraph,
        selected=selected,
        fixed_topic=fixed_topic,
        style=style,
    )


def _add_inline_table(
    cell,
    widths_mm: tuple[float, ...],
    *,
    alignment: WD_TABLE_ALIGNMENT = WD_TABLE_ALIGNMENT.LEFT,
):
    table = cell.add_table(rows=1, cols=len(widths_mm))
    _set_table_geometry(table, widths_mm, alignment=alignment)
    _remove_table_borders(table)
    for nested_cell in table.rows[0].cells:
        _set_cell_margins(nested_cell, left=0, right=0, top=0, bottom=0)
        nested_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        _clear_cell(nested_cell)
    return table


def _write_inline_cell(
    cell,
    runs: str | list[dict[str, Any]],
    *,
    style: dict[str, Any],
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    paragraph = cell.paragraphs[0]
    _format_paragraph(paragraph, alignment=alignment)
    _append_runs(
        paragraph,
        plain_runs(runs) if isinstance(runs, str) else runs,
        style=style,
    )


def _underline_cell(cell) -> None:
    _set_cell_border(
        cell,
        bottom={"val": "single", "sz": 4, "space": 0, "color": "000000"},
    )


def _add_signature_line(cell, label: str, *, blank_width_mm: float) -> None:
    table = cell.add_table(rows=1, cols=3)
    table.autofit = False
    label_width = 42.0
    right_inset = LAYOUT["signature"]["right_inset_mm"]
    _set_table_geometry(
        table,
        (label_width, blank_width_mm, right_inset),
        alignment=WD_TABLE_ALIGNMENT.RIGHT,
    )
    _remove_table_borders(table)
    table.rows[0].cells[0].width = Mm(label_width)
    table.rows[0].cells[1].width = Mm(blank_width_mm)
    table.rows[0].cells[2].width = Mm(right_inset)
    for item in table.rows[0].cells:
        _set_cell_margins(item, left=0, right=0, top=0, bottom=0)
        item.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    _clear_cell(table.rows[0].cells[0])
    paragraph = table.rows[0].cells[0].paragraphs[0]
    _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _append_runs(paragraph, plain_runs(label), style=LAYOUT["typography"]["signature"])
    _clear_cell(table.rows[0].cells[1])
    _clear_cell(table.rows[0].cells[2])
    _underline_cell(table.rows[0].cells[1])


def _topic_lines(cell, topic: dict[str, Any], fixed_topic: dict[str, Any], signature_labels: dict[str, str]) -> None:
    style = LAYOUT["typography"]["topic_body"]
    nature = topic["nature"]
    source = topic["source"]
    nature_options = fixed_topic["nature_options"]
    source_options = fixed_topic["source_options"]
    levels = fixed_topic["research_project_levels"]
    proposer_options = fixed_topic["self_proposed_by_options"]

    usable_width = float(LAYOUT["table"]["width_mm"]) - 2 * float(
        LAYOUT["table"]["horizontal_padding_mm"]
    )
    right_inset = float(LAYOUT["signature"]["right_inset_mm"])
    point_to_mm = 25.4 / 72.0
    project_indent_em = float(
        LAYOUT["signature"].get("topic_project_left_indent_em", 5.0)
    )
    option_indent_em = float(
        LAYOUT["signature"].get("topic_option_left_indent_em", 6.0)
    )
    child_indent_em = float(
        LAYOUT["signature"].get("self_proposed_child_left_indent_em", 7.0833)
    )
    project_indent_mm = style["size_pt"] * project_indent_em * point_to_mm
    option_indent_mm = style["size_pt"] * option_indent_em * point_to_mm
    option_indent_pt = style["size_pt"] * option_indent_em
    child_indent_pt = style["size_pt"] * child_indent_em

    nature_table = _add_inline_table(
        cell,
        (option_indent_mm, usable_width - option_indent_mm),
    )
    _write_inline_cell(
        nature_table.rows[0].cells[0],
        fixed_topic["nature_label"],
        style=style,
    )
    nature_line = nature_table.rows[0].cells[1].paragraphs[0]
    _format_paragraph(nature_line)
    _append_runs(
        nature_line,
        plain_runs(nature_options["graduation_design"]),
        style=style,
    )
    _append_checkbox(
        nature_line,
        selected=nature == "graduation_design",
        fixed_topic=fixed_topic,
        style=style,
    )
    _append_checkbox_option(
        nature_line,
        f"    {nature_options['graduation_thesis']}",
        selected=nature == "graduation_thesis",
        fixed_topic=fixed_topic,
        style=style,
    )

    source_type = source["type"]
    research = source_type == "research_project"
    project = source.get("research_project", {})
    level = project.get("level")
    other_label = project.get("other_level", "")

    other_width = float(LAYOUT["signature"].get("topic_other_blank_width_mm", 32.0))
    research_options_width = (
        usable_width - option_indent_mm - other_width - right_inset
    )
    research_table = _add_inline_table(
        cell,
        (option_indent_mm, research_options_width, other_width, right_inset),
    )
    _write_inline_cell(
        research_table.rows[0].cells[0],
        fixed_topic["source_label"],
        style=style,
    )
    research_options = research_table.rows[0].cells[1].paragraphs[0]
    _format_paragraph(research_options)
    _append_runs(
        research_options,
        plain_runs(
            f"{source_options['research_project']}"
            f"    {levels['national']}"
        ),
        style=style,
    )
    _append_checkbox(
        research_options,
        selected=research and level == "national",
        fixed_topic=fixed_topic,
        style=style,
    )
    _append_checkbox_option(
        research_options,
        f"    {levels['provincial_ministerial']}",
        selected=research and level == "provincial_ministerial",
        fixed_topic=fixed_topic,
        style=style,
    )
    _append_runs(research_options, plain_runs(f"    {levels['other']}"), style=style)
    _write_inline_cell(
        research_table.rows[0].cells[2],
        other_label if research and level == "other" else "",
        style=style,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _underline_cell(research_table.rows[0].cells[2])

    project_label_width = float(
        LAYOUT["signature"].get("topic_project_label_width_mm", 25.0)
    )
    project_blank_width = float(
        LAYOUT["signature"].get("topic_project_blank_width_mm", 65.0)
    )
    project_trailing_width = (
        usable_width - project_indent_mm - project_label_width - project_blank_width
    )
    if project_trailing_width < right_inset:
        raise ValueError(
            "task-book topic project-number field exceeds the usable table width"
        )
    project_table = _add_inline_table(
        cell,
        (
            project_indent_mm,
            project_label_width,
            project_blank_width,
            project_trailing_width,
        ),
    )
    _write_inline_cell(
        project_table.rows[0].cells[1],
        fixed_topic["project_number_label"],
        style=style,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _write_inline_cell(
        project_table.rows[0].cells[2],
        project.get("project_number", "") if research else "",
        style=style,
        alignment=_paragraph_alignment(
            str(
                LAYOUT["signature"].get(
                    "topic_project_value_alignment",
                    "center",
                )
            ),
            path="signature.topic_project_value_alignment",
        ),
    )
    _underline_cell(project_table.rows[0].cells[2])

    self_alignment = _paragraph_alignment(
        str(LAYOUT["signature"].get("self_proposed_alignment", "left")),
        path="signature.self_proposed_alignment",
    )
    practice = cell.add_paragraph()
    _format_paragraph(
        practice,
        alignment=self_alignment,
        left_indent_pt=option_indent_pt,
    )
    _append_checkbox_option(
        practice,
        source_options["practice_project"],
        selected=source_type == "practice_project",
        fixed_topic=fixed_topic,
        style=style,
    )

    self_proposed = source_type == "self_proposed"
    proposer = source.get("self_proposed_by")
    self_heading = cell.add_paragraph()
    _format_paragraph(
        self_heading,
        alignment=self_alignment,
        left_indent_pt=option_indent_pt,
    )
    _append_runs(
        self_heading,
        plain_runs(source_options["self_proposed"]),
        style=style,
    )

    for name in ("teacher", "student"):
        paragraph = cell.add_paragraph()
        _format_paragraph(
            paragraph,
            alignment=self_alignment,
            left_indent_pt=child_indent_pt,
        )
        _append_checkbox_option(
            paragraph,
            proposer_options[name],
            selected=self_proposed and proposer == name,
            fixed_topic=fixed_topic,
            style=style,
        )

    signature_label_width = 42.0
    signature_blank_width = float(LAYOUT["signature"]["advisor_blank_width_mm"])
    option_width = (
        usable_width - signature_label_width - signature_blank_width - right_inset
    )
    final_table = _add_inline_table(
        cell,
        (option_width, signature_label_width, signature_blank_width, right_inset),
    )
    joint = final_table.rows[0].cells[0].paragraphs[0]
    _format_paragraph(
        joint,
        alignment=self_alignment,
        left_indent_pt=child_indent_pt,
    )
    _append_checkbox_option(
        joint,
        proposer_options["teacher_student_joint"],
        selected=self_proposed and proposer == "teacher_student_joint",
        fixed_topic=fixed_topic,
        style=style,
    )
    _write_inline_cell(
        final_table.rows[0].cells[1],
        signature_labels["advisor"],
        style=LAYOUT["typography"]["signature"],
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    _underline_cell(final_table.rows[0].cells[2])


def _add_date_line(cell, signature_labels: dict[str, Any]) -> None:
    style = LAYOUT["typography"]["notice_body"]
    right_inset = float(LAYOUT["signature"]["right_inset_mm"])
    year_blank = float(LAYOUT["signature"].get("date_year_blank_width_mm", 15.0))
    month_blank = float(LAYOUT["signature"].get("date_month_blank_width_mm", 10.0))
    day_blank = float(LAYOUT["signature"].get("date_day_blank_width_mm", 10.0))
    label_width = float(LAYOUT["signature"].get("date_label_width_mm", 5.0))
    date_table = _add_inline_table(
        cell,
        (
            year_blank,
            label_width,
            month_blank,
            label_width,
            day_blank,
            label_width,
            right_inset,
        ),
        alignment=WD_TABLE_ALIGNMENT.RIGHT,
    )
    date_labels = signature_labels.get("date_labels", {})
    for index, label in zip(
        (1, 3, 5),
        (
            date_labels.get("year", "年"),
            date_labels.get("month", "月"),
            date_labels.get("day", "日"),
        ),
    ):
        _write_inline_cell(
            date_table.rows[0].cells[index],
            label,
            style=style,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for index in (0, 2, 4):
        _underline_cell(date_table.rows[0].cells[index])


def _flow_cell_may_exceed_page(cell) -> bool:
    """Release cantSplit only when a flowing row cannot safely fit one page."""

    style = LAYOUT["typography"]["body"]
    line_height_mm = (
        float(style["size_pt"])
        * float(LAYOUT["paragraphs"]["body_line_spacing"])
        * 25.4
        / 72.0
    )
    usable_height_mm = (
        float(LAYOUT["page"]["height_mm"])
        - float(LAYOUT["page"]["top_margin_mm"])
        - float(LAYOUT["page"]["bottom_margin_mm"])
        - 10.0
    )
    max_lines = max(1, math.floor(usable_height_mm / line_height_mm))
    usable_width_mm = float(LAYOUT["table"]["width_mm"]) - 2 * float(
        LAYOUT["table"]["horizontal_padding_mm"]
    )
    cjk_char_width_mm = float(style["size_pt"]) * 25.4 / 72.0
    line_capacity = max(1.0, usable_width_mm / cjk_char_width_mm)
    estimated_lines = 0
    for paragraph in cell.paragraphs:
        width = display_width(paragraph.text)
        estimated_lines += max(1, math.ceil(width / line_capacity))
    for extent in cell._tc.xpath(".//wp:extent"):
        height_emu = extent.get("cy")
        if height_emu:
            estimated_lines += math.ceil((int(height_emu) / 36000.0) / line_height_mm)
    return estimated_lines > max_lines


def _add_task_table(document: Document, data: dict[str, Any], fixed: dict[str, Any], *, data_dir: Path) -> None:
    styles = LAYOUT["typography"]
    body_fixed = fixed["body"]
    labels = body_fixed["section_labels"]
    signature_labels = body_fixed["signature_labels"]
    table = document.add_table(rows=6, cols=1)
    _set_table_geometry(
        table,
        (LAYOUT["table"]["width_mm"],),
        alignment=WD_TABLE_ALIGNMENT.LEFT,
    )
    _set_table_indent(table, 0.0)
    _set_table_borders(table, size_eighth_pt=round(LAYOUT["table"]["border_pt"] * 8))
    for index, row in enumerate(table.rows):
        cell = row.cells[0]
        vertical_padding = (
            LAYOUT["table"]["flow_vertical_padding_mm"]
            if index in {1, 2, 3}
            else LAYOUT["table"]["top_bottom_padding_mm"]
        )
        _set_cell_margins(
            cell,
            left=LAYOUT["table"]["horizontal_padding_mm"],
            right=LAYOUT["table"]["horizontal_padding_mm"],
            top=vertical_padding,
            bottom=vertical_padding,
        )
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _clear_cell(cell)

    cells = [row.cells[0] for row in table.rows]
    _heading(cells[0], labels["title"])
    title = cells[0].paragraphs[0]
    _append_runs(title, data["metadata"]["title"], style=styles["title_row_value"])
    reserved_second_line = cells[0].add_paragraph()
    _format_paragraph(reserved_second_line)
    _set_run_font(reserved_second_line.add_run(), style=styles["title_row_value"])

    _heading(cells[1], labels["basic_content_and_requirements"])
    _append_content_blocks(cells[1], data["sections"]["basic_content_and_requirements"], data_dir=data_dir)

    _heading(cells[2], labels["schedule"])
    _append_schedule(cells[2], data["sections"]["schedule"])

    _heading(cells[3], labels["required_materials_and_references"])
    _append_content_blocks(cells[3], data["sections"]["required_materials"], data_dir=data_dir)
    _append_references(cells[3], data["sections"]["references"])

    _heading(cells[4], labels["topic_information"])
    _topic_lines(
        cells[4],
        data["sections"]["topic_information"],
        body_fixed["topic_information"],
        signature_labels,
    )

    _heading(cells[5], signature_labels["college_opinion"])
    cells[5].paragraphs[0].paragraph_format.left_indent = Mm(
        float(LAYOUT["signature"].get("college_opinion_label_indent_mm", 0.0))
    )
    for _ in range(LAYOUT["signature"].get("college_blank_paragraphs", 4)):
        paragraph = cells[5].add_paragraph()
        _format_paragraph(paragraph)
    _add_signature_line(
        cells[5],
        signature_labels["college_leader"],
        blank_width_mm=LAYOUT["signature"]["college_leader_blank_width_mm"],
    )
    _add_date_line(cells[5], signature_labels)

    heights = LAYOUT["section_min_heights_mm"]
    for index, key in (
        (0, "title"),
        (4, "topic_information"),
        (5, "college_leader_opinion"),
    ):
        if key in heights:
            _set_row_height(table.rows[index], float(heights[key]))
    for row in table.rows:
        _set_row_cant_split(row, True)
    for index in (1, 2, 3):
        if _flow_cell_may_exceed_page(cells[index]):
            _set_row_cant_split(table.rows[index], False)


def render(template: Path, data_path: Path, output: Path, *, overwrite: bool) -> Path:
    if output.exists() and not overwrite:
        raise FileExistsError(f"output exists; pass --overwrite: {output}")
    data = validate_data(json.loads(data_path.read_text(encoding="utf-8")))
    resolve_font_files(
        LAYOUT["required_font_files"],
    )
    _require_checkbox_symbol_font()
    fixed = _fixed_content()
    document = Document(template)
    _remove_all_body_content(document)
    _set_page_system(document)
    _add_cover(document, data, fixed)
    _add_notice(document, fixed)
    _add_task_table(document, data, fixed, data_dir=data_path.resolve().parent)
    document.core_properties.title = "深圳技术大学本科毕业论文（设计）任务书"
    document.core_properties.author = "SZTU Thesis Template"
    document.core_properties.last_modified_by = "SZTU Thesis Template"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    reopened = Document(output)
    if len(reopened.tables) != 3 or len(reopened.tables[-1].rows) != 6:
        raise RuntimeError("generated DOCX failed structural reopen validation")
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=Path(__file__).with_name("official-template.docx"))
    parser.add_argument("--data", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    try:
        output = render(args.template, args.data, args.output, overwrite=args.overwrite)
    except (DataError, OSError, RuntimeError, ValueError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
