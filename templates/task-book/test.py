#!/usr/bin/env python3
"""Run structural and rendered regression checks for the task-book template."""

from __future__ import annotations

import argparse
import copy
import importlib.util
import json
import math
import re
import subprocess
import sys
from pathlib import Path
from typing import Any, Callable, Iterable

import pdfplumber
from docx import Document
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


TEMPLATES_DIR = Path(__file__).resolve().parents[1]
if str(TEMPLATES_DIR) not in sys.path:
    sys.path.insert(0, str(TEMPLATES_DIR))

from common.python.content import (  # noqa: E402
    display_width,
    runs_text,
    starts_with_calendar_date,
)
from common.python.process_form import load_process_document_layout  # noqa: E402
from common.python.font_files import resolve_font_files  # noqa: E402
from common.python.list_layout import marker_advance_pt  # noqa: E402
from common.python.regression_fixtures import (  # noqa: E402
    load_page_range_recipes,
    materialize_page_range_fixture,
)


FIXTURE_NAMES = (
    "minimal",
    "normal",
    "feedback-regression",
    "layout-stress",
    "long",
    "page-break",
    "rich-text-list",
    "nested-list",
    "list-wrap",
    "cover-subscript",
    "reference-overflow",
    "image-page-break",
    "structured-content",
)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MM_TO_PT = 72.0 / 25.4


def run(command: list[str], *, cwd: Path) -> str:
    result = subprocess.run(
        command,
        cwd=cwd,
        check=True,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
    )
    return result.stdout


def load_renderer(path: Path, module_name: str):
    spec = importlib.util.spec_from_file_location(module_name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load renderer: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def expect_invalid(action: Callable[[], Any], label: str) -> None:
    try:
        action()
    except (OSError, RuntimeError, ValueError):
        return
    raise AssertionError(f"invalid task-book input was accepted: {label}")


def visible_runs(paragraph) -> list:
    return [item for item in paragraph.runs if item.text.strip()]


def direct_cell_runs(cell) -> list:
    return [item for paragraph in cell.paragraphs for item in visible_runs(paragraph)]


def assert_run_role(
    item,
    style: dict[str, Any],
    *,
    bold: bool | None = None,
) -> None:
    r_pr = item._r.rPr
    assert r_pr is not None and r_pr.rFonts is not None
    assert r_pr.rFonts.get(qn("w:eastAsia")) == style["cjk_word_family"]
    assert r_pr.rFonts.get(qn("w:ascii")) == style["latin_word_family"]
    assert r_pr.rFonts.get(qn("w:hAnsi")) == style["latin_word_family"]
    assert item.font.size is not None
    assert abs(item.font.size.pt - float(style["size_pt"])) <= 0.01
    assert bool(item.bold) is (bool(style["bold"]) if bold is None else bold)
    assert item.font.color.rgb is not None
    assert str(item.font.color.rgb) == "000000"


def assert_paragraph_role(paragraph, style: dict[str, Any]) -> None:
    items = visible_runs(paragraph)
    assert items, f"expected visible runs in paragraph: {paragraph.text!r}"
    for item in items:
        assert_run_role(item, style)


def _dxa(element) -> int:
    return int(element.get(qn("w:w")))


def table_grid_dxa(table) -> list[int]:
    return [_dxa(column) for column in table._tbl.tblGrid]


def assert_table_geometry(table) -> None:
    tbl_w = table._tbl.tblPr.first_child_found_in("w:tblW")
    assert tbl_w is not None and tbl_w.get(qn("w:type")) == "dxa"
    grid = table_grid_dxa(table)
    # OOXML stores each fractional column and the total table width as separate
    # integer twip values, so multi-column sums may differ by a few rounding twips.
    assert grid and abs(_dxa(tbl_w) - sum(grid)) <= math.ceil(len(grid) / 2)
    layout = table._tbl.tblPr.first_child_found_in("w:tblLayout")
    assert layout is not None and layout.get(qn("w:type")) == "fixed"
    assert table.autofit is False
    for row in table.rows:
        cursor = 0
        raw_cells = row._tr.findall(qn("w:tc"))
        assert raw_cells
        for raw_cell in raw_cells:
            tc_pr = raw_cell.find(qn("w:tcPr"))
            assert tc_pr is not None
            tc_w = tc_pr.find(qn("w:tcW"))
            assert tc_w is not None and tc_w.get(qn("w:type")) == "dxa"
            grid_span = tc_pr.find(qn("w:gridSpan"))
            span = 1 if grid_span is None else int(grid_span.get(qn("w:val")))
            assert span >= 1 and cursor + span <= len(grid)
            assert _dxa(tc_w) == sum(grid[cursor : cursor + span])
            cursor += span
        assert cursor == len(grid)


def nested_tables(table) -> Iterable:
    for row in table.rows:
        for cell in row.cells:
            for nested in cell.tables:
                yield nested
                yield from nested_tables(nested)


def assert_main_table_borders(table, *, border_pt: float) -> None:
    borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
    assert borders is not None
    expected_size = str(round(border_pt * 8))
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        assert element is not None
        assert element.get(qn("w:val")) == "single"
        assert element.get(qn("w:sz")) == expected_size
        assert element.get(qn("w:space")) == "0"
        assert element.get(qn("w:color")) == "000000"
    spacing = table._tbl.tblPr.first_child_found_in("w:tblCellSpacing")
    assert spacing is None or _dxa(spacing) == 0


def cell_margin_dxa(cell, edge: str) -> int:
    margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
    assert margins is not None
    item = margins.find(qn(f"w:{edge}"))
    assert item is not None and item.get(qn("w:type")) == "dxa"
    return _dxa(item)


def twips_from_mm(value: float) -> int:
    return round(value / 25.4 * 1440)


def points_or_zero(value) -> float:
    return 0.0 if value is None else float(value.pt)


def raw_row_spans(row) -> list[int]:
    spans: list[int] = []
    for raw_cell in row._tr.findall(qn("w:tc")):
        tc_pr = raw_cell.find(qn("w:tcPr"))
        assert tc_pr is not None
        grid_span = tc_pr.find(qn("w:gridSpan"))
        spans.append(1 if grid_span is None else int(grid_span.get(qn("w:val"))))
    return spans


def assert_bottom_rule(cell, *, border_pt: float = 0.5) -> None:
    borders = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcBorders")
    assert borders is not None
    bottom = borders.find(qn("w:bottom"))
    assert bottom is not None
    assert bottom.get(qn("w:val")) == "single"
    assert bottom.get(qn("w:sz")) == str(round(border_pt * 8))
    assert bottom.get(qn("w:space")) == "0"
    assert bottom.get(qn("w:color")) == "000000"


def assert_no_bottom_rule(cell) -> None:
    borders = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcBorders")
    if borders is None:
        return
    bottom = borders.find(qn("w:bottom"))
    assert bottom is None or bottom.get(qn("w:val")) in {"nil", "none"}


def table_text(table) -> str:
    return "\n".join(cell.text for row in table.rows for cell in row.cells)


def row_has_cant_split(row) -> bool:
    tr_pr = row._tr.trPr
    return tr_pr is not None and tr_pr.find(qn("w:cantSplit")) is not None


def paragraph_has_content(paragraph) -> bool:
    return bool(paragraph.text.strip() or paragraph._p.xpath(".//w:drawing"))


def assert_signature_table(table, *, blank_mm: float, right_mm: float) -> None:
    assert len(table.rows) == 1 and len(table.columns) == 3
    assert_table_geometry(table)
    assert table.alignment == WD_TABLE_ALIGNMENT.RIGHT
    expected = [
        twips_from_mm(42.0),
        twips_from_mm(blank_mm),
        twips_from_mm(right_mm),
    ]
    assert table_grid_dxa(table) == expected
    assert table.rows[0].cells[1].text == ""
    assert table.rows[0].cells[2].text == ""
    for cell in table.rows[0].cells:
        assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    assert_bottom_rule(table.rows[0].cells[1])
    assert_no_bottom_rule(table.rows[0].cells[2])


def assert_cover_structure(
    document: Document,
    normalized: dict[str, Any],
    layout: dict[str, Any],
    fixed: dict[str, Any],
) -> None:
    title_table, info_table, _main_table = document.tables
    title_style = layout["typography"]["cover_title_value"]
    line_capacity = (float(layout["cover"]["title_underline_width_mm"])
        - 2 * float(layout["cover"]["title_cell_horizontal_padding_mm"])) / (
        float(title_style["size_pt"]) * 25.4 / 72.0
    )
    expected_title_lines = sum(
        max(1, math.ceil(display_width(segment) / line_capacity))
        for segment in runs_text(normalized["metadata"]["title"]).splitlines()
    )
    assert 1 <= expected_title_lines <= int(layout["cover"]["title_max_lines"])
    assert len(title_table.rows) == expected_title_lines
    assert len(title_table.columns) == 2
    title_label_width = (
        float(layout["cover"]["content_width_mm"])
        - float(layout["cover"]["title_underline_width_mm"])
    )
    assert table_grid_dxa(title_table) == [
        twips_from_mm(title_label_width),
        twips_from_mm(layout["cover"]["title_underline_width_mm"]),
    ]
    assert raw_row_spans(title_table.rows[0]) == [1, 1]
    assert title_table.rows[0].cells[0].text == fixed["cover"]["field_labels"]["title"]
    rendered_title_parts: list[str] = []
    for index, row in enumerate(title_table.rows):
        value_cell = row.cells[1] if index == 0 else row.cells[0]
        if index:
            assert raw_row_spans(row) == [2]
        expected_alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
            if index == 0
            else WD_ALIGN_PARAGRAPH.LEFT
        )
        assert value_cell.paragraphs[0].alignment == expected_alignment
        assert value_cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
        usable_twips = int(value_cell._tc.tcPr.tcW.w) - sum(
            cell_margin_dxa(value_cell, edge) for edge in ("left", "right")
        )
        assert display_width(value_cell.text) * float(title_style["size_pt"]) * 20 <= usable_twips
        assert_bottom_rule(value_cell)
        assert cell_margin_dxa(value_cell, "bottom") >= twips_from_mm(
            layout["cover"]["title_underline_clearance_mm"]
        )
        rendered_title_parts.append(value_cell.text)
    assert "".join(rendered_title_parts) == runs_text(normalized["metadata"]["title"])

    assert (len(info_table.rows), len(info_table.columns)) == (4, 5)
    short_width = float(layout["cover"]["short_field_underline_width_mm"])
    gap_width = float(layout["cover"]["metadata_column_gap_mm"])
    label_width = (
        float(layout["cover"]["content_width_mm"]) - 2 * short_width - gap_width
    ) / 2
    assert table_grid_dxa(info_table) == [
        twips_from_mm(label_width),
        twips_from_mm(short_width),
        twips_from_mm(gap_width),
        twips_from_mm(label_width),
        twips_from_mm(short_width),
    ]
    assert sum(table_grid_dxa(info_table)[1:]) == twips_from_mm(
        layout["cover"]["full_field_underline_width_mm"]
    )
    labels = fixed["cover"]["field_labels"]
    full_rows = (
        (labels["college"], normalized["metadata"]["college"]),
        (labels["major"], normalized["metadata"]["major"]),
    )
    for row, (label, value) in zip(info_table.rows[:2], full_rows):
        assert raw_row_spans(row) == [1, 4]
        assert row.cells[0].text == label
        assert row.cells[1].text == value
        assert row.cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
        assert row.cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert row.cells[0].vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
        assert row.cells[1].vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
        assert_bottom_rule(row.cells[1])
    paired_rows = (
        (
            labels["class_name"],
            normalized["metadata"]["class_name"],
            labels["student_id"],
            normalized["metadata"]["student_id"],
        ),
        (
            labels["student_name"],
            normalized["metadata"]["student_name"],
            labels["advisor"],
            normalized["metadata"]["advisor"],
        ),
    )
    for row, values in zip(info_table.rows[2:], paired_rows):
        assert raw_row_spans(row) == [1, 1, 1, 1, 1]
        assert [row.cells[index].text for index in (0, 1, 3, 4)] == list(values)
        assert row.cells[2].text == ""
        for index in (0, 3):
            assert row.cells[index].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
        for index in (1, 4):
            assert row.cells[index].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
            assert_bottom_rule(row.cells[index])
        for cell in row.cells:
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER


def assert_main_table_position(table, layout: dict[str, Any]) -> None:
    assert table.alignment == WD_TABLE_ALIGNMENT.LEFT
    table_properties = table._tbl.tblPr
    justification = table_properties.first_child_found_in("w:jc")
    assert justification is not None and justification.get(qn("w:val")) == "left"
    indent = table_properties.first_child_found_in("w:tblInd")
    assert indent is not None
    assert indent.get(qn("w:type")) == "dxa"
    assert _dxa(indent) == 0
    expected_width = twips_from_mm(layout["table"]["width_mm"])
    assert table_grid_dxa(table) == [expected_width]
    table_width = table_properties.first_child_found_in("w:tblW")
    assert table_width is not None and _dxa(table_width) == expected_width


def assert_date_table(
    table,
    layout: dict[str, Any],
    fixed: dict[str, Any],
) -> None:
    assert len(table.rows) == 1 and len(table.columns) == 7
    assert_table_geometry(table)
    assert table.alignment == WD_TABLE_ALIGNMENT.RIGHT
    label_width = float(layout["signature"].get("date_label_width_mm", 5.0))
    expected = [
        float(layout["signature"]["date_year_blank_width_mm"]),
        label_width,
        float(layout["signature"]["date_month_blank_width_mm"]),
        label_width,
        float(layout["signature"]["date_day_blank_width_mm"]),
        label_width,
        float(layout["signature"]["right_inset_mm"]),
    ]
    assert table_grid_dxa(table) == [twips_from_mm(value) for value in expected]
    labels = fixed["body"]["signature_labels"]["date_labels"]
    assert [table.rows[0].cells[index].text for index in (1, 3, 5)] == [
        labels["year"],
        labels["month"],
        labels["day"],
    ]
    for index in (0, 2, 4):
        assert table.rows[0].cells[index].text == ""
        assert_bottom_rule(table.rows[0].cells[index])
    assert table.rows[0].cells[6].text == ""
    assert_no_bottom_rule(table.rows[0].cells[6])
    for cell in table.rows[0].cells:
        assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.BOTTOM


def assert_word_typography(
    document: Document,
    normalized: dict[str, Any],
    layout: dict[str, Any],
    fixed: dict[str, Any],
) -> None:
    styles = layout["typography"]
    assert_paragraph_role(document.paragraphs[0], styles["cover_number"])
    assert_paragraph_role(document.paragraphs[1], styles["school_name"])
    assert_paragraph_role(document.paragraphs[2], styles["document_title"])
    assert_paragraph_role(document.paragraphs[3], styles["cohort"])
    assert document.paragraphs[3].text == (
        f"（ {normalized['metadata']['graduation_year']} 届）"
    )
    for run in visible_runs(document.paragraphs[3]):
        assert run.bold is True

    title_table, info_table, main_table = document.tables
    assert_paragraph_role(
        title_table.rows[0].cells[0].paragraphs[0], styles["cover_label"]
    )
    for index, row in enumerate(title_table.rows):
        value_cell = row.cells[1] if index == 0 else row.cells[0]
        assert_paragraph_role(value_cell.paragraphs[0], styles["cover_title_value"])
    for row in info_table.rows[:2]:
        assert_paragraph_role(row.cells[0].paragraphs[0], styles["cover_label"])
        assert_paragraph_role(row.cells[1].paragraphs[0], styles["cover_value"])
    for row in info_table.rows[2:]:
        assert_paragraph_role(row.cells[0].paragraphs[0], styles["cover_label"])
        assert_paragraph_role(row.cells[1].paragraphs[0], styles["cover_value"])
        assert_paragraph_role(row.cells[3].paragraphs[0], styles["cover_label"])
        assert_paragraph_role(row.cells[4].paragraphs[0], styles["cover_value"])

    notice_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text == "本科生毕业论文（设计）须知"
    )
    assert_paragraph_role(document.paragraphs[notice_index], styles["notice_title"])
    for paragraph in document.paragraphs[notice_index + 1 : notice_index + 7]:
        assert_paragraph_role(paragraph, styles["notice_body"])
        spacing = paragraph._p.get_or_add_pPr().find(qn("w:spacing"))
        assert spacing is not None
        assert spacing.get(qn("w:line")) == "360"
        assert spacing.get(qn("w:lineRule")) == "auto"
        assert spacing.get(qn("w:after")) in {None, "0"}

    title_paragraph = main_table.rows[0].cells[0].paragraphs[0]
    title_runs = visible_runs(title_paragraph)
    assert title_runs
    assert title_paragraph.text == (
        fixed["body"]["section_labels"]["title"]
        + runs_text(normalized["metadata"]["title"])
    )
    assert_run_role(title_runs[0], styles["section_title"])
    for item in title_runs[1:]:
        assert_run_role(item, styles["title_row_value"])
    reserved = main_table.rows[0].cells[0].paragraphs[1]
    assert reserved.text == "" and reserved.runs
    for item in reserved.runs:
        assert_run_role(item, styles["title_row_value"])
    for row in main_table.rows[1:]:
        assert_paragraph_role(row.cells[0].paragraphs[0], styles["section_title"])
    for row_index in (1, 2, 3):
        for paragraph in main_table.rows[row_index].cells[0].paragraphs[1:]:
            if paragraph.text.strip():
                assert_paragraph_role(paragraph, styles["body"])
    for paragraph in main_table.rows[4].cells[0].paragraphs[1:]:
        if paragraph.text.strip():
            assert_paragraph_role(paragraph, styles["topic_body"])

    topic_tables = main_table.rows[4].cells[0].tables
    for nested in topic_tables[:-1]:
        for row in nested.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        assert_paragraph_role(paragraph, styles["topic_body"])
    advisor_signature = topic_tables[-1]
    assert_paragraph_role(
        advisor_signature.rows[0].cells[0].paragraphs[0], styles["topic_body"]
    )
    assert_paragraph_role(
        advisor_signature.rows[0].cells[1].paragraphs[0], styles["signature"]
    )
    college_signature = main_table.rows[5].cells[0].tables[0]
    assert_paragraph_role(
        college_signature.rows[0].cells[0].paragraphs[0], styles["signature"]
    )
    date_table = main_table.rows[5].cells[0].tables[1]
    for index in (1, 3, 5):
        assert_paragraph_role(
            date_table.rows[0].cells[index].paragraphs[0], styles["notice_body"]
        )


def assert_word_indents(
    document: Document,
    normalized: dict[str, Any],
    layout: dict[str, Any],
) -> set[bool]:
    main = document.tables[-1]
    body_size = float(layout["typography"]["body"]["size_pt"])
    first_line = body_size * float(layout["paragraphs"]["first_line_indent_em"])
    list_first = body_size * float(
        layout["paragraphs"]["list_first_level_indent_em"]
    )
    assert abs(list_first - body_size * 2.0) <= 0.01
    list_level = body_size * float(layout["paragraphs"]["list_level_indent_em"])
    list_hanging = body_size * float(
        layout["paragraphs"]["list_hanging_indent_em"]
    )
    unordered_markers = layout["paragraphs"]["unordered_list_markers"]
    fonts = resolve_font_files(["Times New Roman", "SimSun"])

    def assert_list(block: dict[str, Any], cell, *, depth: int = 1) -> None:
        for index, item in enumerate(block["items"], start=1):
            if block["type"] == "ordered_list":
                explicit_marker = item.get("marker")
                marker = explicit_marker or f"{index}."
                text = f"{marker}\u2009{runs_text(item['runs'])}"
            else:
                marker = unordered_markers[depth - 1]
                text = f"{marker}\u2009{runs_text(item['runs'])}"
            paragraph = next(item for item in cell.paragraphs if item.text == text)
            marker_indent = list_first + (depth - 1) * list_level
            actual_hanging = max(
                list_hanging,
                marker_advance_pt(
                    marker,
                    size_pt=body_size,
                    latin_font=fonts["Times New Roman"],
                    cjk_font=fonts["SimSun"],
                ),
            )
            assert abs(
                points_or_zero(paragraph.paragraph_format.left_indent)
                - marker_indent
                - actual_hanging
            ) <= 0.05
            assert abs(
                points_or_zero(paragraph.paragraph_format.first_line_indent)
                + actual_hanging
            ) <= 0.05
            if item["children"]:
                assert_list(item["children"], cell, depth=depth + 1)

    for title in main.rows[0].cells[0].paragraphs[:2]:
        assert abs(points_or_zero(title.paragraph_format.first_line_indent)) <= 0.01
        assert abs(points_or_zero(title.paragraph_format.left_indent)) <= 0.01

    basic_cell = main.rows[1].cells[0]
    for block in normalized["sections"]["basic_content_and_requirements"]:
        if block["type"] != "paragraph":
            continue
        text = runs_text(block["runs"])
        paragraph = next(item for item in basic_cell.paragraphs if item.text == text)
        assert abs(points_or_zero(paragraph.paragraph_format.first_line_indent) - first_line) <= 0.01
    for block in normalized["sections"]["basic_content_and_requirements"]:
        if block["type"] in {"ordered_list", "unordered_list"}:
            assert_list(block, basic_cell)

    schedule_paragraphs = [
        item for item in main.rows[2].cells[0].paragraphs[1:] if item.text.strip()
    ]
    assert len(schedule_paragraphs) == len(normalized["sections"]["schedule"])
    date_prefix_cases: set[bool] = set()
    for paragraph, item in zip(
        schedule_paragraphs, normalized["sections"]["schedule"]
    ):
        expected_text = f"{item['period']}：{runs_text(item['content'])}"
        assert paragraph.text == expected_text
        is_date_prefix = starts_with_calendar_date(item["period"])
        date_prefix_cases.add(is_date_prefix)
        expected_indent = 0.0 if is_date_prefix else first_line
        assert abs(points_or_zero(paragraph.paragraph_format.left_indent)) <= 0.01
        assert abs(points_or_zero(paragraph.paragraph_format.first_line_indent) - expected_indent) <= 0.01

    materials_cell = main.rows[3].cells[0]
    for block in normalized["sections"]["required_materials"]:
        if block["type"] != "paragraph":
            continue
        text = runs_text(block["runs"])
        paragraph = next(item for item in materials_cell.paragraphs if item.text == text)
        assert abs(points_or_zero(paragraph.paragraph_format.first_line_indent) - first_line) <= 0.01
    for block in normalized["sections"]["required_materials"]:
        if block["type"] in {"ordered_list", "unordered_list"}:
            assert_list(block, materials_cell)

    reference_texts = [runs_text(item) for item in normalized["sections"]["references"]]
    rendered_references = [
        paragraph
        for paragraph in materials_cell.paragraphs
        if paragraph.text in reference_texts
    ]
    assert [paragraph.text for paragraph in rendered_references] == reference_texts
    for paragraph in rendered_references:
        assert abs(points_or_zero(paragraph.paragraph_format.left_indent)) <= 0.01
        assert abs(points_or_zero(paragraph.paragraph_format.first_line_indent)) <= 0.01
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    return date_prefix_cases


def assert_word_topic_choices(
    document: Document,
    topic: dict[str, Any],
    layout: dict[str, Any],
    fixed: dict[str, Any],
) -> None:
    topic_cell = document.tables[-1].rows[4].cells[0]
    paragraphs = topic_cell.paragraphs
    lines = [paragraph.text for paragraph in paragraphs if paragraph.text.strip()]
    topic_fixed = fixed["body"]["topic_information"]
    signatures = fixed["body"]["signature_labels"]
    def symbols_in(element) -> list[str]:
        return [
            str(item.get(qn("w:char"))).upper()
            for item in element.xpath(".//w:sym")
        ]

    checked = "0052"
    unchecked = "00A3"
    assert len(topic_cell.tables) == 4
    nature_table, research_table, project_table, final_table = topic_cell.tables
    assert nature_table.rows[0].cells[0].text == topic_fixed["nature_label"]
    nature_paragraph = nature_table.rows[0].cells[1].paragraphs[0]
    nature_options = topic_fixed["nature_options"]
    assert nature_options["graduation_design"] in nature_paragraph.text
    assert nature_options["graduation_thesis"] in nature_paragraph.text
    assert symbols_in(nature_paragraph._p) == [
        checked if topic["nature"] == "graduation_design" else unchecked,
        checked if topic["nature"] == "graduation_thesis" else unchecked,
    ]

    source = topic["source"]
    source_type = source["type"]
    assert research_table.rows[0].cells[0].text == topic_fixed["source_label"]
    research_cell = research_table.rows[0].cells[1]
    research_line = research_cell.text
    other_cell = research_table.rows[0].cells[2]
    project_value_cell = project_table.rows[0].cells[2]
    source_options = topic_fixed["source_options"]
    levels = topic_fixed["research_project_levels"]
    proposer_options = topic_fixed["self_proposed_by_options"]
    assert research_line.startswith(source_options["research_project"])
    assert levels["national"] in research_line
    assert levels["provincial_ministerial"] in research_line
    assert symbols_in(research_cell._tc) == [
        checked
        if source_type == "research_project"
        and source["research_project"]["level"] == "national"
        else unchecked,
        checked
        if source_type == "research_project"
        and source["research_project"]["level"] == "provincial_ministerial"
        else unchecked,
    ]
    other_selected = (
        source_type == "research_project"
        and source["research_project"]["level"] == "other"
    )
    other_text = (
        source["research_project"].get("other_level", "")
        if source_type == "research_project"
        else ""
    )
    assert other_cell.text == (other_text if other_selected else "")
    assert_bottom_rule(other_cell)
    assert_no_bottom_rule(research_table.rows[0].cells[3])

    assert project_table.rows[0].cells[1].text == topic_fixed["project_number_label"]
    expected_project_number = (
        source["research_project"]["project_number"]
        if source_type == "research_project"
        else ""
    )
    assert project_value_cell.text == expected_project_number
    assert_bottom_rule(project_value_cell)
    assert project_value_cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert points_or_zero(project_value_cell.paragraphs[0].paragraph_format.space_after) == 0.0
    expected_position = -round(
        float(layout["signature"]["topic_project_value_baseline_shift_mm"])
        * 72.0
        / 25.4
        * 2.0
    )
    if expected_project_number:
        assert project_value_cell.paragraphs[0].runs
        assert all(
            run._r.get_or_add_rPr().find(qn("w:position")) is not None
            and int(run._r.get_or_add_rPr().find(qn("w:position")).get(qn("w:val")))
            == expected_position
            for run in project_value_cell.paragraphs[0].runs
        )
    assert_no_bottom_rule(project_table.rows[0].cells[3])

    practice_paragraph = next(
        item
        for item in paragraphs
        if source_options["practice_project"] in item.text
    )
    self_heading_paragraph = next(
        item
        for item in paragraphs
        if source_options["self_proposed"] in item.text
    )
    assert symbols_in(practice_paragraph._p) == [
        checked if source_type == "practice_project" else unchecked
    ]
    assert symbols_in(self_heading_paragraph._p) == []
    proposer = source.get("self_proposed_by")
    teacher_paragraph = next(
        item for item in paragraphs if proposer_options["teacher"] in item.text
    )
    student_paragraph = next(
        item for item in paragraphs if proposer_options["student"] in item.text
    )
    joint_paragraph = final_table.rows[0].cells[0].paragraphs[0]
    for name, paragraph in {
        "teacher": teacher_paragraph,
        "student": student_paragraph,
        "teacher_student_joint": joint_paragraph,
    }.items():
        assert symbols_in(paragraph._p) == [
            checked
            if source_type == "self_proposed" and proposer == name
            else unchecked
        ]
    selected_project_level = (
        source_type == "research_project"
        and source["research_project"]["level"] in {"national", "provincial_ministerial"}
    )
    expected_total = 1 + int(
        selected_project_level or source_type in {"practice_project", "self_proposed"}
    )
    all_symbols = topic_cell._tc.xpath(".//w:sym")
    symbol_states = [str(item.get(qn("w:char"))).upper() for item in all_symbols]
    assert symbol_states.count(checked) == expected_total
    assert len(symbol_states) == 8
    assert set(symbol_states) <= {checked, unchecked}
    topic_text = "\n".join(lines + [table_text(table) for table in topic_cell.tables])
    for forbidden in (
        "■",
        "□",
        "☑",
        "☒",
        "✅",
        "☐",
        "outline_box",
        "outline_box_with_tick",
        "\ufe0f",
    ):
        assert forbidden not in topic_text
    for symbol in all_symbols:
        assert symbol.get(qn("w:font")) == "Wingdings 2"
        run = symbol.getparent()
        r_pr = run.find(qn("w:rPr"))
        assert r_pr is not None
        size = r_pr.find(qn("w:sz"))
        color = r_pr.find(qn("w:color"))
        assert size is not None and size.get(qn("w:val")) == "21"
        assert color is not None and color.get(qn("w:val")) == "000000"
        assert r_pr.find(qn("w:highlight")) is None

    usable_width = float(layout["table"]["width_mm"]) - 2 * float(
        layout["table"]["horizontal_padding_mm"]
    )
    topic_style = layout["typography"]["topic_body"]
    right_inset = float(layout["signature"]["right_inset_mm"])
    other_width = float(layout["signature"].get("topic_other_blank_width_mm", 32.0))
    option_indent_em = float(
        layout["signature"].get("topic_option_left_indent_em", 6.0)
    )
    option_indent_mm = (
        float(topic_style["size_pt"]) * option_indent_em * 25.4 / 72.0
    )
    assert table_grid_dxa(nature_table) == [
        twips_from_mm(option_indent_mm),
        twips_from_mm(usable_width - option_indent_mm),
    ]
    assert table_grid_dxa(research_table) == [
        twips_from_mm(option_indent_mm),
        twips_from_mm(usable_width - option_indent_mm - other_width - right_inset),
        twips_from_mm(other_width),
        twips_from_mm(right_inset),
    ]
    project_indent_mm = (
        float(topic_style["size_pt"])
        * float(layout["signature"].get("topic_project_left_indent_em", 5.0))
        * 25.4
        / 72.0
    )
    project_label_width = float(
        layout["signature"].get("topic_project_label_width_mm", 25.0)
    )
    project_blank_width = float(
        layout["signature"].get("topic_project_blank_width_mm", 65.0)
    )
    assert table_grid_dxa(project_table) == [
        twips_from_mm(project_indent_mm),
        twips_from_mm(project_label_width),
        twips_from_mm(project_blank_width),
        twips_from_mm(
            usable_width - project_indent_mm - project_label_width - project_blank_width
        ),
    ]
    assert project_table.rows[0].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT

    signature_label_width = 42.0
    advisor_blank = float(layout["signature"]["advisor_blank_width_mm"])
    assert table_grid_dxa(final_table) == [
        twips_from_mm(
            usable_width - signature_label_width - advisor_blank - right_inset
        ),
        twips_from_mm(signature_label_width),
        twips_from_mm(advisor_blank),
        twips_from_mm(right_inset),
    ]
    assert final_table.rows[0].cells[1].text == signatures["advisor"]
    option_indent = float(topic_style["size_pt"]) * float(
        layout["signature"]["topic_option_left_indent_em"]
    )
    child_indent = float(topic_style["size_pt"]) * float(
        layout["signature"]["self_proposed_child_left_indent_em"]
    )
    assert abs(
        points_or_zero(practice_paragraph.paragraph_format.left_indent)
        - option_indent
    ) <= 0.01
    assert abs(
        points_or_zero(self_heading_paragraph.paragraph_format.left_indent)
        - option_indent
    ) <= 0.01
    for paragraph in (
        teacher_paragraph,
        student_paragraph,
        joint_paragraph,
    ):
        assert abs(
            points_or_zero(paragraph.paragraph_format.left_indent) - child_indent
        ) <= 0.03
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert final_table.rows[0].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert final_table.rows[0].cells[2].text == ""
    assert_bottom_rule(final_table.rows[0].cells[2])
    assert_no_bottom_rule(final_table.rows[0].cells[3])
    for table in topic_cell.tables:
        for cell in table.rows[0].cells:
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.BOTTOM


def normalized_visible_text(value: str) -> str:
    return re.sub(r"\s+", "", value)


def normalized_hyphenation_text(value: str) -> str:
    """Normalize TeX discretionary line-break hyphens for text-presence checks."""

    return re.sub(r"[\s-]+", "", value)


def content_texts(normalized: dict[str, Any]) -> list[str]:
    values = [
        runs_text(normalized["metadata"]["title"]),
        normalized["metadata"]["student_name"],
        normalized["metadata"]["student_id"],
        normalized["metadata"]["college"],
        normalized["metadata"]["major"],
        normalized["metadata"]["class_name"],
        normalized["metadata"]["advisor"],
    ]

    def add_blocks(blocks: list[dict[str, Any]]) -> None:
        def add_list(block: dict[str, Any]) -> None:
            for item in block["items"]:
                marker = item.get("marker") if block["type"] == "ordered_list" else ""
                values.append((marker or "") + runs_text(item["runs"]))
                if item["children"]:
                    add_list(item["children"])

        for block in blocks:
            if block["type"] == "paragraph":
                values.append(runs_text(block["runs"]))
            elif block["type"] in {"ordered_list", "unordered_list"}:
                add_list(block)
            elif block["type"] == "image" and block.get("caption_runs"):
                values.append(runs_text(block["caption_runs"]))
            elif block["type"] == "figure_group":
                values.append(runs_text(block["caption_runs"]))
                for item in block["items"]:
                    if item["caption_runs"]:
                        values.append(runs_text(item["caption_runs"]))
            elif block["type"] == "data_table":
                values.extend(runs_text(column["header_runs"]) for column in block["columns"])
                values.extend(runs_text(cell) for row in block["rows"] for cell in row)
                if block["caption_runs"]:
                    values.append(runs_text(block["caption_runs"]))

    add_blocks(normalized["sections"]["basic_content_and_requirements"])
    for item in normalized["sections"]["schedule"]:
        values.extend((item["period"], runs_text(item["content"])))
    add_blocks(normalized["sections"]["required_materials"])
    values.extend(runs_text(item) for item in normalized["sections"]["references"])
    return [value for value in values if value.strip()]


def _long_horizontal_edges(page) -> list[dict[str, Any]]:
    return [
        edge
        for edge in page.edges
        if abs(float(edge["bottom"]) - float(edge["top"])) <= 0.2
        and float(edge["x1"]) - float(edge["x0"]) >= 420.0
    ]


def _long_vertical_edges(page) -> list[dict[str, Any]]:
    return [
        edge
        for edge in page.edges
        if abs(float(edge["x1"]) - float(edge["x0"])) <= 0.2
        and float(edge["bottom"]) - float(edge["top"]) >= 15.0
    ]


def assert_continuous_side(
    edges: list[dict[str, Any]],
    *,
    x: float,
    top: float,
    bottom: float,
) -> None:
    intervals = sorted(
        (
            float(edge["top"]),
            float(edge["bottom"]),
        )
        for edge in edges
        if abs(float(edge["x0"]) - x) <= 0.7
    )
    assert intervals
    cursor = top
    for start, end in intervals:
        if end < cursor:
            continue
        assert start <= cursor + 1.2, f"vertical border gap at x={x}: {cursor}..{start}"
        cursor = max(cursor, end)
    assert cursor >= bottom - 1.2


def assert_pdf_form_geometry(pdf, *, fixture_name: str) -> None:
    title_page = next(
        index
        for index, page in enumerate(pdf.pages)
        if "题目名称" in (page.extract_text() or "")
    )
    assert title_page == 2
    body_pages = pdf.pages[title_page:]
    assert len(body_pages) >= 1
    for relative_index, page in enumerate(body_pages):
        horizontal = _long_horizontal_edges(page)
        vertical = _long_vertical_edges(page)
        assert horizontal, f"{fixture_name} body page {relative_index + 3} lacks closing rules"
        assert vertical, f"{fixture_name} body page {relative_index + 3} lacks side borders"
        left = min(float(edge["x0"]) for edge in vertical)
        right = max(float(edge["x0"]) for edge in vertical)
        assert abs(left - 90.0) <= 1.0
        assert abs(right - 516.2) <= 1.0
        top = min(float(edge["top"]) for edge in horizontal)
        bottom = max(float(edge["top"]) for edge in horizontal)
        assert top <= (83.0 if relative_index == 0 else 82.5)
        assert bottom > top + 100.0
        assert_continuous_side(vertical, x=left, top=top, bottom=bottom)
        assert_continuous_side(vertical, x=right, top=top, bottom=bottom)
    if fixture_name == "page-break":
        first_body = body_pages[0]
        bottom = max(float(edge["top"]) for edge in _long_horizontal_edges(first_body))
        assert bottom >= float(first_body.height) - 90.0
        next_top = min(float(edge["top"]) for edge in _long_horizontal_edges(body_pages[1]))
        assert next_top <= 82.5


def assert_pdf_closing_bundle(
    pdf,
    layout: dict[str, Any],
    normalized: dict[str, Any],
) -> None:
    found: dict[str, tuple[int, dict[str, Any]]] = {}
    for page_index, page in enumerate(pdf.pages, start=1):
        for word in page.extract_words():
            text = word["text"]
            if text in {"四、选题信息：", "指导教师签名：", "学院领导意见：", "签名："}:
                found[text] = (page_index, word)
    assert set(found) == {"四、选题信息：", "指导教师签名：", "学院领导意见：", "签名："}
    topic_page = found["四、选题信息："][0]
    teacher_page = found["指导教师签名："][0]
    college_page = found["学院领导意见："][0]
    college_signature_page = found["签名："][0]
    assert topic_page == teacher_page
    assert college_page == college_signature_page == len(pdf.pages)
    assert topic_page <= college_page

    topic_words = pdf.pages[topic_page - 1].extract_words()
    topic_top = float(found["四、选题信息："][1]["top"])

    def topic_word(text: str) -> dict[str, Any]:
        return next(
            word
            for word in topic_words
            if word["text"] == text and float(word["top"]) > topic_top
        )

    primary_lefts = [
        float(topic_word(text)["x0"])
        for text in ("设计", "1.", "2.", "3.")
    ]
    assert max(primary_lefts) - min(primary_lefts) <= 0.8
    self_proposed_left = float(topic_word("自拟题目")["x0"])
    child_lefts = [
        float(topic_word(text)["x0"])
        for text in ("教师自拟", "学生自拟", "师生共拟")
    ]
    assert max(abs(left - self_proposed_left) for left in child_lefts) <= 0.8

    topic_render_page = pdf.pages[topic_page - 1]
    research_project = topic_word("科研项目")
    project_number = topic_word("项目编号：")
    assert abs(float(project_number["x0"]) - float(research_project["x0"])) <= 0.8
    expected_project_rule_width = (
        float(layout["signature"]["topic_project_blank_width_mm"]) * MM_TO_PT
    )
    project_rules = [
        edge
        for edge in topic_render_page.lines
        if abs(float(edge["bottom"]) - float(edge["top"])) <= 0.2
        and float(project_number["bottom"]) <= float(edge["top"]) <= float(project_number["bottom"]) + 5.0
    ]
    project_rule = min(
        project_rules,
        key=lambda edge: abs(float(edge["width"]) - expected_project_rule_width),
    )
    assert abs(float(project_rule["width"]) - expected_project_rule_width) <= 0.6
    source = normalized["sections"]["topic_information"]["source"]
    if source["type"] == "research_project":
        project_value = topic_word(source["research_project"]["project_number"])
        project_value_gap = float(project_rule["top"]) - float(project_value["bottom"])
        assert 0.2 <= project_value_gap <= 2.5

    practice_project = topic_word("实践项目")
    practice_center = (
        float(practice_project["top"]) + float(practice_project["bottom"])
    ) / 2
    checkbox_rects = [
        rect
        for rect in topic_render_page.rects
        if 5.0 <= float(rect["width"]) <= 10.0
        and 5.0 <= float(rect["height"]) <= 10.0
        and float(rect["x0"]) >= float(practice_project["x1"])
        and abs(
            (float(rect["top"]) + float(rect["bottom"])) / 2 - practice_center
        ) <= 4.0
    ]
    practice_checkbox = min(checkbox_rects, key=lambda rect: float(rect["x0"]))
    checkbox_center = (
        float(practice_checkbox["top"]) + float(practice_checkbox["bottom"])
    ) / 2
    assert abs(checkbox_center - practice_center) <= 0.6
    checkbox_gap = float(practice_checkbox["x0"]) - float(practice_project["x1"])
    assert 1.5 <= checkbox_gap <= 3.2

    topic_form_right = max(
        float(edge["x0"])
        for edge in _long_vertical_edges(pdf.pages[topic_page - 1])
    )
    college_form_right = max(
        float(edge["x0"])
        for edge in _long_vertical_edges(pdf.pages[college_page - 1])
    )
    teacher_gap = topic_form_right - float(found["指导教师签名："][1]["x1"])
    college_gap = college_form_right - float(found["签名："][1]["x1"])
    right_inset = float(layout["signature"]["right_inset_mm"])
    teacher_required = (
        float(layout["signature"]["advisor_blank_width_mm"]) + right_inset
    ) * MM_TO_PT
    college_required = (
        float(layout["signature"]["college_leader_blank_width_mm"]) + right_inset
    ) * MM_TO_PT
    assert teacher_gap >= teacher_required - 2.0
    assert college_gap >= college_required - 2.0


def assert_pdf_notice_geometry(pdf) -> None:
    page = pdf.pages[1]
    words = page.extract_words()
    title = next(
        word for word in words if word["text"] == "本科生毕业论文（设计）须知"
    )
    assert abs(float(title["top"]) - 119.99) <= 1.0
    expected_item_tops = (201.14, 222.74, 244.34, 265.94, 309.14, 352.34)
    observed_line_tops: list[float] = []
    for index, expected_top in enumerate(expected_item_tops, start=1):
        item = next(
            word for word in words if word["text"].startswith(f"{index}．")
        )
        assert abs(float(item["top"]) - expected_top) <= 1.0
        assert abs(float(item["x0"]) - 90.1) <= 1.0
        observed_line_tops.append(float(item["top"]))
    for prefix, expected_top in (("与要求", 287.54), ("擅自带离", 330.74)):
        continuation = next(
            word for word in words if word["text"].startswith(prefix)
        )
        assert abs(float(continuation["top"]) - expected_top) <= 1.0
        assert abs(float(continuation["x0"]) - 107.75) <= 1.0
        observed_line_tops.append(float(continuation["top"]))

    # The official OOXML uses w:line=360 with lineRule=auto.  At small-four
    # (12 pt), each rendered line therefore advances by a true 1.5 line height
    # of 21.6 pt, including natural continuation lines inside long items.
    observed_line_tops.sort()
    for first, second in zip(observed_line_tops, observed_line_tops[1:]):
        assert abs((second - first) - 21.6) <= 0.5


def assert_pdf_cover_geometry(
    pdf,
    *,
    normalized: dict[str, Any],
    layout: dict[str, Any],
    title_line_count: int,
) -> None:
    page = pdf.pages[0]
    cover = layout["cover"]
    title_top = float(cover["title_field_top_mm"]) * MM_TO_PT
    horizontal = [
        edge
        for edge in page.lines
        if abs(float(edge["bottom"]) - float(edge["top"])) <= 0.2
        and title_top + 10.0 <= float(edge["top"]) <= title_top + 55.0
    ]
    short_width = float(cover["title_underline_width_mm"]) * MM_TO_PT
    short_rule = min(horizontal, key=lambda edge: abs(float(edge["width"]) - short_width))
    assert abs(float(short_rule["width"]) - short_width) <= 0.6

    words = page.extract_words(extra_attrs=["fontname", "size"])
    title_label = next(word for word in words if word["text"] == "题目：")
    label_to_rule_gap = float(short_rule["x0"]) - float(title_label["x1"])
    assert label_to_rule_gap >= -0.2
    assert label_to_rule_gap <= (
        float(cover["title_label_to_rule_gap_max_mm"]) * MM_TO_PT + 0.2
    )
    title_words = [
        word
        for word in words
        if abs(float(word["size"]) - 16.0) <= 0.1
        and title_top - 4.0 <= float(word["top"]) <= title_top + 42.0
        and word["text"] != "题目："
    ]
    grouped: list[list[dict[str, Any]]] = []
    for word in sorted(title_words, key=lambda item: float(item["top"])):
        for group in grouped:
            if abs(float(group[0]["top"]) - float(word["top"])) <= 2.0:
                group.append(word)
                break
        else:
            grouped.append([word])
    assert len(grouped) == title_line_count
    spans = [
        (
            min(float(word["x0"]) for word in group),
            max(float(word["x1"]) for word in group),
        )
        for group in grouped
    ]
    first_center = (spans[0][0] + spans[0][1]) / 2
    rule_center = (float(short_rule["x0"]) + float(short_rule["x1"])) / 2
    assert abs(first_center - rule_center) <= 1.0
    if title_line_count == 2:
        full_width = float(cover["content_width_mm"]) * MM_TO_PT
        full_rule = min(
            horizontal,
            key=lambda edge: abs(float(edge["width"]) - full_width),
        )
        assert abs(float(full_rule["width"]) - full_width) <= 0.6
        assert abs(spans[1][0] - float(full_rule["x0"])) <= 0.5

    rules = [short_rule]
    if title_line_count == 2:
        rules.append(full_rule)
    minimum_clearance = (
        float(cover["title_underline_clearance_mm"]) * MM_TO_PT
    )
    previous_rule_top = title_top - 4.0
    for rule in rules:
        rule_top = float(rule["top"])
        line_chars = [
            char
            for char in page.chars
            if char.get("text", "").strip()
            and float(char["x0"]) >= float(rule["x0"]) - 0.5
            and float(char["x1"]) <= float(rule["x1"]) + 0.5
            and float(char["top"]) >= previous_rule_top
            and float(char["top"]) < rule_top
        ]
        assert line_chars
        lowest_glyph = max(float(char["bottom"]) for char in line_chars)
        visible_rule_top = rule_top - float(rule["linewidth"]) / 2
        assert visible_rule_top - lowest_glyph >= minimum_clearance - 0.2
        previous_rule_top = rule_top + 0.2

    cohort_top = float(cover["cohort_top_mm"]) * MM_TO_PT
    cohort_chars = [
        char
        for char in page.chars
        if char.get("text", "").strip()
        and abs(float(char["top"]) - cohort_top) <= 2.0
    ]
    cohort_text = "".join(
        char["text"] for char in sorted(cohort_chars, key=lambda item: float(item["x0"]))
    )
    expected_cohort = f"（{normalized['metadata']['graduation_year']}届）"
    assert cohort_text == expected_cohort
    ordered_cohort_chars = sorted(cohort_chars, key=lambda item: float(item["x0"]))
    year_text = str(normalized["metadata"]["graduation_year"])
    opening_parenthesis = ordered_cohort_chars[0]
    first_year_digit = ordered_cohort_chars[1]
    last_year_digit = ordered_cohort_chars[len(year_text)]
    cohort_label = ordered_cohort_chars[len(year_text) + 1]
    expected_inner_gap = float(cover["cohort_inner_gap_em"]) * 14.0
    opening_gap = float(first_year_digit["x0"]) - float(opening_parenthesis["x1"])
    label_gap = float(cohort_label["x0"]) - float(last_year_digit["x1"])
    assert abs(opening_gap - expected_inner_gap) <= 0.8
    assert abs(label_gap - expected_inner_gap) <= 0.8
    for char in cohort_chars:
        assert "SimHei" in char["fontname"]
        assert abs(float(char["size"]) - 14.0) <= 0.05


def assert_pdf_title_row_geometry(pdf, layout: dict[str, Any]) -> None:
    page = next(
        item for item in pdf.pages if "题目名称" in (item.extract_text() or "")
    )
    positions = sorted(
        float(edge["top"])
        for edge in _long_horizontal_edges(page)
        if float(edge["top"]) < 160.0
    )
    groups: list[list[float]] = []
    for position in positions:
        if not groups or position - groups[-1][-1] > 1.0:
            groups.append([position])
        else:
            groups[-1].append(position)
    assert len(groups) >= 2
    row_top = min(groups[0])
    row_bottom = max(groups[1])
    expected_height = float(layout["section_min_heights_mm"]["title"]) * MM_TO_PT
    assert abs((row_bottom - row_top) - expected_height) <= 0.6

    row_words = [
        word
        for word in page.extract_words()
        if row_top <= float(word["top"]) and float(word["bottom"]) <= row_bottom
    ]
    assert row_words
    assert min(float(word["top"]) for word in row_words) - row_top >= 2.0
    assert row_bottom - max(float(word["bottom"]) for word in row_words) >= 2.0


def assert_pdf_fonts(font_report: str) -> None:
    lines = [line for line in font_report.splitlines()[2:] if line.strip()]
    assert lines
    for line in lines:
        columns = line.split()
        assert columns[-5:-2] == ["yes", "yes", "yes"], line
    names = "\n".join(lines)
    for family in ("STXingkai", "SimHei", "SimSun", "TimesNewRoman"):
        assert family in names


def assert_negative_model_cases(
    word_renderer,
    latex_renderer,
    minimal: dict[str, Any],
    task_dir: Path,
) -> None:
    cases: list[tuple[str, Callable[[dict[str, Any]], None]]] = [
        ("schema version", lambda data: data.__setitem__("schema_version", "1.0")),
        ("boolean graduation year", lambda data: data["metadata"].__setitem__("graduation_year", True)),
        ("graduation year range", lambda data: data["metadata"].__setitem__("graduation_year", 1900)),
        ("advisor title", lambda data: data["metadata"].__setitem__("advisor", "李华教授")),
        ("title capacity", lambda data: data["metadata"].__setitem__("title", "超长任务书题目" * 30)),
        ("empty content", lambda data: data["sections"].__setitem__("basic_content_and_requirements", [])),
        ("empty references", lambda data: data["sections"].__setitem__("references", [])),
        ("invalid nature", lambda data: data["sections"]["topic_information"].__setitem__("nature", ["graduation_design", "graduation_thesis"])),
        ("invalid source", lambda data: data["sections"]["topic_information"]["source"].__setitem__("type", "both")),
    ]
    for label, mutate in cases:
        candidate = copy.deepcopy(minimal)
        mutate(candidate)
        expect_invalid(lambda candidate=candidate: word_renderer.validate_data(candidate), label)

    missing_metadata = copy.deepcopy(minimal)
    missing_metadata["metadata"].pop("student_id")
    expect_invalid(lambda: word_renderer.validate_data(missing_metadata), "missing metadata")
    unknown_metadata = copy.deepcopy(minimal)
    unknown_metadata["metadata"]["teacher_comment"] = "not allowed"
    expect_invalid(lambda: word_renderer.validate_data(unknown_metadata), "unknown metadata")

    invalid_research = copy.deepcopy(minimal)
    invalid_research["sections"]["topic_information"]["source"] = {
        "type": "research_project",
        "research_project": {"level": "national"},
    }
    expect_invalid(lambda: word_renderer.validate_data(invalid_research), "research number")
    invalid_other = copy.deepcopy(minimal)
    invalid_other["sections"]["topic_information"]["source"] = {
        "type": "research_project",
        "research_project": {"level": "other", "project_number": "X"},
    }
    expect_invalid(lambda: word_renderer.validate_data(invalid_other), "other level label")
    mixed_source = copy.deepcopy(minimal)
    mixed_source["sections"]["topic_information"]["source"] = {
        "type": "self_proposed",
        "self_proposed_by": "student",
        "research_project": {"level": "national", "project_number": "X"},
    }
    expect_invalid(lambda: word_renderer.validate_data(mixed_source), "mutually exclusive source")

    too_deep_list: dict[str, Any] = {
        "type": "unordered_list",
        "items": [{"content": "第五级"}],
    }
    for level in ("第四级", "第三级", "第二级", "第一级"):
        too_deep_list = {
            "type": "ordered_list",
            "items": [{"content": level, "children": too_deep_list}],
        }
    too_deep = copy.deepcopy(minimal)
    too_deep["sections"]["basic_content_and_requirements"] = [too_deep_list]
    expect_invalid(lambda: word_renderer.validate_data(too_deep), "fifth list level")
    expect_invalid(lambda: latex_renderer._load_model(task_dir / "latex").validate_data(too_deep), "fifth list level in LaTeX model")

    unordered_marker = copy.deepcopy(minimal)
    unordered_marker["sections"]["basic_content_and_requirements"] = [
        {
            "type": "unordered_list",
            "items": [{"marker": "-", "content": "无序列表不能自定义标记"}],
        }
    ]
    expect_invalid(
        lambda: word_renderer.validate_data(unordered_marker),
        "marker on unordered list",
    )

    remote = copy.deepcopy(minimal)
    remote["sections"]["basic_content_and_requirements"] = [
        {"type": "image", "path": "https://example.com/a.png", "alt": "remote"}
    ]
    normalized_remote = word_renderer.validate_data(remote)
    path_text = normalized_remote["sections"]["basic_content_and_requirements"][0]["path"]
    expect_invalid(
        lambda: word_renderer._resolve_image(path_text, task_dir / "fixtures"),
        "remote Word image",
    )
    expect_invalid(
        lambda: latex_renderer._resolve_image(path_text, task_dir / "fixtures"),
        "remote LaTeX image",
    )
    expect_invalid(
        lambda: word_renderer._resolve_image(
            "assets/does-not-exist.png", task_dir / "fixtures"
        ),
        "missing Word image",
    )
    expect_invalid(
        lambda: latex_renderer._resolve_image(
            "assets/does-not-exist.png", task_dir / "fixtures"
        ),
        "missing LaTeX image",
    )


def assert_cover_title_capacity(word_renderer, layout: dict[str, Any]) -> None:
    style = layout["typography"]["cover_title_value"]
    line_capacity = float(layout["cover"]["title_underline_width_mm"]) / (
        float(style["size_pt"]) * 25.4 / 72.0
    )
    too_long = "超" * (math.floor(line_capacity * 2) + 1)
    assert display_width(too_long) <= float(word_renderer.MODEL.TITLE_DISPLAY_WIDTH_LIMIT)
    expect_invalid(
        lambda: word_renderer._split_runs_for_cover_title(
            [
                {
                    "text": too_long,
                    "script": "normal",
                    "italic": False,
                    "bold": False,
                }
            ],
            line_capacity=line_capacity,
            max_lines=int(layout["cover"]["title_max_lines"]),
        ),
        "third cover title line",
    )


def assert_word_cant_split_exceptions(
    *,
    minimal: dict[str, Any],
    task_dir: Path,
    project_dir: Path,
    output_root: Path,
) -> None:
    """Exercise each flowing row with synthetic, non-private over-page content."""

    cases: dict[int, Callable[[dict[str, Any]], None]] = {
        1: lambda data: data["sections"].__setitem__(
            "basic_content_and_requirements", ["超长流内容" * 600]
        ),
        2: lambda data: data["sections"].__setitem__(
            "schedule",
            [
                {"period": "第1阶段", "content": "进度内容" * 340},
                {"period": "第2阶段", "content": "进度内容" * 340},
            ],
        ),
        3: lambda data: data["sections"].__setitem__(
            "required_materials", ["超长流内容" * 600]
        ),
    }
    for target_row, mutate in cases.items():
        candidate = copy.deepcopy(minimal)
        mutate(candidate)
        case_dir = output_root / "cant-split" / f"row-{target_row}"
        data_path = case_dir / "synthetic.json"
        output = case_dir / "task-book.docx"
        case_dir.mkdir(parents=True, exist_ok=True)
        data_path.write_text(
            json.dumps(candidate, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        run(
            [
                sys.executable,
                str(task_dir / "word" / "render.py"),
                "--data",
                str(data_path),
                "--output",
                str(output),
                "--overwrite",
            ],
            cwd=project_dir,
        )
        table = Document(output).tables[-1]
        assert not row_has_cant_split(table.rows[target_row])
        for row_index, row in enumerate(table.rows):
            if row_index != target_row:
                assert row_has_cant_split(row)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-pdf", action="store_true")
    args = parser.parse_args()

    task_dir = Path(__file__).resolve().parent
    project_dir = task_dir.parents[1]
    output_root = project_dir / "tmp" / "task-book-tests"
    output_root.mkdir(parents=True, exist_ok=True)
    layout = load_process_document_layout(task_dir / "spec" / "layout.json")
    word_renderer = load_renderer(task_dir / "word" / "render.py", "task_book_word_test")
    latex_renderer = load_renderer(task_dir / "latex" / "render.py", "task_book_latex_test")
    fixed = word_renderer._fixed_content()
    latex_source = (task_dir / "latex" / "main.tex").read_text(encoding="utf-8")
    assert r"\newcommand{\CohortText}[1]{{\heitiBold\heitiLatin" in latex_source
    assert r"\newcommand{\TaskCheckbox}" in latex_source
    assert r"\newcommand{\TaskListItem}[3]" in latex_source
    assert r"\input{sztu-process-form.tex}" in latex_source
    assert r"\newcommand{\TaskTightJoin}" not in latex_source
    assert r"\RequiredMaterialsReferenceGap" in latex_source
    assert r"rectangle (\SZTUCheckboxSize,\SZTUCheckboxSize)" in latex_source
    assert r"\TaskCoverTitleValueWidth][c]" in latex_source
    assert r"\SZTUCoverContentWidth][l]" in latex_source
    for forbidden in (r"\blacksquare", "TaskNoticeLongItem", "3.25"):
        assert forbidden not in latex_source
    emphasis_tex = latex_renderer.rich_runs(
        [
            {"text": "中文加粗", "bold": True, "italic": False, "script": "normal"},
            {"text": "中文倾斜", "bold": False, "italic": True, "script": "normal"},
            {"text": "中文加粗倾斜", "bold": True, "italic": True, "script": "normal"},
        ]
    )
    emphasis_tex = emphasis_tex.replace(r"\nobreak{}", "")
    assert r"\textbf{中文加粗}" in emphasis_tex
    assert r"\textit{中文倾斜}" in emphasis_tex
    assert latex_renderer._render_rich_run(
        {
            "text": "中文加粗倾斜",
            "bold": True,
            "italic": True,
            "script": "normal",
        },
        "中文加粗倾斜",
    ) == r"\textbf{\textit{中文加粗倾斜}}"
    font_dir = output_root / "synthetic-fonts"
    font_tex = latex_renderer.fonts_tex(
        {
            "STXingkai": font_dir / "school.ttf",
            "SimHei": font_dir / "simhei.ttf",
            "SimSun": font_dir / "simsun.ttf",
            "Times New Roman": font_dir / "times.ttf",
            "Times New Roman Bold": font_dir / "timesbd.ttf",
            "Times New Roman Italic": font_dir / "timesi.ttf",
            "Times New Roman Bold Italic": font_dir / "timesbi.ttf",
        }
    )
    assert "AutoFakeBold=3,AutoFakeSlant=0.2" in font_tex
    assert font_tex.count("AutoFakeSlant=0.2") >= 8

    assert tuple(layout["table"]["column_widths_mm"]) == (150.32,)
    assert layout["table"]["border_pt"] == 0.5
    assert layout["table"]["flow_vertical_padding_mm"] == 1.0
    assert layout["table"]["top_bottom_padding_mm"] == 1.0
    assert layout["table"]["flow_end_space_mm"] == 0.0
    assert layout["table"]["section_title_content_gap_mm"] == 1.0
    assert layout["cover"]["title_max_lines"] == 2
    assert layout["cover"]["title_second_line_alignment"] == "left"
    assert layout["cover"]["title_underline_clearance_mm"] == 0.3
    assert layout["cover"]["title_latex_single_line_rule_offset_mm"] == 1.1
    assert layout["cover"]["title_latex_single_line_subscript_rule_offset_mm"] == 1.4
    assert layout["cover"]["title_latex_two_line_rule_offset_mm"] == 0.6
    assert layout["cover"]["title_latex_two_line_subscript_rule_offset_mm"] == 1.5
    assert layout["cover"]["title_cell_vertical_padding_mm"] == 0.5
    assert layout["paragraphs"]["notice_line_spacing"] == 1.5
    assert layout["paragraphs"]["list_first_level_indent_em"] == 2.0
    assert layout["paragraphs"]["list_level_indent_em"] == 2.0
    assert layout["paragraphs"]["list_hanging_indent_em"] == 0.0
    assert layout["paragraphs"]["list_marker_gap_em"] == 0.25
    assert layout["paragraphs"]["list_max_depth"] == 4
    assert layout["paragraphs"]["unordered_list_markers"] == ["•", "◦", "▪", "▫"]
    assert layout["notice"]["item_space_after_pt"] == 0.0
    assert layout["section_min_heights_mm"]["title"] == 12.0
    assert layout["checkbox"]["style"] == "outline_with_tick"
    assert layout["typography"]["cover_title_value"]["cjk_word_family"] == "黑体"
    assert layout["typography"]["cover_title_value"]["size_pt"] == 16.0
    assert layout["typography"]["cover_value"]["cjk_word_family"] == "宋体"
    assert layout["typography"]["cover_value"]["size_pt"] == 16.0
    assert layout["typography"]["title_row_value"]["cjk_word_family"] == "黑体"
    assert layout["typography"]["title_row_value"]["size_pt"] == 12.0
    assert layout["typography"]["body"]["cjk_word_family"] == "宋体"
    assert layout["typography"]["body"]["size_pt"] == 12.0
    assert layout["typography"]["topic_body"]["cjk_word_family"] == "宋体"
    assert layout["typography"]["topic_body"]["size_pt"] == 10.5
    assert layout["typography"]["signature"]["cjk_word_family"] == "黑体"
    assert layout["typography"]["signature"]["size_pt"] == 12.0
    assert layout["signature"]["advisor_blank_width_mm"] == 45.0
    assert layout["signature"]["college_leader_blank_width_mm"] == 50.0
    assert layout["signature"]["right_inset_mm"] == 7.5
    assert layout["signature"]["topic_project_value_baseline_shift_mm"] == 0.4
    assert layout["required_font_files"] == [
        "STXingkai",
        "SimHei",
        "SimSun",
        "Times New Roman",
        "Times New Roman Bold",
        "Times New Roman Bold Italic",
        "Times New Roman Italic",
    ]
    assert starts_with_calendar_date("2026年1—2月")
    assert starts_with_calendar_date("2025年10月—11月")
    assert starts_with_calendar_date("2026年1月5日至2026年1月18日")
    assert starts_with_calendar_date("10—11月")
    assert starts_with_calendar_date("2026-01")
    assert starts_with_calendar_date("2026/01")
    assert starts_with_calendar_date("2026.01")
    assert not starts_with_calendar_date("第1—2周")
    assert not starts_with_calendar_date("1—2周")
    assert not starts_with_calendar_date("2026年第1—2周")
    assert not starts_with_calendar_date("2026年1—2周")
    assert not starts_with_calendar_date("2026年13月")
    assert_cover_title_capacity(word_renderer, layout)

    fixture_data: dict[str, dict[str, Any]] = {}
    normalized_data: dict[str, dict[str, Any]] = {}
    schedule_indent_cases: set[bool] = set()
    page_range_recipes = load_page_range_recipes(
        task_dir.parent / "common" / "fixtures" / "page-range-stress-recipes.json"
    )
    page_range_recipe = page_range_recipes["task-book"]
    page_range_fixture = materialize_page_range_fixture(
        task_dir / "fixtures" / page_range_recipe["base_fixture"],
        page_range_recipe,
        output_root / "page-range-maximum" / "input.json",
    )
    for name in (*FIXTURE_NAMES, "page-range-maximum"):
        fixture = (
            page_range_fixture
            if name == "page-range-maximum"
            else task_dir / "fixtures" / f"{name}.json"
        )
        raw = json.loads(fixture.read_text(encoding="utf-8"))
        normalized = word_renderer.validate_data(raw)
        if name == "minimal":
            assert normalized["sections"]["required_materials"] == []
        fixture_data[name] = raw
        normalized_data[name] = normalized
        assert normalized == latex_renderer._load_model(task_dir / "latex").validate_data(raw)

        docx_output = output_root / name / "task-book.docx"
        run(
            [
                sys.executable,
                str(task_dir / "word" / "render.py"),
                "--data",
                str(fixture),
                "--output",
                str(docx_output),
                "--overwrite",
            ],
            cwd=project_dir,
        )
        document = Document(docx_output)
        assert len(document.sections) == 1
        section = document.sections[0]
        assert abs(section.page_width.mm - 210.0) <= 0.02
        assert abs(section.page_height.mm - 297.0) <= 0.02
        assert abs(section.top_margin.mm - 25.4) <= 0.02
        assert abs(section.bottom_margin.mm - 25.4) <= 0.02
        assert abs(section.left_margin.mm - 31.75) <= 0.02
        assert abs(section.right_margin.mm - 31.75) <= 0.02
        assert len(document.element.xpath('.//w:br[@w:type="page"]')) == 1
        assert len(document.element.xpath(".//w:pageBreakBefore")) == 1
        assert len(document.tables) == 3
        assert_cover_structure(document, normalized, layout, fixed)
        main_table = document.tables[2]
        assert (len(main_table.rows), len(main_table.columns)) == (6, 1)
        for table in document.tables:
            assert_table_geometry(table)
            for nested in nested_tables(table):
                assert_table_geometry(nested)
        assert_main_table_position(main_table, layout)
        assert_main_table_borders(main_table, border_pt=layout["table"]["border_pt"])

        expected_left_right = twips_from_mm(layout["table"]["horizontal_padding_mm"])
        for index, row in enumerate(main_table.rows):
            cell = row.cells[0]
            assert cell_margin_dxa(cell, "left") == expected_left_right
            assert cell_margin_dxa(cell, "right") == expected_left_right
            vertical_mm = layout["table"]["flow_vertical_padding_mm"] if index in {1, 2, 3} else layout["table"]["top_bottom_padding_mm"]
            assert cell_margin_dxa(cell, "top") == twips_from_mm(vertical_mm)
            assert cell_margin_dxa(cell, "bottom") == twips_from_mm(vertical_mm)
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.TOP
            if index in {0, 4, 5}:
                key = {
                    0: "title",
                    4: "topic_information",
                    5: "college_leader_opinion",
                }[index]
                assert row.height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
                assert row.height is not None
                assert abs(row.height.mm - layout["section_min_heights_mm"][key]) <= 0.02
            else:
                tr_pr = row._tr.trPr
                assert tr_pr is None or tr_pr.find(qn("w:trHeight")) is None
                assert row.height is None and row.height_rule is None
            expected_cant_split = not (
                index in {1, 2, 3}
                and word_renderer._flow_cell_may_exceed_page(cell)
            )
            assert row_has_cant_split(row) is expected_cant_split
            if index in {1, 2, 3}:
                assert paragraph_has_content(cell.paragraphs[-1])

        advisor_table = main_table.rows[4].cells[0].tables[-1]
        college_table = main_table.rows[5].cells[0].tables[0]
        assert advisor_table.rows[0].cells[1].text == fixed["body"]["signature_labels"]["advisor"]
        assert_signature_table(
            college_table,
            blank_mm=layout["signature"]["college_leader_blank_width_mm"],
            right_mm=layout["signature"]["right_inset_mm"],
        )
        assert_date_table(main_table.rows[5].cells[0].tables[1], layout, fixed)
        schedule_indent_cases.update(assert_word_indents(document, normalized, layout))
        if name == "minimal":
            materials_text = [
                paragraph.text
                for paragraph in main_table.rows[3].cells[0].paragraphs
                if paragraph.text.strip()
            ]
            assert materials_text == [
                fixed["body"]["section_labels"]["required_materials_and_references"],
                runs_text(normalized["sections"]["references"][0]),
            ]
        schedule_paragraphs = [
            item
            for item in main_table.rows[2].cells[0].paragraphs[1:]
            if item.text.strip()
        ]
        if name == "layout-stress":
            assert [item.text.split("：", 1)[0] for item in schedule_paragraphs] == [
                "2026年1—2月",
                "2025年10月—11月",
            ]
            assert all(
                abs(points_or_zero(item.paragraph_format.first_line_indent)) <= 0.01
                for item in schedule_paragraphs
            )
        if name == "normal":
            assert schedule_paragraphs[0].text.startswith("第1—2周：")
            assert abs(
                points_or_zero(schedule_paragraphs[0].paragraph_format.first_line_indent)
                - layout["typography"]["body"]["size_pt"]
                * layout["paragraphs"]["first_line_indent_em"]
            ) <= 0.01
        assert_word_topic_choices(
            document,
            normalized["sections"]["topic_information"],
            layout,
            fixed,
        )
        if name == "normal":
            assert_word_typography(document, normalized, layout, fixed)
        if name == "rich-text-list":
            image_alt = [
                item.get("descr")
                for item in document.element.xpath(".//wp:docPr")
            ]
            assert image_alt == ["虚构的材料制备、表征、反应和结果分析流程示意图"]
            all_runs = [
                item
                for table in document.tables
                for row in table.rows
                for cell in row.cells
                for item in direct_cell_runs(cell)
            ]
            assert any(item.text == "3" and item.font.subscript for item in all_runs)
            superscript_text = "".join(
                item.text for item in all_runs if item.font.superscript
            )
            assert "−3" in superscript_text
            assert "−1" in superscript_text
            assert any(item.text == "k" and item.bold and item.italic for item in all_runs)
            assert any(item.text == "实验要求：" and item.bold for item in all_runs)
            assert any(item.text == "中文加粗" and item.bold for item in all_runs)
            assert any(item.text == "中文倾斜" and item.italic for item in all_runs)
            assert any(
                item.text == "中文加粗倾斜" and item.bold and item.italic
                for item in all_runs
            )

        if name == "reference-overflow":
            assert normalized["metadata"]["student_name"].startswith("测试")
            assert normalized["metadata"]["student_id"].startswith("202600000")
            assert len(normalized["sections"]["references"]) == 24

        if name == "image-page-break":
            assert normalized["metadata"]["student_name"].startswith("测试")
            assert normalized["metadata"]["student_id"].startswith("202600000")
            image_alt = [
                item.get("descr")
                for item in document.element.xpath(".//wp:docPr")
            ]
            assert image_alt == ["虚构的材料制备、表征、反应和结果分析流程示意图"]
            basic_paragraphs = main_table.rows[1].cells[0].paragraphs
            caption_index = next(
                index
                for index, paragraph in enumerate(basic_paragraphs)
                if paragraph.text == "图 1  合成实验流程与数据复核路径"
            )
            trailing_index = next(
                index
                for index, paragraph in enumerate(basic_paragraphs)
                if paragraph.text.startswith("图片之后的正文必须继续")
            )
            assert trailing_index == caption_index + 1
            trailing = basic_paragraphs[trailing_index]
            assert abs(points_or_zero(trailing.paragraph_format.left_indent)) <= 0.01
            assert abs(
                points_or_zero(trailing.paragraph_format.first_line_indent)
                - layout["typography"]["body"]["size_pt"]
                * layout["paragraphs"]["first_line_indent_em"]
            ) <= 0.01

        if name == "structured-content":
            basic_blocks = normalized["sections"]["basic_content_and_requirements"]
            table_block = next(item for item in basic_blocks if item["type"] == "data_table")
            equation = next(item for item in basic_blocks if item["type"] == "equation")
            assert len(table_block["columns"]) == 4
            assert len(table_block["rows"]) == 3
            assert equation["expression"]["type"] == "row"
            assert equation["equation_label"] == "(1-1)"
            assert "{{eq:" not in basic_blocks[0]["runs"][0]["text"]
            assert "(1-1)" in basic_blocks[0]["runs"][0]["text"]
            group = next(
                item
                for item in normalized["sections"]["required_materials"]
                if item["type"] == "figure_group"
            )
            assert group["figure_label"] == "图 1-1"
            assert [item["subfigure_label"] for item in group["items"]] == ["（a）", "（b）"]
            embedded = main_table.rows[1].cells[0].tables[0]
            assert len(embedded.rows) == 4 and len(embedded.columns) == 4
            assert embedded.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
            assert all(row_has_cant_split(row) for row in embedded.rows)
            group_table = main_table.rows[3].cells[0].tables[0]
            assert len(group_table.rows) == 1 and len(group_table.columns) == 2
            assert row_has_cant_split(group_table.rows[0])
            image_alt = [item.get("descr") for item in document.element.xpath(".//wp:docPr")]
            assert image_alt == [
                "虚构的结构化数据输入流程图",
                "虚构的双路文档生成流程图",
            ]
            for math_tag in ("m:oMathPara", "m:f", "m:sSub", "m:sSup", "m:rad"):
                assert f"<{math_tag}" in document.element.xml
            assert "效率等于输出浓度" in document.element.xml
            equation_table = main_table.rows[1].cells[0].tables[1]
            assert equation_table.rows[0].cells[2].text == "(1-1)"
            assert row_has_cant_split(equation_table.rows[0])
            assert all(
                border.get(qn("w:val")) == "nil"
                for border in equation_table._tbl.tblPr.find(qn("w:tblBorders"))
            )
            assert equation_table.rows[0].cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

        if name == "nested-list":
            basic_paragraphs = main_table.rows[1].cells[0].paragraphs
            basic_texts = [item.text for item in basic_paragraphs if item.text.strip()]
            for expected in (
                "•\u2009一级无序：建立公共内容模型",
                "（1）\u2009二级有序：表达 H2O 富文本",
                "▪\u2009三级无序：保持双路渲染一致",
                "A.\u2009四级有序：验证深度上限",
                "•\u2009一级无序：验证同级项目连续排列",
                "列表结束后，本段恢复为普通正文首行缩进。",
            ):
                assert expected in basic_texts
            nested_runs = [
                item
                for paragraph in basic_paragraphs
                for item in visible_runs(paragraph)
            ]
            assert any(item.text == "2" and item.font.subscript for item in nested_runs)
            trailing = next(
                item
                for item in basic_paragraphs
                if item.text == "列表结束后，本段恢复为普通正文首行缩进。"
            )
            assert abs(points_or_zero(trailing.paragraph_format.left_indent)) <= 0.01
            assert abs(
                points_or_zero(trailing.paragraph_format.first_line_indent)
                - layout["typography"]["body"]["size_pt"]
                * layout["paragraphs"]["first_line_indent_em"]
            ) <= 0.01

            materials_texts = [
                item.text
                for item in main_table.rows[3].cells[0].paragraphs
                if item.text.strip()
            ]
            for expected in (
                "1.\u2009收集官方模板与规范",
                "2.\u2009整理结构化验收记录",
                "◦\u2009记录 Word 结果",
                "◦\u2009记录 LaTeX 结果",
                "资料列表结束后继续填写普通段落。",
            ):
                assert expected in materials_texts

            rendered_blocks = latex_renderer.render_blocks(
                normalized["sections"]["basic_content_and_requirements"],
                data_dir=fixture.parent,
                assets_dir=output_root / name / "latex-structure-assets",
            ).replace(r"\nobreak{}", "")
            for expected in (
                r"\TaskListItem{2}{•}{一级无序：建立公共内容模型",
                r"\TaskListItem{4}{（1）}{二级有序：表达 H\textsubscript{2}O",
                r"\TaskListItem{6}{▪}{三级无序：保持双路渲染一致",
                r"\TaskListItem{8}{A.}{四级有序：验证深度上限",
                r"\TaskBodyParagraph{列表结束后",
            ):
                assert expected in rendered_blocks

        if args.skip_pdf:
            continue

        latex_dir = output_root / name / "latex"
        build_output = run(
            [
                sys.executable,
                str(task_dir / "latex" / "render.py"),
                "--data",
                str(fixture),
                "--output-dir",
                str(latex_dir),
                "--overwrite",
                "--compile",
            ],
            cwd=project_dir,
        )
        log_text = (latex_dir / "main.log").read_text(encoding="utf-8", errors="replace")
        for warning in (
            "Overfull",
            "Underfull",
            "Missing character",
            "LaTeX Font Warning",
            "Some font shapes were not available",
        ):
            assert warning not in build_output
            assert warning not in log_text

        typography_tex = (latex_dir / "task-book-typography.tex").read_text(encoding="utf-8")
        fonts_tex = (latex_dir / "task-book-fonts.tex").read_text(encoding="utf-8")
        data_tex = (latex_dir / "task-book-data.tex").read_text(encoding="utf-8")
        process_form_tex = (latex_dir / "sztu-process-form.tex").read_text(
            encoding="utf-8"
        )
        assert r"\newfontfamily\heitiLatin" in fonts_tex
        assert r"\newCJKfontfamily\heiti[" in fonts_tex
        assert r"\newCJKfontfamily\songti[" in fonts_tex
        assert fonts_tex.count("AutoFakeBold=3") >= 5
        assert r"\newcommand{\SZTUFormRuleWidth}{0.5pt}" in typography_tex
        assert r"\newcommand{\SZTUTitleSingleLineUnderlineOffset}{1.1mm}" in typography_tex
        assert r"\newcommand{\SZTUTitleSingleLineSubscriptUnderlineOffset}{1.4mm}" in typography_tex
        assert r"\newcommand{\SZTUTitleTwoLineUnderlineOffset}{0.6mm}" in typography_tex
        assert r"\newcommand{\SZTUTitleTwoLineSubscriptUnderlineOffset}{1.5mm}" in typography_tex
        assert r"\newcommand{\SZTUFlowVerticalPadding}{1mm}" in typography_tex
        assert r"\newcommand{\SZTUFixedVerticalPadding}{1mm}" in typography_tex
        assert r"\newcommand{\SZTUSectionTitleContentGap}{1mm}" in typography_tex
        assert "sztuformflow/.style" in process_form_tex
        assert r"\newcommand{\SZTUTitleRowMinHeight}{12mm}" in typography_tex
        assert r"\newcommand{\SZTUNoticeItemGap}{0pt}" in typography_tex
        assert r"\newcommand{\SZTUListHangingIndent}{0em}" in typography_tex
        assert r"\newcommand{\SZTUTeacherSignatureBlank}{45mm}" in typography_tex
        assert r"\newcommand{\SZTUCollegeSignatureBlank}{50mm}" in typography_tex
        assert r"\newcommand{\SZTUSignatureRightInset}{7.5mm}" in typography_tex
        assert r"\newcommand{\SZTUTopicProjectValueShift}{0.4mm}" in typography_tex
        if name == "layout-stress":
            assert r"\TaskScheduleDateParagraph{2026年1—2月：" in data_tex
            assert r"\TaskScheduleDateParagraph{2025年10月—11月：" in data_tex
        if name == "normal":
            assert r"\TaskBodyParagraph{第1—2周：" in data_tex
        if name == "minimal":
            assert r"\long\def\RequiredMaterialsContent{}" in data_tex
            assert r"\long\def\RequiredMaterialsReferenceGap{}" in data_tex
        if name == "nested-list":
            nested_data_tex = data_tex.replace(r"\nobreak{}", "")
            for expected in (
                r"\TaskListItem{2}{•}{一级无序：建立公共内容模型",
                r"\TaskListItem{4}{（1）}{二级有序：表达 H\textsubscript{2}O",
                r"\TaskListItem{6}{▪}{三级无序：保持双路渲染一致",
                r"\TaskListItem{8}{A.}{四级有序：验证深度上限",
                r"\TaskBodyParagraph{列表结束后",
            ):
                assert expected in nested_data_tex
        if name == "cover-subscript":
            assert r"\TaskCoverTitleLineCount{2}" in data_tex
            assert r"\TaskCoverTitleLineOneHasSubscript{1}" in data_tex
            assert r"\TaskCoverTitleLineTwoHasSubscript{1}" in data_tex
            assert data_tex.count(r"\textsubscript{") >= 4
        if name == "structured-content":
            assert r"\begin{tblr}" in data_tex
            assert r"\begin{minipage}[t]{0.4894\linewidth}" in data_tex
            assert r"图 1\mbox{-}1" in data_tex
            assert "（a）" in data_tex and "（b）" in data_tex
            assert r"\frac{" in data_tex and r"_{out}" in data_tex
            assert r"\sqrt{{x}^{2}}" in data_tex
            assert r"\input" not in data_tex
            assert "(1-1)" in data_tex and "{{eq:" not in data_tex

        pdf_path = latex_dir / "main.pdf"
        info = run(["pdfinfo", str(pdf_path)], cwd=project_dir)
        assert "Page size:       595.28 x 841.89 pts (A4)" in info
        page_match = re.search(r"^Pages:\s+(\d+)$", info, re.MULTILINE)
        assert page_match is not None
        page_count = int(page_match.group(1))
        if name == "page-range-maximum":
            minimum, maximum = page_range_recipe["expected_pdf_pages"]
            assert minimum <= page_count <= maximum
        else:
            assert 2 <= page_count <= 12
        fonts = run(["pdffonts", str(pdf_path)], cwd=project_dir)
        assert_pdf_fonts(fonts)
        extracted = run(["pdftotext", str(pdf_path), "-"], cwd=project_dir)
        compact = normalized_visible_text(extracted)
        for expected in content_texts(normalized):
            if normalized_visible_text(expected) not in compact:
                assert (
                    normalized_hyphenation_text(expected)
                    in normalized_hyphenation_text(extracted)
                ), expected
        for fixed_text in (
            "深圳技术大学",
            "本科生毕业论文（设计）须知",
            "题目名称：",
            "一、毕业论文(设计)基本内容与要求：",
            "二、进度安排：",
            "三、需收集的资料和指导性参考文献：",
            "四、选题信息：",
            "学院领导意见：",
        ):
            assert normalized_visible_text(fixed_text) in compact
        gate = run(
            [
                sys.executable,
                str(project_dir / "scripts" / "validate_cjk_render.py"),
                str(pdf_path),
            ],
            cwd=project_dir,
        )
        assert "PASS" in gate and "100.0%" in gate

        with pdfplumber.open(pdf_path) as pdf:
            line_count_match = re.search(
                r"\\TaskCoverTitleLineCount\{([12])\}", data_tex
            )
            assert line_count_match is not None
            assert_pdf_cover_geometry(
                pdf,
                normalized=normalized,
                layout=layout,
                title_line_count=int(line_count_match.group(1)),
            )
            assert_pdf_notice_geometry(pdf)
            assert_pdf_title_row_geometry(pdf, layout)
            assert_pdf_form_geometry(pdf, fixture_name=name)
            assert_pdf_closing_bundle(pdf, layout, normalized)
            if name == "list-wrap":
                lines = [
                    line
                    for page in pdf.pages
                    for line in page.extract_text_lines(return_chars=True)
                ]
                for prefix, marker_chars in (
                    ("• 一级无序：验证项目", 1),
                    ("1.收集官方模板", 2),
                ):
                    first_index = next(
                        index for index, line in enumerate(lines)
                        if line["text"].startswith(prefix)
                    )
                    first_line = lines[first_index]
                    continuation = lines[first_index + 1]
                    assert 10.0 <= continuation["top"] - first_line["top"] <= 24.0
                    first_body_x = float(first_line["chars"][marker_chars]["x0"])
                    continuation_x = float(continuation["chars"][0]["x0"])
                    assert abs(first_body_x - continuation_x) <= 1.5, (
                        prefix, first_body_x, continuation_x
                    )
            chars = [
                char
                for page in pdf.pages
                for char in page.chars
                if char.get("text", "").strip()
            ]
            assert any("STXingkai" in char["fontname"] and abs(float(char["size"]) - 36.0) <= 0.05 for char in chars)
            assert any("SimHei" in char["fontname"] and abs(float(char["size"]) - 24.0) <= 0.05 for char in chars)
            assert any("SimSun" in char["fontname"] and abs(float(char["size"]) - 10.5) <= 0.05 for char in chars)
            if name == "rich-text-list":
                assert any(page.images for page in pdf.pages)
                small_scripts = [
                    char
                    for char in chars
                    if char["text"] in {"−", "1", "2", "3", "4"}
                    and float(char["size"]) < 8.0
                ]
                assert small_scripts
                assert r"\textsubscript{3}" in data_tex
                assert r"\textsubscript{4}" in data_tex
                assert r"\textsuperscript{−3}" in data_tex
                assert r"\textsuperscript{−1}" in data_tex
                assert r"\textbf{实验要求：}" in data_tex
                assert r"\TaskFigure{assets/" in data_tex
            if name == "reference-overflow":
                reference_pages = []
                for reference in normalized["sections"]["references"]:
                    expected = normalized_hyphenation_text(runs_text(reference))[:50]
                    page_index = next(
                        index
                        for index, page in enumerate(pdf.pages)
                        if expected
                        in normalized_hyphenation_text(page.extract_text() or "")
                    )
                    reference_pages.append(page_index)
                assert reference_pages == sorted(reference_pages)
                assert len(set(reference_pages)) >= 2
                topic_page = next(
                    index
                    for index, page in enumerate(pdf.pages)
                    if normalized_visible_text("四、选题信息：")
                    in normalized_visible_text(page.extract_text() or "")
                )
                assert topic_page >= reference_pages[-1]
            if name == "image-page-break":
                caption_text = normalized_visible_text(
                    "图 1  合成实验流程与数据复核路径"
                )
                trailing_text = normalized_visible_text(
                    "图片之后的正文必须继续保持正常两字符首行缩进"
                )
                caption_page = next(
                    index
                    for index, page in enumerate(pdf.pages)
                    if caption_text in normalized_visible_text(page.extract_text() or "")
                )
                trailing_page = next(
                    index
                    for index, page in enumerate(pdf.pages)
                    if trailing_text in normalized_visible_text(page.extract_text() or "")
                )
                assert trailing_page >= caption_page
                assert pdf.pages[caption_page].images
                figure = max(
                    pdf.pages[caption_page].images,
                    key=lambda item: float(item["x1"]) - float(item["x0"]),
                )
                assert float(figure["x0"]) >= 90.0
                assert float(figure["x1"]) <= 516.2
            if name == "structured-content":
                assert any(page.images for page in pdf.pages)
                assert "图1-1" in normalized_visible_text(extracted)
                assert "（a）数据输入" in extracted
                assert "（b）双路生成" in extracted

    assert schedule_indent_cases == {False, True}
    assert_word_cant_split_exceptions(
        minimal=fixture_data["minimal"],
        task_dir=task_dir,
        project_dir=project_dir,
        output_root=output_root,
    )

    assert_negative_model_cases(
        word_renderer,
        latex_renderer,
        fixture_data["minimal"],
        task_dir,
    )

    unified_output = output_root / "unified-minimal"
    run(
        [
            sys.executable,
            str(task_dir / "render.py"),
            "--data",
            str(task_dir / "fixtures" / "minimal.json"),
            "--output-dir",
            str(unified_output),
            "--format",
            "all",
            "--overwrite",
        ],
        cwd=project_dir,
    )
    assert (unified_output / "task-book.docx").is_file()
    assert (unified_output / "latex" / "main.tex").is_file()

    print("task-book template regression checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
