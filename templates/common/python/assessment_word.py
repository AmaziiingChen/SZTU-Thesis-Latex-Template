"""Source-preserving DOCX renderer for fixed one-page official assessment forms."""
from __future__ import annotations
import argparse,json,sys,math,tempfile
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt,Mm
from docx.enum.table import WD_ROW_HEIGHT_RULE,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .assessment_forms import *

def font(run,size=10.5,family='宋体',bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    f=run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii','hAnsi','cs'):f.set(qn('w:'+key),'Times New Roman')
    f.set(qn('w:eastAsia'),family)

def append(p,runs,size=10.5,family='宋体'):
    for item in runs:
        r=p.add_run(item['text']);font(r,size,family,item.get('bold',False));r.italic=item.get('italic',False)
        if item.get('script')=='super':r.font.superscript=True
        elif item.get('script')=='sub':r.font.subscript=True

def paragraph(p,*,size=10.5,center=False,indent=0):
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(14 if size==10.5 else size*1.1);p.paragraph_format.first_line_indent=Pt(indent*size);p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    # Remove inherited character-unit and legacy tab indentation.
    ppr=p._p.get_or_add_pPr()
    for item in list(ppr):
        if item.tag==qn('w:ind'):
            for k in ('firstLineChars','startChars','endChars'):item.attrib.pop(qn('w:'+k),None)

def clear_cell(cell):
    for p in list(cell._tc):
        if p.tag!=qn('w:tcPr'):cell._tc.remove(p)
    p=OxmlElement('w:p');cell._tc.append(p)

def set_cell(cell,label,runs,size=14,center=False):
    clear_cell(cell);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;p=cell.paragraphs[0];paragraph(p,size=size,center=center)
    if label:append(p,plain(label),size,'黑体')
    append(p,runs,size)

def explicit_line_spacing(p,points):
    """Use measured physical points without a second inherited Word grid."""
    p.paragraph_format.line_spacing=Pt(points)
    ppr=p._p.get_or_add_pPr()
    for tag in ('autoSpaceDE','autoSpaceDN','snapToGrid'):
        old=ppr.find(qn('w:'+tag))
        if old is not None:ppr.remove(old)
        node=OxmlElement('w:'+tag);node.set(qn('w:val'),'false')
        spacing=ppr.find(qn('w:spacing'));ppr.insert(list(ppr).index(spacing),node)

def add_blocks(cell,content,spec):
    for marker,runs,depth in iter_paragraphs(content):
        p=cell.add_paragraph();paragraph(p,indent=2 if depth==0 else 0)
        p.paragraph_format.line_spacing=Pt(spec['paragraphs']['body_line_spacing_pt'])
        if depth:
            size=spec['typography']['body']['size_pt'];fonts=resolved_fonts(spec);advance=marker_advance_pt(marker,size_pt=size,latin_font=fonts['Times New Roman'],cjk_font=fonts['SimSun']);p.paragraph_format.left_indent=Pt(depth*2*size+advance);p.paragraph_format.first_line_indent=Pt(-advance);append(p,plain(marker+'\u2009'))
        append(p,runs)

def empty_anchor(cell):
    """A visible-free 1pt paragraph makes following cell spacing explicit."""
    p=cell.add_paragraph();paragraph(p,size=1);p.paragraph_format.line_spacing=Pt(1)
    mark=p._p.get_or_add_pPr();rpr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'2');rpr.append(sz);mark.append(rpr)

def section_cell(cell,intro_xml,content,spec,index,signature=None,date='',guidance=None):
    clear_cell(cell);cell._tc.remove(cell.paragraphs[0]._p)
    prompt_key='questions-prompt' if spec['kind']=='defense-record' else 'division-label' if spec['kind']=='advisor-review' and index==4 else 'prompt'
    intro_region=next(r for r in spec['regions'] if r['key']==prompt_key)
    intro=None;v2_height=None
    if guidance is not None:
        field,notes=guidance;heading=v2_heading(intro_region,field)
        value=guidance_runs(heading,{'sections':{'writing_guidance':{field:notes}}})
        v2_height=guidance_height(heading,value,Measure(resolved_fonts(spec)))
        if value:
            cell._tc.append(deepcopy(intro_xml));intro=cell.paragraphs[0];intro.clear()
            paragraph(intro,size=heading['size_pt'],indent=heading.get('first_line_indent_em',0));intro.paragraph_format.left_indent=Pt(0);intro.paragraph_format.right_indent=Pt(0)
            intro.paragraph_format.line_spacing=Pt(heading['line_spacing_pt'])
            if heading['label']:
                fixed=plain(heading['label'])
                for run in fixed:run['bold']=heading['guidance_style']['fixed_label']['bold']
                append(intro,fixed,heading.get('label_size_pt',14),'黑体' if heading['label_font'].startswith('heiti') else '宋体')
            append(intro,notes,heading['size_pt'],'宋体')
            if heading.get('suffix'):
                fixed=plain(heading['suffix']['text'])
                for run in fixed:run['bold']=heading['guidance_style']['fixed_suffix']['bold']
                append(intro,fixed,heading['suffix']['size_pt'],'黑体' if heading['suffix']['font'].startswith('heiti') else '宋体')
    elif intro_xml is not None:
        cell._tc.append(deepcopy(intro_xml));intro=cell.paragraphs[0]
        paragraph(intro,indent=intro_region.get('first_line_indent_em',0));intro.paragraph_format.line_spacing=Pt(intro_region['line_spacing_pt'])
        for r in intro.runs:
            family=r._element.rPr.rFonts.get(qn('w:eastAsia'),'宋体') if r._element.rPr is not None and r._element.rPr.rFonts is not None else '宋体'
            size=r.font.size.pt if r.font.size else 10.5;font(r,size,'黑体' if '黑体' in family else '宋体')
    add_blocks(cell,content,spec)
    anchor_height=0
    if intro is None and not content:
        empty_anchor(cell);anchor_height=25.4/72
    if signature:
        region=next(r for r in spec['regions'] if r['key']=='signature');dr=next(r for r in spec['regions'] if r['key']=='date');pr=next(r for r in spec['regions'] if r['key']=='prompt')
        p=cell.add_paragraph();paragraph(p,size=region['size_pt'])
        measure=Measure(resolved_fonts(spec));intro_height=v2_height if v2_height is not None and intro is not None else measure.lines(plain(intro.text),pr['width_mm'],pr['size_pt'])*pr['line_spacing_pt']*25.4/72 if intro is not None else anchor_height
        body_height=measure.blocks_height(content,pr['width_mm'],spec)
        row_top=spec['table']['top_mm']+sum(spec['table']['rows_mm'][:index])
        p.paragraph_format.space_before=Mm(max(0,region['y_mm']-row_top-intro_height-body_height-spec['word'].get('signature_paragraph_vertical_adjust_mm',0)))
        p.paragraph_format.left_indent=Mm(region['x_mm']-spec['table']['left_mm']-spec['table']['horizontal_padding_mm'])
        p.paragraph_format.tab_stops.add_tab_stop(Mm(dr['x_mm']-spec['table']['left_mm']-spec['table']['horizontal_padding_mm']))
        append(p,plain(signature),region['size_pt'],'黑体' if region['font'].startswith('heiti') else '宋体')
        append(p,plain('\t'+(format_date(date) if date else '年     月    日')),dr['size_pt'],'黑体' if dr['font'].startswith('heiti') else '宋体')
    if not cell.paragraphs:cell.add_paragraph()

def canonicalize_package_fonts(path):
    """The legacy .doc conversion lists unrelated missing fonts in styles/themes.
    Declare only the three required form families; never allow theme fallback.
    """
    with ZipFile(path) as archive: entries={i.filename:(i,archive.read(i)) for i in archive.infolist()}
    for name,(info,content) in list(entries.items()):
        if not name.endswith('.xml') or not name.startswith('word/'):continue
        root=etree.fromstring(content);changed=False
        title_paragraphs=set(root.xpath('./w:body/w:p[position()<=2]',namespaces={'w':qn('w:p').split('}')[0][1:]})) if name=='word/document.xml' else set()
        for element in root.iter():
            tag=etree.QName(element).localname
            if tag=='rFonts':
                cjk=element.get(qn('w:eastAsia'),'宋体');cjk='黑体' if ('黑体' in cjk or cjk=='SimHei') else '宋体'
                for attr in ('ascii','hAnsi','cs'):element.set(qn('w:'+attr),'Times New Roman')
                element.set(qn('w:eastAsia'),cjk)
                ancestor=element.getparent()
                while ancestor is not None and ancestor.tag!=qn('w:p'):ancestor=ancestor.getparent()
                if ancestor in title_paragraphs:
                    for attr in ('ascii','hAnsi','cs'):element.set(qn('w:'+attr),'黑体')
                for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme'):element.attrib.pop(qn('w:'+attr),None)
                changed=True
            elif name.startswith('word/theme/') and 'typeface' in element.attrib:
                element.set('typeface','宋体' if tag=='ea' or element.get('script') in ('Hans','Hant','Jpan','Hang') else 'Times New Roman');changed=True
        if name=='word/fontTable.xml':
            for el in list(root):root.remove(el)
            for family in ('宋体','黑体','Times New Roman'):
                el=etree.SubElement(root,qn('w:font'));el.set(qn('w:name'),family)
            changed=True
        if changed:entries[name]=(info,etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True))
    with ZipFile(path,'w',compression=ZIP_DEFLATED) as archive:
        for info,content in entries.values():archive.writestr(info,content)

def fixed_footer(doc,table,footer,spec,metadata,date):
    """Keep the official signature line independent of inherited document grids.

    Legacy .doc conversions represent the table-to-signature gap as an empty
    paragraph. Its implicit line grid differs among Office engines. Use the
    measured gap as explicit spacing and disable CJK auto-spacing so a real
    date has the same glyph advances as the fixed-region preflight.
    """
    fs=spec['word']['footer']
    body=doc._element.body
    between=False
    for node in list(body):
        if node is table._tbl:between=True;continue
        if node is footer._p:break
        if between and node.tag==qn('w:p') and not ''.join(node.xpath('.//w:t/text()')):
            body.remove(node)
    after=False
    for node in list(body):
        if node is footer._p:after=True;continue
        if after and node.tag==qn('w:p') and not ''.join(node.xpath('.//w:t/text()')):
            body.remove(node)
    footer.clear();paragraph(footer)
    footer.paragraph_format.left_indent=Mm(fs['left_indent_mm'])
    footer.paragraph_format.right_indent=Mm(fs['right_indent_mm'])
    footer.paragraph_format.space_before=Mm(fs['space_before_mm'])
    footer.paragraph_format.widow_control=False
    ppr=footer._p.get_or_add_pPr()
    for tag in ('autoSpaceDE','autoSpaceDN','adjustRightInd','snapToGrid'):
        old=ppr.find(qn('w:'+tag))
        if old is not None:ppr.remove(old)
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),'false')
        spacing=ppr.find(qn('w:spacing'))
        ppr.insert(list(ppr).index(spacing) if spacing is not None else len(ppr),el)
    tabs=ppr.find(qn('w:tabs'))
    if tabs is not None:ppr.remove(tabs)
    footer.paragraph_format.tab_stops.add_tab_stop(Mm(fs['leader_tab_mm']))
    footer.paragraph_format.tab_stops.add_tab_stop(Mm(fs['date_tab_mm']))
    append(footer,plain('记录员：'+metadata['recorder']+'\t答辩小组（答辩委员会）组长：'+metadata['leader']))
    append(footer,plain('\t'+(format_date(date) or '年  月  日')))

def minimal_terminal_paragraph(doc,table,spec):
    """Word requires a paragraph after a terminal table; reserve its real height.
    Only empty conversion paragraphs are affected. It prevents the newly full
    handwriting table from acquiring an extra page due to an inherited grid.
    """
    after=False
    for node in list(doc._element.body):
        if node is table._tbl:after=True;continue
        if not after or node.tag!=qn('w:p') or ''.join(node.xpath('.//w:t/text()')):continue
        p=next(p for p in doc.paragraphs if p._p is node);paragraph(p,size=1)
        p.paragraph_format.line_spacing=Pt(spec['word']['trailing_empty_paragraph_line_pt'])
        ppr=p._p.get_or_add_pPr()
        old=ppr.find(qn('w:snapToGrid'))
        if old is not None:ppr.remove(old)
        el=OxmlElement('w:snapToGrid');el.set(qn('w:val'),'false')
        spacing=ppr.find(qn('w:spacing'));ppr.insert(list(ppr).index(spacing),el)
        for r in p.runs:font(r,1)
        mark=ppr.find(qn('w:rPr'))
        if mark is None:mark=OxmlElement('w:rPr');ppr.append(mark)
        sz=mark.find(qn('w:sz'))
        if sz is None:sz=OxmlElement('w:sz');mark.append(sz)
        sz.set(qn('w:val'),'2')

def render(kind,template,data_path,output,*,overwrite=False):
    output=Path(output);data_path=Path(data_path);template=Path(template)
    if output.exists() and not overwrite:raise FileExistsError(f'output exists; pass --overwrite: {output}')
    data=validate_data(json.loads(data_path.read_text(encoding='utf-8')),kind);spec=layout(kind);fonts=preflight(data,kind,spec)
    spec=metadata_row_plan(data,title_row_plan(data,identity_row_plan(data,spec,fonts),fonts),fonts)
    # Use actual TeX font shaping and region boxes before publishing any fixed-height DOCX.
    from .assessment_latex import render_bundle
    with tempfile.TemporaryDirectory(prefix='assessment-preflight-') as tmp:
        render_bundle(kind,data_path,Path(tmp),overwrite=True,compile_pdf=True)
    doc=Document(template);table=doc.tables[0];m=data['metadata'];s=data['sections']
    for sec in doc.sections:
        sec.page_width=Mm(210);sec.page_height=Mm(297)
    # Normalize family aliases while retaining the source package's fixed labels/grid.
    for part in (doc._element,doc.styles.element):
        for f in part.xpath('.//w:rFonts'):
            for a,v in list(f.attrib.items()):
                if '黑体' in v or v=='SimHei':f.set(a,'黑体')
    for row,h in zip(table.rows,spec['table']['rows_mm']):
        row.height=Mm(h);row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
        trpr=row._tr.get_or_add_trPr()
        if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
    if kind!='defense-record':
        college=doc.paragraphs[2];college.clear();paragraph(college,size=16);college.paragraph_format.left_indent=Mm(spec['word']['prelude']['college_left_indent_mm']);college.paragraph_format.space_before=Mm(spec['word']['prelude']['college_space_before_mm']);college.paragraph_format.space_after=Mm(spec['word']['prelude']['college_space_after_mm']);college.paragraph_format.tab_stops.add_tab_stop(Mm(spec['word']['prelude']['college_tab_mm']))
        college_label=plain('学院：');major_label=plain('\t专业：')
        for key,runs in (('college',college_label),('major',major_label)):
            if next(r for r in spec['regions'] if r['key']==key)['label_font'].endswith('Bold'):
                for run in runs:run['bold']=True
        append(college,college_label,16,'黑体');append(college,plain(m['college']),16);append(college,major_label,16,'黑体');append(college,plain(m['major']),16)
        set_cell(table.cell(0,0),'论文（设计）题目：',m['title']);set_cell(table.cell(1,0),'姓名：',plain(m['student_name']));set_cell(table.cell(1,1),'学号：',plain(m['student_id']))
        if kind in ('advisor-review','reviewer-review'):
            field='advisor' if kind=='advisor-review' else 'reviewer';set_cell(table.cell(2,0),'指导教师姓名：' if field=='advisor' else '评阅教师姓名：',plain(m[field]));set_cell(table.cell(2,1),'职称：',plain(m[field+'_title']));idx=3;signature='指导教师签名：' if field=='advisor' else '签名：'
        else:idx=3;signature='组长签名：'
        cell=table.cell(idx,0);intro=deepcopy(cell.paragraphs[0]._p);content=s['comments'];guide=None
        if has_guidance(data,'comments'):guide=('comments',s['writing_guidance']['comments'])
        elif s.get('editable_headings',False):
            matched,content=split_writing_prompt(content,spec,'prompt')
            if not matched:intro=None
        section_cell(cell,intro,content,spec,idx,signature,s['date'],guide)
        if kind=='advisor-review':
            c=table.cell(4,0);intro=deepcopy(c.paragraphs[0]._p);content=s['division_of_work'];guide=None
            if has_guidance(data,'division_of_work'):guide=('division_of_work',s['writing_guidance']['division_of_work'])
            elif s.get('editable_headings',False):
                matched,content=split_writing_prompt(content,spec,'division-label')
                if not matched:intro=None
            section_cell(c,intro,content,spec,4,guidance=guide)
        if kind=='grade-assessment':
            set_cell(table.cell(2,0),'答辩小组成员签名：',[]);table.cell(2,0).vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell=table.cell(4,0);clear_cell(cell);p=cell.paragraphs[0];grade=next(r for r in spec['regions'] if r['key']=='final_grade')
            paragraph(p,size=grade['size_pt']);explicit_line_spacing(p,grade['line_spacing_pt'])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in region_runs(grade,data):append(p,[run],run['size_pt'],'黑体' if run['role']=='SimHei' else '宋体')
        else:set_cell(table.cell(len(table.rows)-1,0),'建议成绩：',plain(s['suggested_grade']))
        minimal_terminal_paragraph(doc,table,spec)
    else:
        # Original title value is merged grid 1–2; source labels remain unchanged.
        for row,col,key in ((0,1,'title'),(0,4,'advisor'),(0,6,'advisor_title'),(1,1,'student_name'),(1,3,'major'),(1,5,'student_id')):set_cell(table.cell(row,col),'',m[key] if key=='title' else plain(m[key]),size=12,center=True)
        time=doc.paragraphs[3];time.clear();paragraph(time,size=12);time.paragraph_format.space_after=Mm(spec['word']['prelude']['time_space_after_mm']);time.paragraph_format.tab_stops.add_tab_stop(Mm(spec['word']['prelude']['time_tab_mm']))
        append(time,plain('答辩时间：'),12,'黑体');append(time,plain(m['defense_time']),12)
        append(time,plain('\t地点：'),12,'黑体');append(time,plain(m['defense_location']),12)
        cell=table.cell(2,0);intro=deepcopy(cell.paragraphs[0]._p);content=questions_content(data);intro_kept=True;guide=None
        if has_guidance(data,'questions'):
            guide=('questions',s['writing_guidance']['questions']);intro_kept=bool(guide[1]);content=s['comments']+content
        elif s.get('editable_headings',False):
            intro_kept,editable_intro=split_writing_prompt(s['comments'],spec,'questions-prompt')
            if not intro_kept:intro=None
            content=editable_intro+content
        section_cell(cell,intro,content,spec,2,guidance=guide)
        p=cell.add_paragraph();paragraph(p);measure=Measure(resolved_fonts(spec));prompt_region=next(r for r in spec['regions'] if r['key']=='questions-prompt');member_region=next(r for r in spec['regions'] if r['key']=='members-label')
        header_height=guidance_height(v2_heading(prompt_region,'questions'),guidance_runs(v2_heading(prompt_region,'questions'),data),measure) if guide is not None else prompt_region['height_mm']
        used=(header_height if intro_kept else 0 if content else 25.4/72)+measure.blocks_height(content,prompt_region['width_mm'],spec);row_top=spec['table']['top_mm']+sum(spec['table']['rows_mm'][:2]);p.paragraph_format.space_before=Mm(max(0,member_region['y_mm']-row_top-used-spec['word']['members_paragraph_vertical_adjust_mm']));append(p,plain(member_region['value']))
        p=cell.add_paragraph();paragraph(p);append(p,plain('、'.join(s['members'])))
        footer=next(p for p in doc.paragraphs if '记录员：' in p.text)
        fixed_footer(doc,table,footer,spec,m,s['date'])
    if spec.get('title_row'):
        seen=set()
        for cell in table.rows[0].cells:
            if cell._tc in seen:continue
            seen.add(cell._tc);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                explicit_line_spacing(p,spec['title_row']['line_step_pt'])
                # Align glyphs within each line as well as the complete cell.
                ppr=p._p.get_or_add_pPr();alignment=ppr.find(qn('w:textAlignment'))
                if alignment is None:alignment=OxmlElement('w:textAlignment');ppr.append(alignment)
                alignment.set(qn('w:val'),'center')
    doc.core_properties.author='SZTU Thesis Template';doc.core_properties.last_modified_by='SZTU Thesis Template';doc.core_properties.title=spec['title']
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output);canonicalize_package_fonts(output)
    reopened=Document(output)
    if len(reopened.tables)!=1 or len(reopened.tables[0].rows)!=len(spec['word']['row_heights_twips']):raise RuntimeError('DOCX table structure failed')
    return output

def word_cli(kind,base):
    p=argparse.ArgumentParser();p.add_argument('--template',type=Path,default=base/'word/official-template.docx');p.add_argument('--data',type=Path,required=True);p.add_argument('--output',type=Path,required=True);p.add_argument('--overwrite',action='store_true');a=p.parse_args()
    try:result=render(kind,a.template,a.data,a.output,overwrite=a.overwrite)
    except (DataError,ValueError,RuntimeError,OSError) as e:print('error: '+str(e),file=sys.stderr);return 2
    print(result);return 0
