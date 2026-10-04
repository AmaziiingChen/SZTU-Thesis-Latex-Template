#!/usr/bin/env python3
"""Contract, fixed-region and optional end-to-end tests for four assessment forms.
Run with the Codex bundled Python. --render emits only synthetic QA artifacts.
"""
from __future__ import annotations
import argparse,copy,json,math,re,subprocess,sys
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

TEMPLATES=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(TEMPLATES))
from common.python.assessment_forms import (KINDS,DataError,validate_data,layout,preflight,final_grade,format_date,resolved_fonts,Measure,plain,writing_prompt,writing_plan,split_writing_prompt,iter_paragraphs,runs_text,questions_content,has_guidance,title_row_plan,metadata_row_plan,region_runs)
from common.python.assessment_latex import render_bundle
from common.python.assessment_word import render as render_word

PASS_FIXTURES=('blank','normal','rich-text-list','capacity-boundary','layout-stress')
EDITABLE_FIXTURES=('editable-default','editable-deleted','editable-modified','editable-styled','editable-empty','editable-heading-only-deleted')
GUIDANCE_FIXTURES=('guidance-default','guidance-deleted','guidance-modified','guidance-styled','guidance-empty')

def fixture(kind,name):return json.loads((TEMPLATES/kind/'fixtures'/f'{name}.json').read_text())

def rejected(raw,kind,token=None):
    try:preflight(validate_data(raw,kind),kind)
    except (DataError,ValueError) as e:
        if token:assert token in str(e),str(e)
    else:raise AssertionError(f'{kind} accepted invalid or over-capacity content')

def contract_tests():
    count=0
    for kind in KINDS:
        for name in PASS_FIXTURES:
            preflight(validate_data(fixture(kind,name),kind),kind);count+=1
        for name in ('overflow','overflow-body'):
            rejected(fixture(kind,name),kind,'SINGLE_PAGE_OVERFLOW');count+=1
        raw=fixture(kind,'normal');raw['metadata']['student_id']='1'*1000;rejected(raw,kind,'SINGLE_PAGE_OVERFLOW');count+=1
        raw=fixture(kind,'normal');raw['metadata']['advisor' if kind=='defense-record' else 'college']='长'*1000;rejected(raw,kind,'SINGLE_PAGE_OVERFLOW');count+=1
        for value in ('-1','NaN','Infinity','1e2','89.5oops'):
            raw=fixture(kind,'normal');raw['sections']['final_grade']=value;rejected(raw,kind);count+=1
        raw=fixture(kind,'normal');raw['sections']['date']='2026-02-31';rejected(raw,kind);count+=1
        raw=fixture(kind,'normal');raw['sections']['comments']=[{'type':'image','path':'ignored.png','alt':'image'}];rejected(raw,kind);count+=1
        raw=fixture(kind,'normal');raw['sections']['final_grade']='9'*40
        if kind=='grade-assessment':rejected(raw,kind,'SINGLE_PAGE_OVERFLOW')
        else:preflight(validate_data(raw,kind),kind)
        count+=1
        # Document sizes are verified against the shared font policy, not copied values.
        sp=layout(kind)
        assert len(sp['table']['rows_mm'])==len(sp['word']['row_heights_twips'])
        assert sp['paragraphs']['list_max_depth']==4
        assert all(v['size_pt']>0 for v in sp['typography'].values())
    assert final_grade('89.5')=='90' and final_grade('89.49')=='89' and final_grade('0.5')=='1' and final_grade('')==''
    assert format_date('2026-06-01')=='2026年6月1日'
    preflight(validate_data(fixture('defense-record','footer-single-line'),'defense-record'),'defense-record')
    print(f'PASS contract/capacity tests: {count+6}',flush=True)

def identity_contract_tests():
    for kind in ('grade-assessment','advisor-review','reviewer-review'):
        sp=layout(kind);fonts=resolved_fonts(sp)
        baseline=copy.deepcopy(sp)
        for name in ('blank','normal'):
            plan,_=writing_plan(validate_data(fixture(kind,name),kind),kind,sp,fonts)
            assert plan==sp,('short identity moved',name)
        raw=fixture(kind,'long-major');data=validate_data(raw,kind);preflight(data,kind)
        plan,_=writing_plan(data,kind,sp,fonts)
        original={r['key']:r for r in sp['regions']};adjusted={r['key']:r for r in plan['regions']}
        assert adjusted['major']['x_mm']<original['major']['x_mm']
        assert abs(adjusted['major']['x_mm']+adjusted['major']['width_mm']-178.25)<.001
        delta=adjusted['major']['x_mm']-original['major']['x_mm']
        assert abs(plan['word']['prelude']['college_tab_mm']-sp['word']['prelude']['college_tab_mm']-delta)<.001
        assert plan['table']==sp['table'] and plan['page']==sp['page'] and sp==baseline
        assert all(adjusted[k]==r for k,r in original.items() if k not in ('college','major'))
        for k in ('college','major'):
            assert adjusted[k]['size_pt']==16 and adjusted[k]['height_mm']==7.6
        long_college=fixture(kind,'normal');long_college['metadata']['college']='合成学院'*3
        preflight(validate_data(long_college,kind),kind)
        for value in ('过长专业名称'*10,'合成\n专业'):
            raw=fixture(kind,'normal');raw['metadata']['major']=value;rejected(raw,kind,'metadata.major')
        print('PASS identity row: long names, source positions, shared tab, fixed geometry and overflow',flush=True)

def identity_render_tests(output,kind='grade-assessment'):
    name='long-major';source=TEMPLATES/kind/'fixtures'/f'{name}.json'
    bundle=output/f'{kind}-{name}'
    render_bundle(kind,source,bundle,overwrite=True,compile_pdf=True);inspect_pdf(bundle/'main.pdf',kind,name)
    bbox=ET.fromstring(subprocess.check_output(['pdftotext','-bbox',str(bundle/'main.pdf'),'-']))
    words=[w for w in bbox.iter() if w.tag.endswith('word')]
    source_data=fixture(kind,name)
    header=[]
    for key,label in (('college','学院：'),('major','专业：')):
        matches=[w for w in words if w.text==label+source_data['metadata'][key]]
        assert len(matches)==1,('identity must occupy one complete line',key)
        header.append({k:float(v)*25.4/72 for k,v in matches[0].attrib.items()})
    assert header[1]['xMax']<=178.25 and header[1]['xMin']-header[0]['xMax']>=5.645
    assert abs(header[0]['yMin']-header[1]['yMin'])<.1
    word=output/f'{kind}-{name}.docx';render_word(kind,TEMPLATES/kind/'word/official-template.docx',source,word,overwrite=True)
    inspect_docx(word,kind,fixture(kind,name))
    data=validate_data(fixture(kind,name),kind);sp=layout(kind);plan,_=writing_plan(data,kind,sp,resolved_fonts(sp))
    paragraph=Document(word).paragraphs[2]
    stops=list(paragraph.paragraph_format.tab_stops)
    assert len(stops)==1 and abs(stops[0].position.mm-plan['word']['prelude']['college_tab_mm'])<.03
    assert data['metadata']['major'] in paragraph.text
    assert all(r.font.size.pt==16 for r in paragraph.runs if r.text)
    print('PASS identity PDF and DOCX structure; Word/WPS visual acceptance remains separate',flush=True)

def editable_contract_tests():
    for kind in KINDS:
        sp=layout(kind);fonts=resolved_fonts(sp);key='questions-prompt' if kind=='defense-record' else 'prompt'
        for name in EDITABLE_FIXTURES:
            d=validate_data(fixture(kind,name),kind);preflight(d,kind)
            plan,overrides=writing_plan(d,kind,sp,fonts)
            keys=[r['key'] for r in plan['regions']]
            assert (key in keys)==(name=='editable-default')
            if kind=='advisor-review':assert ('division-label' in keys)==(name=='editable-default')
            for fixed in ('signature','date','members-label','members'):
                original=next((r for r in sp['regions'] if r['key']==fixed),None)
                assert next((r for r in plan['regions'] if r['key']==fixed),None)==original
            if kind!='defense-record' and name!='editable-styled':
                assert not any(writing_prompt(sp,key) in ''.join(r['text'] for r in b.get('runs',[])) for b in overrides['comments'])
            legacy=fixture(kind,'normal');legacy['sections']['editable_headings']=False
            legacy=validate_data(legacy,kind);assert writing_plan(legacy,kind,sp,fonts)==(metadata_row_plan(legacy,title_row_plan(legacy,sp,fonts),fonts),{})
        for bad in ('true',1,None):
            d=fixture(kind,'normal');d['sections']['editable_headings']=bad;rejected(d,kind,'must be a boolean')
        d=fixture(kind,'editable-modified');d['sections']['comments']=['长'*10000];rejected(d,kind,'SINGLE_PAGE_OVERFLOW')
    print('PASS editable-heading contracts: 24 canonical/deleted/modified/styled/empty/heading-only-deleted fixtures, fixed signature/member boxes, legacy compatibility, boolean validation and overflow',flush=True)

def guidance_contract_tests():
    for kind in KINDS:
        sp=layout(kind);fonts=resolved_fonts(sp);field='questions' if kind=='defense-record' else 'comments'
        for name in GUIDANCE_FIXTURES:
            d=validate_data(fixture(kind,name),kind);preflight(d,kind);plan,overrides=writing_plan(d,kind,sp,fonts)
            key='questions-prompt' if kind=='defense-record' else 'prompt';heading=next(r for r in plan['regions'] if r['key']==key)
            assert heading['source']=='guidance'
            source_heading=next(r for r in sp['regions'] if r['key']==key)
            if name in ('guidance-default','guidance-deleted','guidance-empty'):
                assert heading.get('punct_style')==source_heading.get('punct_style')
                assert 'punct_optimize_margin' not in heading
            elif 'punctuation' in source_heading['guidance_style']:
                assert heading['punct_optimize_margin'] is False
            if kind!='defense-record':
                assert heading['label_font']=='heiti' and heading['label_size_pt']==14 and heading['first_line_indent_em']==0
                assert heading['guidance_style']['fixed_label']['bold'] is False
                assert overrides['comments']==d['sections']['comments']
            for fixed in ('signature','date','members-label','members'):
                assert next((r for r in plan['regions'] if r['key']==fixed),None)==next((r for r in sp['regions'] if r['key']==fixed),None)
        for empty in (None,'',{'type':'paragraph','runs':[]}):
            d=fixture(kind,'guidance-default');d['sections']['writing_guidance'][field]=empty
            normalized=validate_data(d,kind);assert normalized['sections']['writing_guidance'][field]==[];preflight(normalized,kind)
        d=fixture(kind,'guidance-default');d['sections']['writing_guidance'][field]='长'*10000;rejected(d,kind,'SINGLE_PAGE_OVERFLOW')
        for bad in (1,False,[],{'type':'image','path':'ignored.png'}):
            d=fixture(kind,'guidance-default');d['sections']['writing_guidance'][field]=bad;rejected(d,kind)
        d=fixture(kind,'guidance-default');d['sections']['writing_guidance']['unknown']=None;rejected(d,kind)
        # Map/key absence preserves old callers; explicit keys override v1.
        legacy=fixture(kind,'editable-modified');legacy['sections']['writing_guidance']={}
        normalized=validate_data(legacy,kind);without=copy.deepcopy(normalized);without['sections'].pop('writing_guidance')
        assert writing_plan(normalized,kind,sp,fonts)==writing_plan(without,kind,sp,fonts)
        if kind!='defense-record':
            d=fixture(kind,'guidance-deleted');d['sections']['comments']=[writing_prompt(sp,'prompt')]
            d=validate_data(d,kind);_,override=writing_plan(d,kind,sp,fonts);assert override['comments']==d['sections']['comments']
        boundary=fixture(kind,'capacity-boundary');boundary['sections']['writing_guidance']=fixture(kind,'guidance-default')['sections']['writing_guidance']
        if kind=='defense-record':boundary['sections']['comments']=[]
        preflight(validate_data(boundary,kind),kind)
    print('PASS v2 guidance contracts: 20 states, explicit empty forms, fixed titles/regions, no body guessing, legacy/v1 and capacity compatibility',flush=True)


def inspect_docx(path,kind,data):
    doc=Document(path);sp=layout(kind);sp=metadata_row_plan(validate_data(data,kind),title_row_plan(validate_data(data,kind),sp,resolved_fonts(sp)),resolved_fonts(sp));assert len(doc.tables)==1
    table=doc.tables[0]
    assert [c.w.twips for c in table._tbl.tblGrid.gridCol_lst]==sp['word']['table_grid_twips']
    assert len(table.rows)==len(sp['table']['rows_mm'])
    for row,expected in zip(table.rows,sp['word']['row_heights_twips']):
        assert abs(row.height.twips-expected)<=1
        assert row._tr.trPr.trHeight.get(qn('w:hRule'))=='exact'
        assert row._tr.trPr.find(qn('w:cantSplit')) is not None
    sec=doc.sections[0]
    assert abs(sec.page_width.mm-210)<.02 and abs(sec.page_height.mm-297)<.02
    with ZipFile(path) as archive:
        for name in ('word/document.xml','word/styles.xml','word/fontTable.xml'):
            root=ET.fromstring(archive.read(name))
            for el in root.iter():
                if el.tag.endswith('rFonts'):
                    for key in ('ascii','hAnsi','eastAsia','cs'):
                        value=el.get(qn('w:'+key))
                        assert value in (None,'宋体','黑体','Times New Roman'),(name,value)
                if el.tag.endswith('font') and name=='word/fontTable.xml':assert el.get(qn('w:name')) in ('宋体','黑体','Times New Roman')
    for p in doc.paragraphs[:2]:
        for f in p._p.xpath('.//w:rFonts'):
            assert all(f.get(qn('w:'+key))=='黑体' for key in ('ascii','hAnsi','eastAsia','cs'))
    if kind!='defense-record':
        p=next(p for p in doc._element.xpath('.//w:p') if '签名：' in ''.join(p.xpath('.//w:t/text()')) and not '成员' in ''.join(p.xpath('.//w:t/text()')))
        assert any(r.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val'))=='28' for r in p.findall(qn('w:r')) if r.find(qn('w:rPr')) is not None and r.find(qn('w:rPr')).find(qn('w:sz')) is not None)
    if kind=='grade-assessment':assert '周示例' not in table.cell(2,0).text
    identity_cells=[table.cell(0,1),table.cell(0,4),table.cell(0,6),table.cell(1,1),table.cell(1,3),table.cell(1,5)] if kind=='defense-record' else [table.cell(0,0),table.cell(1,0),table.cell(1,1)]
    assert all(c._tc.tcPr.vAlign.get(qn('w:val'))=='center' for c in identity_cells)
    if kind=='defense-record':
        assert table.cell(0,0).text=='题 目' and table.cell(0,3).text=='指导教师姓 名'
        time=next(p for p in doc.paragraphs if p.text.startswith('答辩时间：'))
        for r in time.runs:
            if '答辩时间：' in r.text or '地点：' in r.text:assert r._element.rPr.rFonts.get(qn('w:eastAsia'))=='黑体'
        prompt=table.cell(2,0).paragraphs[0]
        if not data['sections'].get('editable_headings') or split_writing_prompt(validate_data(data,kind)['sections']['comments'],sp,'questions-prompt')[0]:
            assert abs(prompt.paragraph_format.first_line_indent.pt-26.25)<.1
        # The .doc conversion's empty paragraphs and implicit line/character
        # grids formerly moved the date/signature line to page 2 in WPS.
        footer=next(p for p in doc.paragraphs if p.text.startswith('记录员：'))
        assert footer._p.getprevious() is table._tbl
        assert footer._p.getnext().tag==qn('w:sectPr')
        fs=sp['word']['footer'];ppr=footer._p.pPr
        assert abs(footer.paragraph_format.space_before.mm-fs['space_before_mm'])<.02
        assert abs(footer.paragraph_format.line_spacing.pt-14)<.01
        assert footer.paragraph_format.keep_with_next is False
        assert footer.paragraph_format.widow_control is False
        assert abs(footer.paragraph_format.right_indent.mm-fs['right_indent_mm'])<.02
        for tag in ('snapToGrid','autoSpaceDE','autoSpaceDN','adjustRightInd'):
            assert ppr.find(qn('w:'+tag)).get(qn('w:val'))=='false'
        ordered=['autoSpaceDE','autoSpaceDN','adjustRightInd','snapToGrid','spacing','ind']
        positions=[list(ppr).index(ppr.find(qn('w:'+tag))) for tag in ordered]
        assert positions==sorted(positions)
        assert len(ppr.tabs)==2
        assert '\n' not in footer.text and '\v' not in footer.text
        if data['sections']['date']:
            date=format_date(data['sections']['date']);assert footer.text.endswith(date)
            region=next(r for r in sp['regions'] if r['key']=='date')
            assert Measure(resolved_fonts(sp)).width(plain(date),region['size_pt'])*25.4/72<=region['width_mm']
        for side in ('top','bottom','left','right'):
            assert abs(getattr(sec,side+'_margin').mm-sp['page'][side+'_margin_mm'])<.02
    else:
        for p in doc.paragraphs:
            if p._p.getprevious() is not table._tbl:continue
            assert not p.text
            assert p.paragraph_format.line_spacing.pt==1
            assert p._p.pPr.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val'))=='2'
            assert p._p.pPr.find(qn('w:snapToGrid')).get(qn('w:val'))=='false'
    active_guidance='questions' if kind=='defense-record' else 'comments'
    if data['sections'].get('editable_headings') and active_guidance not in data['sections'].get('writing_guidance',{}):
        body=''.join(doc._element.xpath('.//w:t/text()'));key='questions-prompt' if kind=='defense-record' else 'prompt'
        normalized=validate_data(data,kind)
        active=[normalized['sections']['comments']]
        if kind=='advisor-review':active.append(normalized['sections']['division_of_work'])
        if kind=='defense-record':active.append(questions_content(normalized))
        for blocks in active:
            for _,runs,_ in iter_paragraphs(blocks):
                assert runs_text(runs) in body,('editable body lost from Word',kind,runs_text(runs))
        prefix=next(r for r in sp['regions'] if r['key']==key);prefix=prefix['label'] or prefix['value'][:20]
        source=data['sections']['comments'];expected=bool(source and (source[0]==writing_prompt(sp,key) or isinstance(source[0],dict)))
        assert body.count(prefix)==int(expected),(kind,prefix,body)
        if kind=='advisor-review':
            source=data['sections']['division_of_work'];expected=bool(source and (source[0]==writing_prompt(sp,'division-label') or isinstance(source[0],dict)))
            assert body.count('毕业论文（设计）的分工情况：')==int(expected)
        for field in ('comments','division_of_work'):
            for item in data['sections'][field]:
                if isinstance(item,str) and item.startswith('自定义'):assert item in body
        if isinstance(data['sections']['comments'][0] if data['sections']['comments'] else None,dict):
            styled=next(p for p in doc._element.xpath('.//w:p') if ''.join(p.xpath('.//w:t/text()')).startswith('指导教师评语' if kind=='advisor-review' else '评阅教师评语' if kind=='reviewer-review' else '答辩小组意见' if kind=='grade-assessment' else '答辩小组（学院答辩委员会）对学生'))
            rs=styled.findall(qn('w:r'));assert rs[0].find(qn('w:rPr')).find(qn('w:b')).get(qn('w:val')) in (None,'1')
            assert rs[1].find(qn('w:rPr')).find(qn('w:i')).get(qn('w:val')) in (None,'1')
            assert rs[2].find(qn('w:rPr')).find(qn('w:vertAlign')).get(qn('w:val'))=='superscript'
            assert all(r.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val'))=='21' for r in rs)
        if not data['sections']['comments']:
            cell=table.cell(2,0) if kind=='defense-record' else table.cell(3,0)
            if kind!='defense-record' or not data['sections']['questions']:
                p=cell.paragraphs[0];assert not p.text and p.paragraph_format.line_spacing.pt==1
                assert p._p.pPr.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val'))=='2'


def pdf_words(pdf):
    root=ET.fromstring(subprocess.check_output(['pdftotext','-bbox',str(pdf),'-']))
    pages=[el for el in root.iter() if el.tag.endswith('page')]
    words=[el for el in root.iter() if el.tag.endswith('word')]
    return pages,words


def inspect_pdf(pdf,kind,name):
    sp=layout(kind);source_basic_height=sum(sp['table']['rows_mm'][:2])
    data=validate_data(fixture(kind,name),kind);fonts=resolved_fonts(sp)
    sp=metadata_row_plan(data,title_row_plan(data,sp,fonts),fonts)
    title_growth=sum(sp['table']['rows_mm'][:2])-source_basic_height
    pages,words=pdf_words(pdf);assert len(pages)==1
    assert abs(float(pages[0].attrib['width'])*25.4/72-210)<.02
    assert abs(float(pages[0].attrib['height'])*25.4/72-297)<.02
    assert any(re.search('[\u4e00-\u9fff]',w.text or '') for w in words)
    # Independent official-source glyph bounds: raw measurements, not values
    # read back from layout.json. Local title spacing must match Word's heading.
    heading=[w for w in words if 27<float(w.attrib['yMin'])*25.4/72<35]
    expected_x=(54.24,155.84) if kind=='defense-record' else (57.415,152.665)
    actual=(min(float(w.attrib['xMin'])*25.4/72 for w in heading),max(float(w.attrib['xMax'])*25.4/72 for w in heading))
    assert all(abs(a-b)<.5 for a,b in zip(actual,expected_x)),('official heading bounds',kind,actual,expected_x)
    assert all(abs(float(w.attrib['yMin'])*25.4/72-27.883)<.25 and abs(float(w.attrib['yMax'])*25.4/72-34.227)<.25 for w in heading)
    second=[w for w in words if 38<float(w.attrib['yMin'])*25.4/72<46]
    assert second and all(abs(float(w.attrib['yMin'])*25.4/72-38.890)<.25 for w in second)
    prompt_expectations={
        'advisor-review': [('（从选题价值',96.390,'创新'),('性、撰写水平',107.397,'评述）')],
        'reviewer-review': [('（从选题价值',96.410,'撰写的'),('水平与规范性',107.417,'评述）')],
        'grade-assessment': [('（从选题价值',117.397,'严密性'),('和论文撰写',125.158,'评价）')],
        'defense-record': [('答辩小组（学院答辩委员会）对学生',94.591,'回答'),('情况：',100.094,'情况：')],
    }
    editable=name in EDITABLE_FIXTURES;guidance=name in GUIDANCE_FIXTURES
    for prefix,y,suffix in (prompt_expectations[kind] if not editable and not guidance or name in ('editable-default','guidance-default') else []):
        y+=title_growth
        line=[w for w in words if (w.text or '').startswith(prefix)]
        assert len(line)==1 and (line[0].text or '').endswith(suffix),(kind,prefix,[w.text for w in line])
        assert abs(float(line[0].attrib['yMin'])*25.4/72-y)<.25,(kind,prefix,line[0].attrib,y)
    result=subprocess.run([sys.executable,str(TEMPLATES.parent/'scripts/validate_cjk_render.py'),str(pdf)],text=True,capture_output=True)
    assert result.returncode==0,result.stdout+result.stderr
    text=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True)
    if editable:
        source_key='questions-prompt' if kind=='defense-record' else 'prompt'
        source_label=next(r for r in sp['regions'] if r['key']==source_key)
        unique_prefix=source_label['label'] or source_label['value'][:20]
        flat=re.sub(r'\s+','',text)
        if name!='editable-styled':assert flat.count(unique_prefix)==(1 if name=='editable-default' else 0)
        if name=='editable-modified':assert '自定义可编辑评价说明。' in flat
        if kind=='advisor-review':
            if name!='editable-styled':assert flat.count('毕业论文（设计）的分工情况：')==(1 if name=='editable-default' else 0)
            if name=='editable-modified':assert '自定义可编辑分工说明。' in flat
        if name=='editable-heading-only-deleted':
            assert ('采用实验分析与理论计算相结合的方法。' if kind=='defense-record' else '该论文选题合理，研究方法清晰，材料完整，论证充分。') in flat
            if kind=='advisor-review':assert '独立完成研究与撰写。' in flat
    if guidance:
        flat=re.sub(r'\s+','',text);data=fixture(kind,name)
        if kind!='defense-record':
            label=next(r for r in sp['regions'] if r['key']=='prompt')['label']
            assert flat.count(label)==1
            if name in ('guidance-deleted','guidance-empty'):
                fixed=label+next(r for r in sp['regions'] if r['key']=='prompt').get('suffix',{}).get('text','')
                assert fixed in flat and '从选题价值' not in flat
            label_words=[w for w in words if (w.text or '').startswith(label)]
            assert label_words and abs(float(label_words[0].attrib['xMin'])*25.4/72-31.708)<.3
        elif name in ('guidance-deleted','guidance-empty'):assert '对学生毕业设计' not in flat
        if name!='guidance-empty':
            target='补充答辩记录说明。' if kind=='defense-record' else '该论文选题合理，研究方法清晰，材料完整，论证充分。'
            assert target in flat
            body=[w for w in words if (w.text or '').startswith(target[:6])]
            assert body and abs(float(body[0].attrib['xMin'])*25.4/72-(31.708+21*25.4/72))<.3
    if kind=='grade-assessment' and name=='normal':assert '90' in text and '89.5' not in text
    if name=='blank':assert '签名' in text or kind=='defense-record'
    if name=='normal':
        assert '2026' in text and '年' in text
        title=next(r for r in sp['regions'] if r['key']=='title')
        # Actual first-row text stays clear of horizontal rules after vertical centering.
        title_words=[w for w in words if float(w.attrib['yMin'])*25.4/72>sp['table']['top_mm'] and float(w.attrib['yMax'])*25.4/72<sp['table']['top_mm']+sp['table']['rows_mm'][0]]
        assert title_words
        assert min(float(w.attrib['yMin'])*25.4/72 for w in title_words)>sp['table']['top_mm']+.5
        assert max(float(w.attrib['yMax'])*25.4/72 for w in title_words)<sp['table']['top_mm']+sp['table']['rows_mm'][0]-.5
    fonts=subprocess.check_output(['pdffonts',str(pdf)],text=True)
    assert 'SimHei' in fonts and 'SimSun' in fonts and 'TimesNewRoman' in fonts,fonts
    # Every actual text box remains on the page and outside the form's border,
    # with positive padding for fields and all four corner ink joins.
    left=sp['table']['left_mm'];right=left+sp['table']['width_mm'];top=sp['table']['top_mm'];bottom=top+sum(sp['table']['rows_mm'])
    bounded_guides=[]
    if name in ('guidance-modified','guidance-styled'):
        guide_data=validate_data(fixture(kind,name),kind);plan,_=writing_plan(guide_data,kind,sp,resolved_fonts(sp))
        bounded_guides=[r for r in plan['regions'] if r['source']=='guidance' and r.get('punct_optimize_margin') is False]
    for word in words:
        x0=float(word.attrib['xMin'])*25.4/72;x1=float(word.attrib['xMax'])*25.4/72
        y0=float(word.attrib['yMin'])*25.4/72;y1=float(word.attrib['yMax'])*25.4/72
        assert 0<=x0<x1<210 and 0<=y0<y1<297
        if top+.5<y0 and y1<bottom-.5:
            # Modified/marked guides preserve source metrics. Their physical
            # glyphs must clear the actual border stroke; the historical .8mm
            # generic field margin was never an official guide measurement.
            in_guide=any(r['y_mm']-.5<=y0 and y1<=r['y_mm']+r['height_mm']+.5 for r in bounded_guides)
            margin=sp['table']['border_pt']*25.4/72/2 if in_guide else .8
            assert x0>left+margin and x1<right-margin,(word.text,x0,x1,margin)
    png=pdf.parent/'page.png';subprocess.run(['pdftoppm','-png','-r','150','-singlefile',str(pdf),str(pdf.parent/'page')],check=True,capture_output=True)
    im=Image.open(png).convert('L');scale=150/25.4
    for x,y in ((left,top),(right,top),(left,bottom),(right,bottom)):
        px,py=round(x*scale),round(y*scale);assert im.crop((px-2,py-2,px+3,py+3)).getextrema()[0]<100
    # Measure the rendered full-width rules independently of spec geometry.
    # User acceptance targets below are deliberately fixed literals: editing
    # layout.json must not make a shrunken form pass its own expected values.
    ix0=round((left+.4)*scale);ix1=round((right-.4)*scale)
    line_rows=[]
    for iy in range(round(55*scale),round(274*scale)):
        pixels=im.crop((ix0,iy,ix1,iy+1)).getdata()
        if sum(v<160 for v in pixels)/max(1,ix1-ix0)>.9:line_rows.append(iy)
    groups=[]
    for iy in line_rows:
        if not groups or iy>groups[-1][-1]+1:groups.append([iy])
        else:groups[-1].append(iy)
    rules=[sum(g)/len(g)/scale for g in groups]
    assert len(rules)==len(sp['table']['rows_mm'])+1,(kind,rules)
    actual_top,actual_bottom=rules[0],rules[-1]
    if kind!='defense-record':
        assert 270<=actual_bottom<=272,('full-page handwriting goal',kind,actual_bottom)
        for label in ('学院：','专业：'):
            line=[w for w in words if (w.text or '').startswith(label)]
            assert line,(kind,label)
            glyph_bottom=max(float(w.attrib['yMax'])*25.4/72 for w in line)
            assert 1.5<=actual_top-glyph_bottom<=3.5,('college-to-table glyph gap',kind,label,actual_top-glyph_bottom)
        signature=[w for w in words if '签名：' in (w.text or '') and '成员' not in (w.text or '')]
        assert signature
        signature_bottom=max(float(w.attrib['yMax'])*25.4/72 for w in signature)
        signature_row_end=rules[4]  # comment row ends at boundary index 4.
        assert 1<=signature_row_end-signature_bottom<=8,('signature anchored near comment bottom',kind,signature_row_end-signature_bottom)
    else:
        assert 257<=actual_bottom<=260,('defense page handwriting region',actual_bottom)
        footer=[w for w in words if '记录员：' in (w.text or '') or '组长：' in (w.text or '')]
        assert footer and 267<=max(float(w.attrib['yMax'])*25.4/72 for w in footer)<=272
    for word in words:
        x0=float(word.attrib['xMin'])*25.4/72;x1=float(word.attrib['xMax'])*25.4/72
        y0=float(word.attrib['yMin'])*25.4/72;y1=float(word.attrib['yMax'])*25.4/72
        if left<x0<x1<right:
            assert not any(y0+.1<rule<y1-.1 for rule in rules),('border crosses glyph',word.text,y0,y1,rules)


def render_tests(output):
    for kind in KINDS:
        for name in PASS_FIXTURES:
            source=TEMPLATES/kind/'fixtures'/f'{name}.json';bundle=output/f'{kind}-{name}'
            render_bundle(kind,source,bundle,overwrite=True,compile_pdf=True)
            inspect_pdf(bundle/'main.pdf',kind,name)
            word=output/f'{kind}-{name}.docx'
            render_word(kind,TEMPLATES/kind/'word/official-template.docx',source,word,overwrite=True)
            inspect_docx(word,kind,fixture(kind,name))
            print(f'PASS render/geometry/font/CJK: {kind}/{name}',flush=True)
    print('PASS 20 PDFs / 20 DOCX generated and structurally inspected; Word/WPS is a separate acceptance gate.',flush=True)
    kind='defense-record';name='footer-single-line'
    source=TEMPLATES/kind/'fixtures'/f'{name}.json';word=output/f'{kind}-{name}.docx'
    render_word(kind,TEMPLATES/kind/'word/official-template.docx',source,word,overwrite=True)
    inspect_docx(word,kind,fixture(kind,name))
    print('PASS Word footer single-line regression: no inherited grids or pagination spacer paragraphs.',flush=True)

def editable_render_tests(output):
    for kind in KINDS:
        for name in EDITABLE_FIXTURES:
            source=TEMPLATES/kind/'fixtures'/f'{name}.json';bundle=output/f'{kind}-{name}'
            render_bundle(kind,source,bundle,overwrite=True,compile_pdf=True);inspect_pdf(bundle/'main.pdf',kind,name)
            word=output/f'{kind}-{name}.docx';render_word(kind,TEMPLATES/kind/'word/official-template.docx',source,word,overwrite=True)
            inspect_docx(word,kind,fixture(kind,name))
            print(f'PASS editable PDF/Word: {kind}/{name}',flush=True)
    print('PASS 24 editable PDFs and 24 DOCX: real default text once, deletion, modification, styled runs, empty cells, body preserved after prompt deletion and fixed geometry.',flush=True)

def inspect_guidance_docx(path,kind,data):
    doc=Document(path);sp=layout(kind);d=validate_data(data,kind);body=''.join(doc._element.xpath('.//w:t/text()'))
    pairs=[('questions','questions-prompt',2)] if kind=='defense-record' else [('comments','prompt',3)]+([('division_of_work','division-label',4)] if kind=='advisor-review' else [])
    for field,key,row in pairs:
        source=next(r for r in sp['regions'] if r['key']==key);notes=d['sections']['writing_guidance'][field]
        fixed=source['label'];suffix=source.get('suffix',{}).get('text','');value=fixed+runs_text(notes)+suffix
        candidates=[p for p in doc.tables[0].cell(row,0).paragraphs if p.text==value]
        if not value:assert not any(source['value'] in p.text for p in doc.tables[0].cell(row,0).paragraphs);continue
        assert len(candidates)==1,(kind,field,value,[p.text for p in doc.tables[0].cell(row,0).paragraphs])
        p=candidates[0];expected_indent=26.25 if kind=='defense-record' else 0
        assert abs(p.paragraph_format.first_line_indent.pt-expected_indent)<.01 and p.paragraph_format.left_indent.pt==0
        assert p.paragraph_format.line_spacing.pt==source['line_spacing_pt']
        runs=p.runs;offset=0
        if fixed:
            style=source['guidance_style']['fixed_label'];r=runs[0];assert r.text==fixed and bool(r.bold)==style['bold'] and r.font.size.pt==style['size_pt'] and r._element.rPr.rFonts.get(qn('w:eastAsia'))==('黑体' if style['font'].startswith('heiti') else '宋体');offset=1
        for expected,r in zip(notes,runs[offset:offset+len(notes)]):
            assert r.text==expected['text'] and r.font.size.pt==10.5 and r._element.rPr.rFonts.get(qn('w:eastAsia'))=='宋体'
            assert bool(r.bold)==expected['bold'] and bool(r.italic)==expected['italic']
            expected_script={'super':'superscript','sub':'subscript','normal':None}[expected['script']]
            el=r._element.rPr.find(qn('w:vertAlign'));assert (el.get(qn('w:val')) if el is not None else None)==expected_script
        if suffix:
            style=source['guidance_style']['fixed_suffix'];r=runs[-1];assert r.text==suffix and bool(r.bold)==style['bold'] and r.font.size.pt==style['size_pt'] and r._element.rPr.rFonts.get(qn('w:eastAsia'))==('黑体' if style['font'].startswith('heiti') else '宋体')
    active=[d['sections']['comments']]+([d['sections']['division_of_work']] if kind=='advisor-review' else [])
    if kind=='defense-record':active.append(questions_content(d))
    paragraphs=doc._element.xpath('.//w:p')
    for blocks in active:
        for _,runs,depth in iter_paragraphs(blocks):
            text=runs_text(runs);assert text in body
            if depth==0:
                found=[p for p in paragraphs if ''.join(p.xpath('.//w:t/text()'))==text];assert found
                ind=found[-1].pPr.find(qn('w:ind'));assert ind.get(qn('w:firstLine'))=='420'
                for r in found[-1].findall(qn('w:r')):
                    assert r.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val'))=='21'
                    assert r.find(qn('w:rPr')).find(qn('w:rFonts')).get(qn('w:eastAsia'))=='宋体'

def guidance_render_tests(output):
    for kind in KINDS:
        for name in GUIDANCE_FIXTURES:
            source=TEMPLATES/kind/'fixtures'/f'{name}.json';bundle=output/f'{kind}-{name}'
            render_bundle(kind,source,bundle,overwrite=True,compile_pdf=True);inspect_pdf(bundle/'main.pdf',kind,name)
            word=output/f'{kind}-{name}.docx';render_word(kind,TEMPLATES/kind/'word/official-template.docx',source,word,overwrite=True)
            inspect_docx(word,kind,fixture(kind,name));inspect_guidance_docx(word,kind,fixture(kind,name))
            print(f'PASS v2 guidance PDF/Word: {kind}/{name}',flush=True)
    print('PASS 20 v2 PDFs / 20 DOCX: official fixed label weight, editable note styles/marks and independent body indents.',flush=True)


def grade_layout_contract_tests():
    kind='grade-assessment';sp=layout(kind);baseline=copy.deepcopy(sp);fonts=resolved_fonts(sp)
    data=validate_data(fixture(kind,'title-two-lines'),kind)
    plan,_=writing_plan(data,kind,sp,fonts);preflight(data,kind)
    assert abs(plan['table']['rows_mm'][0]-17.3594444)<.001
    assert abs(sum(plan['table']['rows_mm'])-sum(sp['table']['rows_mm']))<.001
    for key in ('signature','date','final_grade'):
        assert next(r for r in plan['regions'] if r['key']==key)==next(r for r in sp['regions'] if r['key']==key)
    assert title_row_plan(data,plan,fonts)==plan and sp==baseline
    for count in (3,4):
        raw=fixture(kind,'normal');raw['metadata']['title']='\n'.join(['合成题目']*count)
        plan,_=writing_plan(validate_data(raw,kind),kind,sp,fonts)
        assert abs(plan['table']['rows_mm'][0]-(6+(14+(count-1)*18.2)*25.4/72))<.001
    raw=fixture(kind,'normal');raw['metadata']['title']='过长合成题目'*200;rejected(raw,kind,'metadata.title')
    raw=fixture(kind,'capacity-boundary');raw['metadata']['title']=fixture(kind,'title-two-lines')['metadata']['title'];rejected(raw,kind,'sections.comments')
    for name in ('blank','normal'):
        data=validate_data(fixture(kind,name),kind);score=next(r for r in sp['regions'] if r['key']=='final_grade')
        runs=region_runs(score,data)
        assert ''.join(r['text'] for r in runs)=='毕业论文（设计）成绩（如有小数点，请四舍五入保留整数）：'+final_grade(data['sections']['final_grade'])
        assert [r['size_pt'] for r in runs[:3]]==[14,10.5,14]
        assert score['alignment']=='left' and score['vertical_alignment']=='center'
        assert Measure(fonts).lines(runs,score['width_mm'],14)==1
    print('PASS grade layout: unified title spacing, shared geometry, fixed signatures, single-line mixed-size score and overflow',flush=True)

def grade_layout_render_tests(output):
    kind='grade-assessment'
    for name in ('blank','normal','title-two-lines','layout-stress','rich-text-list','guidance-styled'):
        source=TEMPLATES/kind/'fixtures'/f'{name}.json';bundle=output/f'{kind}-{name}'
        render_bundle(kind,source,bundle,overwrite=True,compile_pdf=True);inspect_pdf(bundle/'main.pdf',kind,name)
        word=bundle/'result.docx';render_word(kind,TEMPLATES/kind/'word/official-template.docx',source,word,overwrite=True)
        inspect_docx(word,kind,fixture(kind,name));doc=Document(word)
        title=doc.tables[0].cell(0,0).paragraphs[0]
        assert abs(title.paragraph_format.line_spacing.pt-18.2)<.01
        assert title._p.pPr.find(qn('w:textAlignment')).get(qn('w:val'))=='center'
        assert title._p.pPr.find(qn('w:snapToGrid')).get(qn('w:val'))=='false'
        score=doc.tables[0].cell(4,0);assert len(score.paragraphs)==1
        p=score.paragraphs[0];assert p.alignment==0 and not p._p.xpath('.//w:br|.//w:tab')
        assert score._tc.tcPr.vAlign.get(qn('w:val'))=='center'
        assert [(r.text,r.font.size.pt) for r in p.runs[:3]]==[('毕业论文（设计）成绩',14),('（如有小数点，请四舍五入保留整数）',10.5),('：',14)]
        _,words=pdf_words(bundle/'main.pdf')
        footer=[w for w in words if float(w.attrib['yMin'])*25.4/72>258.47]
        assert len(footer)==3 and footer[0].text=='毕业论文（设计）成绩'
        assert abs(float(footer[0].attrib['yMin'])-float(footer[2].attrib['yMin']))<.01
        top=min(float(w.attrib['yMin'])*25.4/72 for w in footer);bottom=max(float(w.attrib['yMax'])*25.4/72 for w in footer)
        assert abs((top+bottom)/2-(258.47+271)/2)<.2
        assert abs(float(footer[0].attrib['xMin'])*25.4/72-31.708)<.1
        if name=='title-two-lines':
            lines=[w for w in words if 58.42<float(w.attrib['yMin'])*25.4/72<75.7795]
            assert len(lines)==2
            ys=[float(w.attrib['yMin'])*25.4/72 for w in lines]
            assert abs(ys[1]-ys[0]-6.4206)<.03
            assert abs(ys[0]-58.42-3)<.4
            assert abs(75.7795-float(lines[-1].attrib['yMax'])*25.4/72-3)<.4
        print(f'PASS grade layout PDF/Word: {name}',flush=True)

def defense_major_contract_tests():
    kind='defense-record';sp=layout(kind);original=copy.deepcopy(sp);fonts=resolved_fonts(sp)
    for name in ('blank','normal'):
        assert metadata_row_plan(validate_data(fixture(kind,name),kind),sp,fonts)==sp
    raw=fixture(kind,'normal');raw['metadata']['major']='合成专业测试名称'
    preflight(validate_data(raw,kind),kind)
    assert metadata_row_plan(validate_data(raw,kind),sp,fonts)==sp
    data=validate_data(fixture(kind,'long-major'),kind);preflight(data,kind)
    plan,_=writing_plan(data,kind,sp,fonts)
    assert abs(plan['table']['rows_mm'][1]-(11.994+13.2*25.4/72))<.001
    assert abs(sum(plan['table']['rows_mm'])-sum(sp['table']['rows_mm']))<.001
    assert plan['word']['table_grid_twips']==sp['word']['table_grid_twips']
    for key in ('members-label','members','recorder','leader','date'):
        assert next(r for r in plan['regions'] if r['key']==key)==next(r for r in sp['regions'] if r['key']==key)
    assert metadata_row_plan(data,plan,fonts)==plan and sp==original
    raw['metadata']['major']='过长合成专业'*100;rejected(raw,kind,'metadata.major')
    print('PASS defense metadata: 8/10-character major, source columns/type, shared row growth, fixed members/footer and overflow',flush=True)

def defense_major_render_tests(output,name='long-major'):
    kind='defense-record';source=TEMPLATES/kind/'fixtures'/f'{name}.json';bundle=output/f'{kind}-{name}'
    render_bundle(kind,source,bundle,overwrite=True,compile_pdf=True);inspect_pdf(bundle/'main.pdf',kind,name)
    sp=layout(kind);plan,_=writing_plan(validate_data(fixture(kind,name),kind),kind,sp,resolved_fonts(sp));row_start=plan['table']['top_mm']+plan['table']['rows_mm'][0];row_end=row_start+plan['table']['rows_mm'][1]
    _,words=pdf_words(bundle/'main.pdf')
    major=[w for w in words if 95<float(w.attrib['xMin'])*25.4/72<114 and row_start<float(w.attrib['yMin'])*25.4/72<row_end]
    assert len(major)==3 and ''.join(w.text for w in major)==fixture(kind,name)['metadata']['major']
    assert min(float(w.attrib['yMin'])*25.4/72 for w in major)>row_start+.8
    assert max(float(w.attrib['yMax'])*25.4/72 for w in major)<row_end-.8
    word=bundle/'result.docx';render_word(kind,TEMPLATES/kind/'word/official-template.docx',source,word,overwrite=True)
    inspect_docx(word,kind,fixture(kind,name))
    cell=Document(word).tables[0].cell(1,3)
    assert cell.text==fixture(kind,name)['metadata']['major']
    assert all(r.font.size.pt==12 for r in cell.paragraphs[0].runs)
    print('PASS defense long-major PDF and DOCX: three lines inside source column, no clipping, one page',flush=True)

def title_layout_contract_tests():
    for kind in KINDS:
        sp=layout(kind);fonts=resolved_fonts(sp);original=copy.deepcopy(sp)
        raw=fixture(kind,'normal');raw['metadata']['title']='合成题目'
        plan=title_row_plan(validate_data(raw,kind),sp,fonts)
        assert plan==sp
        for name in ('title-layout-long','title-layout-rich','title-layout-breaks'):
            data=validate_data(fixture(kind,name),kind);preflight(data,kind)
            plan,_=writing_plan(data,kind,sp,fonts)
            assert abs(sum(plan['table']['rows_mm'])-sum(sp['table']['rows_mm']))<.001
            assert plan['table']['width_mm']==sp['table']['width_mm']
            title=next(r for r in plan['regions'] if r['key']=='title')
            assert abs(title['line_spacing_pt']/title['size_pt']-1.3)<.001
            assert title['size_pt']==(12 if kind=='defense-record' else 14)
            for key in (('members-label','members','recorder','leader','date') if kind=='defense-record' else ('signature','date')):
                assert next(r for r in plan['regions'] if r['key']==key)==next(r for r in sp['regions'] if r['key']==key)
            if kind=='defense-record':
                first=[line for line in plan['table']['vertical_segments_mm'] if abs(line[1]-sp['table']['top_mm'])<.001]
                assert len(first)==5 and all(abs(line[2]-(sp['table']['top_mm']+plan['table']['rows_mm'][0]))<.001 for line in first)
                if name=='title-layout-rich':
                    # Office may retain Asian/Latin spacing and wrap to five
                    # lines even when the PDF renderer needs only four.
                    assert plan['table']['rows_mm'][0]>=6+(12+4*15.6)*25.4/72-.001
            assert title_row_plan(data,plan,fonts)==plan and sp==original
        raw['metadata']['title']='过长合成标题'*200;rejected(raw,kind,'metadata.title')
        if kind=='defense-record':
            raw=fixture(kind,'capacity-boundary')
            raw['metadata']['title']=fixture(kind,'title-layout-long')['metadata']['title']
            rejected(raw,kind,'sections.questions')
    print('PASS all title layouts: 1.3 line ratio, padding, rich marks, fixed footer and continuous columns',flush=True)

def title_layout_render_tests(output):
    for kind in KINDS:
        for name in ('normal','title-layout-long','title-layout-rich','title-layout-breaks'):
            source=TEMPLATES/kind/'fixtures'/f'{name}.json';bundle=output/f'{kind}-{name}'
            render_bundle(kind,source,bundle,overwrite=True,compile_pdf=True);inspect_pdf(bundle/'main.pdf',kind,name)
            data=validate_data(fixture(kind,name),kind);sp=layout(kind);plan,_=writing_plan(data,kind,sp,resolved_fonts(sp))
            title=next(r for r in plan['regions'] if r['key']=='title');top=plan['table']['top_mm'];end=top+plan['table']['rows_mm'][0]
            _,words=pdf_words(bundle/'main.pdf')
            title_words=[w for w in words if top<float(w.attrib['yMin'])*25.4/72<end and title['x_mm']-.1<=float(w.attrib['xMin'])*25.4/72 and float(w.attrib['xMax'])*25.4/72<=title['x_mm']+title['width_mm']+.1]
            assert title_words
            y0=min(float(w.attrib['yMin'])*25.4/72 for w in title_words);y1=max(float(w.attrib['yMax'])*25.4/72 for w in title_words)
            assert y0-top>2.2 and end-y1>2.2,('title padding',kind,name,y0-top,end-y1)
            if name!='title-layout-rich':
                lines=sorted(set(round(float(w.attrib['yMin'])*25.4/72,2) for w in title_words))
                for a,b in zip(lines,lines[1:]):
                    if b-a>.1:assert abs((b-a)*72/25.4/title['size_pt']-1.3)<.03,('title leading',kind,name,lines)
                if name=='title-layout-breaks':assert len(lines)==3
            word=bundle/'result.docx';render_word(kind,TEMPLATES/kind/'word/official-template.docx',source,word,overwrite=True)
            inspect_docx(word,kind,fixture(kind,name));doc=Document(word);cell=doc.tables[0].cell(0,1 if kind=='defense-record' else 0)
            p=cell.paragraphs[0];assert abs(p.paragraph_format.line_spacing.pt/title['size_pt']-1.3)<.001
            assert p._p.pPr.find(qn('w:snapToGrid')).get(qn('w:val'))=='false'
            assert p._p.pPr.find(qn('w:textAlignment')).get(qn('w:val'))=='center'
            assert runs_text(data['metadata']['title']) in cell.text
            if name=='title-layout-rich':
                assert sum(bool(r.font.subscript) for r in p.runs)==3
                assert sum(bool(r.font.superscript) for r in p.runs)==1
                assert any(r.bold for r in p.runs) and any(r.italic for r in p.runs)
            print(f'PASS title spacing PDF/DOCX: {kind}/{name}',flush=True)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--render',action='store_true');parser.add_argument('--editable-render',action='store_true');parser.add_argument('--guidance-render',action='store_true');parser.add_argument('--identity-render',action='store_true');parser.add_argument('--grade-layout-render',action='store_true');parser.add_argument('--all-major-render',action='store_true');parser.add_argument('--title-layout-render',action='store_true');parser.add_argument('--output-dir',type=Path,default=TEMPLATES.parent/'tmp/assessment-regression');a=parser.parse_args()
    contract_tests()
    editable_contract_tests()
    guidance_contract_tests()
    identity_contract_tests()
    grade_layout_contract_tests()
    defense_major_contract_tests()
    title_layout_contract_tests()
    if a.title_layout_render:title_layout_render_tests(a.output_dir)
    if a.all_major_render:
        for kind in ('grade-assessment','advisor-review','reviewer-review'):identity_render_tests(a.output_dir,kind)
        defense_major_render_tests(a.output_dir)
        defense_major_render_tests(a.output_dir,'long-major-guidance')
    if a.grade_layout_render:grade_layout_render_tests(a.output_dir)
    if a.identity_render:identity_render_tests(a.output_dir)
    if a.render:render_tests(a.output_dir)
    if a.editable_render:editable_render_tests(a.output_dir)
    if a.guidance_render:guidance_render_tests(a.output_dir)
    return 0
if __name__=='__main__':raise SystemExit(main())
