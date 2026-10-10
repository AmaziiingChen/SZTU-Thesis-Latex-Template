"""Measure advisor glyph margins for one- and two-line college cells."""
import json
import subprocess
import sys
from pathlib import Path

import pdfplumber
from PIL import Image
from docx import Document
from docx.oxml.ns import qn

TEMPLATE = Path(__file__).resolve().parent
ROOT = TEMPLATE.parents[1]


def main():
    layout = json.loads((TEMPLATE / "spec/layout.json").read_text(encoding="utf-8"))
    reports = []
    for name, college in (("single", "合成学院"), ("wrapped", "大数据与互联网学院")):
        folder = ROOT / "tmp/proposal-advisor-tests" / name
        folder.mkdir(parents=True, exist_ok=True)
        data = json.loads((TEMPLATE / "fixtures/short.json").read_text(encoding="utf-8"))
        data["metadata"]["college"] = college
        fixture = folder / "input.json"
        fixture.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        subprocess.run([sys.executable, str(TEMPLATE / "latex/render.py"), "--data", str(fixture), "--output-dir", str(folder), "--compile", "--overwrite"], check=True, capture_output=True)
        subprocess.run([sys.executable, str(TEMPLATE / "word/render.py"), "--data", str(fixture), "--output", str(folder / "proposal.docx"), "--overwrite"], check=True, capture_output=True)
        subprocess.run(["pdftoppm", "-f", "1", "-l", "1", "-singlefile", "-r", "300", "-png", str(folder / "main.pdf"), str(folder / "page")], check=True)
        image = Image.open(folder / "page.png").convert("L")
        with pdfplumber.open(folder / "main.pdf") as pdf:
            page = pdf.pages[0]
            sx, sy = image.width / page.width, image.height / page.height
            edges = [e["top"] for e in page.edges if e["orientation"] == "h"]
            measured = []
            for word in page.extract_words():
                if word["text"] not in {"指导教师", data["metadata"]["advisor"]}:
                    continue
                top = max(y for y in edges if y < word["top"])
                bottom = min(y for y in edges if y > word["bottom"])
                assert abs((bottom - top) * 25.4 / 72 - 12.96) < 0.2
                crop = (int(word["x0"] * sx) - 1, int((top + 1) * sy), int(word["x1"] * sx) + 1, int((bottom - 1) * sy))
                ink = image.crop(crop).point(lambda p: 255 if p < 100 else 0).getbbox()
                assert ink is not None
                upper = ((crop[1] + ink[1]) / sy - top) * 25.4 / 72
                lower = (bottom - (crop[1] + ink[3]) / sy) * 25.4 / 72
                assert abs(upper - lower) / 2 <= 0.3, (name, word["text"], upper, lower)
                measured.append({"text": word["text"], "upper_mm": upper, "lower_mm": lower})
            assert len(measured) == 2
        table = Document(folder / "proposal.docx").tables[0]
        for index, key in ((2, "advisor_label_baseline_shift_pt"), (4, "advisor_value_baseline_shift_pt")):
            for paragraph in table.rows[2].cells[index].paragraphs:
                for run in paragraph.runs:
                    assert run._r.rPr.find(qn("w:position")).get(qn("w:val")) == str(round(layout["table"][key] * 2))
        reports.append({"case": name, "measurements": measured})
    (ROOT / "tmp/proposal-advisor-tests/report.json").write_text(json.dumps(reports, ensure_ascii=False, indent=2), encoding="utf-8")
    print("PASS: advisor optical centering, unchanged row height, and Word baseline settings")


if __name__ == "__main__":
    main()
