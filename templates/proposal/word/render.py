#!/usr/bin/env python3
"""Render the official SZTU proposal DOCX from schema-constrained JSON data."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image as PILImage

TEMPLATES_DIR = Path(__file__).resolve().parents[2]
if str(TEMPLATES_DIR) not in sys.path:
    sys.path.insert(0, str(TEMPLATES_DIR))

from common.python.content import (  # noqa: E402
    ContentDataError as DataError,
    display_width,
    figure_caption_runs,
    list_marker_text,
    normalize_content_block,
    normalize_paragraph,
    plain_runs,
    prepare_equation_content,
    prepare_figure_content,
    require_object,
    require_text,
    resolve_list_marker,
    runs_text,
    split_numbered_subitems,
)
from common.python.process_form import set_word_cell_vertical_padding, load_process_document_layout  # noqa: E402
from common.python.equation import add_numbered_omml_table, append_omml  # noqa: E402


SCHEMA_VERSION = "0.2"
METADATA_LIMITS = {
    "title": 80,
    "student_name": 20,
    "student_id": 30,
    "major": 40,
    "college": 50,
    "advisor": 30,
}
SECTION_KEYS = (
    "significance_and_status",
    "research_content",
    "methods_and_means",
    "research_steps",
    "references",
)
ADVISOR_TITLE_PATTERN = re.compile(r"(?:老师|教授|副教授|讲师|博士|硕士|导师)$")
TITLE_DISPLAY_WIDTH_LIMIT = 64.0
LAYOUT = load_process_document_layout(
    Path(__file__).resolve().parents[1] / "spec" / "layout.json"
)
BODY_STYLE = LAYOUT["typography"]["body"]
BODY_SIZE_PT = BODY_STYLE["size_pt"]
BODY_CJK_FONT = BODY_STYLE["cjk_word_family"]
BODY_LATIN_FONT = BODY_STYLE["latin_word_family"]
LINE_SPACING = LAYOUT["paragraphs"]["line_spacing"]
FIRST_LINE_INDENT_PT = BODY_SIZE_PT * LAYOUT["paragraphs"]["first_line_indent_em"]
NESTED_LIST_LEFT_INDENT_PT = (
    BODY_SIZE_PT * LAYOUT["paragraphs"]["nested_list_left_indent_em"]
)
LIST_LEVEL_INDENT_PT = BODY_SIZE_PT * LAYOUT["paragraphs"]["list_level_indent_em"]
LIST_HANGING_INDENT_PT = BODY_SIZE_PT * LAYOUT["paragraphs"]["list_hanging_indent_em"]
LIST_MAX_DEPTH = LAYOUT["paragraphs"]["list_max_depth"]
UNORDERED_LIST_MARKERS = LAYOUT["paragraphs"]["unordered_list_markers"]


def validate_data(raw: Any) -> dict[str, Any]:
    data = require_object(raw, "root")
    expected_root = {"schema_version", "metadata", "sections"}
    unknown_root = set(data) - expected_root
    if unknown_root:
        raise DataError(f"unknown root fields: {sorted(unknown_root)}")
    if data.get("schema_version") != SCHEMA_VERSION:
        raise DataError(f"schema_version must be {SCHEMA_VERSION!r}")

    metadata = require_object(data.get("metadata"), "metadata")
    unknown_metadata = set(metadata) - set(METADATA_LIMITS)
    if unknown_metadata:
        raise DataError(f"unknown metadata fields: {sorted(unknown_metadata)}")
    clean_metadata: dict[str, Any] = {}
    for key, limit in METADATA_LIMITS.items():
        value = metadata.get(key)
        clean_metadata[key] = (
            normalize_paragraph(value, f"metadata.{key}", limit)
            if key == "title"
            else require_text(value, f"metadata.{key}", limit)
        )
    if display_width(runs_text(clean_metadata["title"])) > TITLE_DISPLAY_WIDTH_LIMIT:
        raise DataError(
            "metadata.title is too wide for the official two-line title cell "
            f"(limit {TITLE_DISPLAY_WIDTH_LIMIT:g} display units)"
        )
    if ADVISOR_TITLE_PATTERN.search(clean_metadata["advisor"]):
        raise DataError("metadata.advisor must contain the name only, without a title")

    sections = require_object(data.get("sections"), "sections")
    unknown_sections = set(sections) - set(SECTION_KEYS)
    if unknown_sections:
        raise DataError(f"unknown section fields: {sorted(unknown_sections)}")
    clean_sections: dict[str, Any] = {}
    for key in SECTION_KEYS:
        value = sections.get(key)
        if not isinstance(value, list) or not value:
            raise DataError(f"sections.{key} must be a non-empty array")
        if key == "significance_and_status" and len(value) < 2:
            raise DataError("sections.significance_and_status must contain at least 2 paragraphs")
        max_length = 1000 if key == "references" else 5000
        if len(value) > 50:
            raise DataError(f"sections.{key} may contain at most 50 items")
        if key in {"research_content", "methods_and_means", "research_steps"}:
            normalized_blocks = [
                normalize_content_block(
                    item,
                    f"sections.{key}[{index}]",
                    max_length,
                    max_list_depth=LIST_MAX_DEPTH,
                )
                for index, item in enumerate(value)
            ]
            clean_sections[key] = normalized_blocks
        else:
            clean_sections[key] = [
                normalize_paragraph(item, f"sections.{key}[{index}]", max_length)
                for index, item in enumerate(value)
            ]

    prepare_figure_content(
        [
            clean_sections["research_content"],
            clean_sections["methods_and_means"],
            clean_sections["research_steps"],
        ]
    )
    prepare_equation_content(
        [
            clean_sections["research_content"],
            clean_sections["methods_and_means"],
            clean_sections["research_steps"],
        ]
    )
    return {
        "schema_version": SCHEMA_VERSION,
        "metadata": clean_metadata,
        "sections": clean_sections,
    }


def _set_run_font(run, *, name: str, size_pt: float, bold: bool = False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    run.bold = bold


def _append_runs(paragraph, runs: list[dict[str, Any]], *, size_pt: float) -> None:
    for rich_run in runs:
        pieces = re.findall(r"[\x00-\x7f]+|[^\x00-\x7f]+", rich_run["text"])
        for piece in pieces:
            run = paragraph.add_run(piece)
            font_name = BODY_LATIN_FONT if ord(piece[0]) < 128 else BODY_CJK_FONT
            _set_run_font(run, name=font_name, size_pt=size_pt, bold=rich_run["bold"])
            run.italic = rich_run["italic"]
            run.font.subscript = rich_run["script"] == "sub"
            run.font.superscript = rich_run["script"] == "super"


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def _fill_info_cell(
    cell,
    text: str | list[dict[str, Any]],
    *,
    size_pt: float = BODY_SIZE_PT,
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.CENTER,
) -> None:
    while len(cell.paragraphs) > 1:
        _remove_paragraph(cell.paragraphs[-1])
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = LINE_SPACING
    runs = (
        [{"text": text, "script": "normal", "italic": False, "bold": False}]
        if isinstance(text, str)
        else text
    )
    _append_runs(paragraph, runs, size_pt=size_pt)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _center_label_cell(cell) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)


def _format_body_paragraph(paragraph, *, references: bool, indent: bool = True) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if references else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = LINE_SPACING
    paragraph.paragraph_format.first_line_indent = (
        None if references or not indent else Pt(FIRST_LINE_INDENT_PT)
    )


def _fill_section(
    cell,
    paragraphs: list[list[dict[str, Any]]],
    *,
    references: bool = False,
) -> None:
    if not cell.paragraphs:
        raise RuntimeError("section cell has no label paragraph")
    # Reuse the official paragraph slots, then remove unused blank slots. The
    # table rows already carry their official minimum heights; retaining every
    # blank paragraph would add its line height on top of wrapped content and
    # can push an otherwise normal document onto a third page in Word.
    slots = list(cell.paragraphs[1:])
    for index, text in enumerate(paragraphs):
        paragraph = slots[index] if index < len(slots) else cell.add_paragraph()
        paragraph.clear()
        _format_body_paragraph(paragraph, references=references)
        _append_runs(paragraph, text, size_pt=BODY_SIZE_PT)
    for paragraph in slots[len(paragraphs) :]:
        _remove_paragraph(paragraph)


def _append_list_block(cell, block: dict[str, Any], *, depth: int = 1) -> None:
    if not 1 <= depth <= LIST_MAX_DEPTH:
        raise AssertionError(f"normalized list depth escaped bounds: {depth}")
    for index, item in enumerate(block["items"], start=1):
        paragraph = cell.add_paragraph()
        _format_body_paragraph(paragraph, references=False, indent=False)
        paragraph.paragraph_format.left_indent = Pt(
            depth * LIST_LEVEL_INDENT_PT + LIST_HANGING_INDENT_PT
        )
        paragraph.paragraph_format.first_line_indent = Pt(-LIST_HANGING_INDENT_PT)
        marker = resolve_list_marker(
            block["type"], item["marker"], index, depth, UNORDERED_LIST_MARKERS
        )
        _append_runs(
            paragraph, plain_runs(list_marker_text(marker)), size_pt=BODY_SIZE_PT
        )
        _append_runs(paragraph, item["runs"], size_pt=BODY_SIZE_PT)
        if item["children"]:
            _append_list_block(cell, item["children"], depth=depth + 1)


def _resolve_image(path_text: str, data_dir: Path) -> Path:
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", path_text):
        raise DataError("remote image URLs are forbidden; use a local file")
    path = Path(path_text).expanduser()
    if not path.is_absolute():
        path = data_dir / path
    path = path.resolve()
    if not path.is_file():
        raise DataError(f"image file does not exist: {path}")
    return path


def _set_picture_alt(run, alt: str) -> None:
    items = run._r.xpath(".//wp:docPr")
    if items:
        items[0].set("descr", alt)


def _fit_image_dimensions(
    image_path: Path,
    requested_width_mm: float,
    *,
    max_height_mm: float,
) -> tuple[float, float]:
    with PILImage.open(image_path) as image:
        pixel_width, pixel_height = image.size
    if pixel_width <= 0 or pixel_height <= 0:
        raise DataError("image dimensions must be positive")
    width_mm = float(requested_width_mm)
    height_mm = width_mm * pixel_height / pixel_width
    if height_mm > max_height_mm:
        scale = max_height_mm / height_mm
        width_mm *= scale
        height_mm = max_height_mm
    return width_mm, height_mm


def _set_table_borders(table, *, border_pt: float | None) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        if border_pt is None:
            element.set(qn("w:val"), "nil")
        else:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(round(border_pt * 8)))
            element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)


def _set_table_cell_margins(table, margin_mm: float) -> None:
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            old = tc_pr.find(qn("w:tcMar"))
            if old is not None:
                tc_pr.remove(old)
            margins = OxmlElement("w:tcMar")
            value = str(round(Mm(margin_mm).twips))
            for edge in ("top", "left", "bottom", "right"):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:w"), value)
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            tc_pr.append(margins)


def _set_table_geometry(table, widths_mm: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total_twips = round(Mm(sum(widths_mm)).twips)
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table._tbl.tblPr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    for grid_col, width_mm in zip(table._tbl.tblGrid.gridCol_lst, widths_mm, strict=True):
        grid_col.set(qn("w:w"), str(round(Mm(width_mm).twips)))
    for row in table.rows:
        for cell, width_mm in zip(row.cells, widths_mm, strict=True):
            cell.width = Mm(width_mm)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(round(Mm(width_mm).twips)))
            tc_w.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _block_alignment(value: str):
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[value]


def _append_data_table(cell, block: dict[str, Any]) -> None:
    config = LAYOUT["embedded_table"]
    weights = [column["width_weight"] for column in block["columns"]]
    total_weight = sum(weights)
    widths = [float(config["width_mm"]) * weight / total_weight for weight in weights]
    table = cell.add_table(rows=1 + len(block["rows"]), cols=len(block["columns"]))
    _set_table_geometry(table, widths)
    _set_table_borders(table, border_pt=float(config["border_pt"]))
    _set_table_cell_margins(table, float(config["cell_padding_mm"]))
    _set_repeat_table_header(table.rows[0])
    for row in table.rows:
        _prevent_row_split(row)
    for column_index, column in enumerate(block["columns"]):
        target = table.rows[0].cells[column_index]
        target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = target.paragraphs[0]
        _format_body_paragraph(paragraph, references=False, indent=False)
        paragraph.alignment = _block_alignment(column["alignment"])
        _append_runs(paragraph, column["header_runs"], size_pt=BODY_SIZE_PT)
        for run in paragraph.runs:
            run.bold = True
    for row_index, values in enumerate(block["rows"], start=1):
        for column_index, runs in enumerate(values):
            target = table.rows[row_index].cells[column_index]
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = target.paragraphs[0]
            _format_body_paragraph(paragraph, references=False, indent=False)
            paragraph.alignment = _block_alignment(block["columns"][column_index]["alignment"])
            _append_runs(paragraph, runs, size_pt=BODY_SIZE_PT)
    if block["caption_runs"]:
        caption = cell.add_paragraph()
        _format_body_paragraph(caption, references=False, indent=False)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(config["caption_space_before_pt"])
        caption.paragraph_format.space_after = Pt(config["caption_space_after_pt"])
        caption.paragraph_format.keep_together = True
        _append_runs(caption, block["caption_runs"], size_pt=BODY_SIZE_PT)


def _append_figure_group(cell, block: dict[str, Any], *, data_dir: Path) -> None:
    config = LAYOUT["figure_group"]
    count = len(block["items"])
    total_width = float(config["width_mm"])
    gap = float(config["column_gap_mm"])
    item_width = (total_width - gap * (count - 1)) / count
    table = cell.add_table(rows=1, cols=count)
    _set_table_geometry(table, [item_width] * count)
    _set_table_borders(table, border_pt=None)
    _set_table_cell_margins(table, gap / 2)
    _prevent_row_split(table.rows[0])
    for index, item in enumerate(block["items"]):
        target = table.rows[0].cells[index]
        target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = target.paragraphs[0]
        _format_body_paragraph(paragraph, references=False, indent=False)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run()
        _set_run_font(run, name=BODY_CJK_FONT, size_pt=BODY_SIZE_PT)
        image_path = _resolve_image(item["path"], data_dir)
        width_mm, height_mm = _fit_image_dimensions(
            image_path,
            item_width - gap,
            max_height_mm=float(config["max_item_height_mm"]),
        )
        run.add_picture(str(image_path), width=Mm(width_mm), height=Mm(height_mm))
        _set_picture_alt(run, item["alt"])
        subcaption = target.add_paragraph()
        _format_body_paragraph(subcaption, references=False, indent=False)
        subcaption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subcaption.paragraph_format.keep_together = True
        _append_runs(subcaption, plain_runs(item["subfigure_label"]), size_pt=BODY_SIZE_PT)
        if item["caption_runs"]:
            _append_runs(subcaption, item["caption_runs"], size_pt=BODY_SIZE_PT)
    caption = cell.add_paragraph()
    _format_body_paragraph(caption, references=False, indent=False)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.keep_together = True
    _append_runs(caption, figure_caption_runs(block), size_pt=BODY_SIZE_PT)


def _append_image(cell, block: dict[str, Any], *, data_dir: Path) -> None:
    paragraph = cell.add_paragraph()
    _format_body_paragraph(paragraph, references=False, indent=False)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    _set_run_font(run, name=BODY_CJK_FONT, size_pt=BODY_SIZE_PT)
    image_path = _resolve_image(block["path"], data_dir)
    width_mm, height_mm = _fit_image_dimensions(
        image_path,
        block["width_mm"],
        max_height_mm=float(LAYOUT["image"]["max_height_mm"]),
    )
    run.add_picture(str(image_path), width=Mm(width_mm), height=Mm(height_mm))
    _set_picture_alt(run, block["alt"])

    # WPS may ignore keep_with_next across paragraphs inside a split table row.
    # Keep the inline image and its caption in one indivisible paragraph.
    paragraph.add_run().add_break()
    paragraph.paragraph_format.space_after = Pt(3)
    _append_runs(paragraph, figure_caption_runs(block), size_pt=BODY_SIZE_PT)


def _append_text_content_blocks(
    cell,
    blocks: list[dict[str, Any]],
    *,
    data_dir: Path,
) -> None:
    for block in blocks:
        if block["type"] == "paragraph":
            paragraph = cell.add_paragraph()
            _format_body_paragraph(paragraph, references=False)
            _append_runs(paragraph, block["runs"], size_pt=BODY_SIZE_PT)
        elif block["type"] in {"ordered_list", "unordered_list"}:
            _append_list_block(cell, block)
        elif block["type"] == "image":
            _append_image(cell, block, data_dir=data_dir)
        elif block["type"] == "data_table":
            _append_data_table(cell, block)
        elif block["type"] == "figure_group":
            _append_figure_group(cell, block, data_dir=data_dir)
        elif block["type"] == "equation":
            label = block.get("equation_label")
            if label is not None:
                _, paragraphs = add_numbered_omml_table(
                    cell,
                    block["expression"],
                    label,
                    width_mm=float(LAYOUT["equation"]["width_mm"]),
                )
                for equation_paragraph in paragraphs:
                    _format_body_paragraph(
                        equation_paragraph, references=False, indent=False
                    )
                    equation_paragraph.paragraph_format.keep_together = True
                paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraphs[2].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                center = paragraphs[1]
                center.paragraph_format.space_before = Pt(
                    LAYOUT["equation"]["space_before_pt"]
                )
                center.paragraph_format.space_after = Pt(
                    LAYOUT["equation"]["space_after_pt"]
                )
                fallback = center.add_run(block["alt"])
                _set_run_font(fallback, name=BODY_CJK_FONT, size_pt=BODY_SIZE_PT)
                fallback.font.hidden = True
                _set_run_font(
                    paragraphs[2].runs[0], name=BODY_CJK_FONT, size_pt=BODY_SIZE_PT
                )
                continue
            paragraph = cell.add_paragraph()
            _format_body_paragraph(paragraph, references=False, indent=False)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(
                LAYOUT["equation"]["space_before_pt"]
            )
            paragraph.paragraph_format.space_after = Pt(
                LAYOUT["equation"]["space_after_pt"]
            )
            paragraph.paragraph_format.keep_together = True
            fallback = paragraph.add_run(block["alt"])
            _set_run_font(fallback, name=BODY_CJK_FONT, size_pt=BODY_SIZE_PT)
            fallback.font.hidden = True
            append_omml(paragraph, block["expression"])
        else:
            raise AssertionError(f"unsupported normalized block: {block['type']}")


def _fill_text_content_section(cell, blocks: list[dict[str, Any]], *, data_dir: Path) -> None:
    if all(block["type"] == "paragraph" for block in blocks):
        _fill_section(cell, [block["runs"] for block in blocks])
        return
    for paragraph in list(cell.paragraphs[1:]):
        _remove_paragraph(paragraph)
    _append_text_content_blocks(cell, blocks, data_dir=data_dir)


def _fill_methods_section(
    cell,
    methods: list[dict[str, Any]],
    steps: list[dict[str, Any]],
    *,
    data_dir: Path,
) -> None:
    if not cell.paragraphs:
        raise RuntimeError("methods cell has no label paragraph")
    for paragraph in list(cell.paragraphs[1:]):
        _remove_paragraph(paragraph)

    groups = (
        ("本课题研究方法、手段如下：", methods),
        ("本课题研究步骤如下：", steps),
    )
    for heading, items in groups:
        heading_paragraph = cell.add_paragraph()
        _format_body_paragraph(heading_paragraph, references=False)
        _append_runs(heading_paragraph, plain_runs(heading), size_pt=BODY_SIZE_PT)
        if not all(item["type"] == "paragraph" for item in items):
            _append_text_content_blocks(cell, items, data_dir=data_dir)
            continue
        for item in items:
            runs = item["runs"]
            split_item = split_numbered_subitems(runs)
            paragraph = cell.add_paragraph()
            _format_body_paragraph(paragraph, references=False)
            lead = split_item[0] if split_item else runs
            _append_runs(paragraph, lead, size_pt=BODY_SIZE_PT)
            if split_item:
                _, subitems, continuations = split_item
                for marker, subitem in subitems:
                    nested = cell.add_paragraph()
                    _format_body_paragraph(nested, references=False)
                    nested.paragraph_format.left_indent = Pt(NESTED_LIST_LEFT_INDENT_PT)
                    _append_runs(
                        nested, plain_runs(f"{marker} "), size_pt=BODY_SIZE_PT
                    )
                    _append_runs(nested, subitem, size_pt=BODY_SIZE_PT)
                for continuation in continuations:
                    trailing = cell.add_paragraph()
                    _format_body_paragraph(trailing, references=False)
                    _append_runs(trailing, continuation, size_pt=BODY_SIZE_PT)


def _prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def render(template: Path, data_path: Path, output: Path, *, overwrite: bool) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"template not found: {template}")
    if output.suffix.lower() != ".docx":
        raise DataError("output path must end with .docx")
    if output.exists() and not overwrite:
        raise FileExistsError(f"output exists; pass --overwrite to replace it: {output}")

    data = validate_data(json.loads(data_path.read_text(encoding="utf-8")))
    document = Document(template)
    if len(document.tables) != 1:
        raise RuntimeError("expected exactly one table in the official template")
    table = document.tables[0]
    if len(table.rows) != 9 or len(table.columns) != 7:
        raise RuntimeError("expected the official 9-row, 7-column table grid")

    # Preserve the official compact metadata rows while allowing long titles,
    # college names, or majors to wrap without being clipped by an exact height.
    for row_index in (0, 1, 2):
        table.rows[row_index].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    for cell in table.rows[0].cells:
        set_word_cell_vertical_padding(cell, LAYOUT["table"]["title_vertical_padding_mm"])

    metadata = data["metadata"]
    for row_index, cell_index in ((0, 0), (1, 0), (1, 2), (1, 5), (2, 0), (2, 2)):
        _center_label_cell(table.rows[row_index].cells[cell_index])

    _fill_info_cell(
        table.rows[0].cells[1],
        metadata["title"],
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _fill_info_cell(table.rows[1].cells[1], metadata["student_name"])
    _fill_info_cell(table.rows[1].cells[3], metadata["student_id"])
    _fill_info_cell(
        table.rows[1].cells[6],
        metadata["major"],
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _fill_info_cell(
        table.rows[2].cells[1],
        metadata["college"],
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _fill_info_cell(table.rows[2].cells[4], metadata["advisor"])

    sections = data["sections"]
    _fill_section(table.rows[3].cells[0], sections["significance_and_status"])
    _fill_text_content_section(
        table.rows[4].cells[0],
        sections["research_content"],
        data_dir=data_path.resolve().parent,
    )
    _fill_methods_section(
        table.rows[5].cells[0],
        sections["methods_and_means"],
        sections["research_steps"],
        data_dir=data_path.resolve().parent,
    )
    _fill_section(table.rows[6].cells[0], sections["references"], references=True)

    # Keep handwritten signature and review regions together when they fit,
    # while allowing all preceding content to flow naturally across pages.
    for row_index, height_mm in {
        7: LAYOUT["signature_regions"]["student_min_height_mm"],
        8: LAYOUT["signature_regions"]["review_min_height_mm"],
    }.items():
        table.rows[row_index].height = Mm(height_mm)
        table.rows[row_index].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _prevent_row_split(table.rows[row_index])

    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.comments = (
        "Generated deterministically from proposal schema v0.2; "
        "signature and review fields intentionally remain blank."
    )
    document.save(output)

    reopened = Document(output)
    if len(reopened.tables) != 1 or len(reopened.tables[0].rows) != 9:
        raise RuntimeError("rendered DOCX failed structural reopening check")


def main() -> int:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, required=True, help="proposal JSON data")
    parser.add_argument("--output", type=Path, required=True, help="output .docx path")
    parser.add_argument(
        "--template",
        type=Path,
        default=script_dir / "official-template.docx",
        help="official DOCX template",
    )
    parser.add_argument("--overwrite", action="store_true", help="replace an existing output")
    args = parser.parse_args()

    try:
        render(args.template, args.data, args.output, overwrite=args.overwrite)
    except (DataError, FileNotFoundError, FileExistsError, json.JSONDecodeError, RuntimeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
