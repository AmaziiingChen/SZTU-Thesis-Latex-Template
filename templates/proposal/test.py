#!/usr/bin/env python3
"""Run structural and rendered regression checks for the proposal prototype."""

from __future__ import annotations

import argparse
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

TEMPLATES_DIR = Path(__file__).resolve().parents[1]
if str(TEMPLATES_DIR) not in sys.path:
    sys.path.insert(0, str(TEMPLATES_DIR))

from common.python.pdf_geometry import (  # noqa: E402
    assert_centered,
    assert_rule_has_raster_ink,
    assert_vertically_centered,
    cell_chars,
    clustered,
)
from common.python.process_form import load_process_document_layout  # noqa: E402
from common.python.list_layout import marker_advance_pt  # noqa: E402
from common.python.regression_fixtures import (  # noqa: E402
    load_page_range_recipes,
    materialize_page_range_fixture,
)


def run(command: list[str], *, cwd: Path) -> str:
    result = subprocess.run(
        command,
        cwd=cwd,
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
    )
    try:
        return result.stdout.decode("utf-8")
    except UnicodeDecodeError:
        return result.stdout.decode("mbcs" if sys.platform == "win32" else "utf-8", errors="replace")


def load_renderer(path: Path, module_name: str):
    spec = importlib.util.spec_from_file_location(module_name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load renderer: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-pdf", action="store_true", help="skip XeLaTeX and CJK checks")
    args = parser.parse_args()
    if not args.skip_pdf:
        import pdfplumber

    proposal_dir = Path(__file__).resolve().parent
    project_dir = proposal_dir.parents[1]
    output_root = project_dir / "tmp" / "proposal-tests"
    output_root.mkdir(parents=True, exist_ok=True)
    layout = load_process_document_layout(proposal_dir / "spec" / "layout.json")
    word_renderer = load_renderer(proposal_dir / "word" / "render.py", "proposal_word_test")
    latex_renderer = load_renderer(proposal_dir / "latex" / "render.py", "proposal_latex_test")
    emphasis_runs = [
        {"text": "中文加粗", "bold": True, "italic": False, "script": "normal"},
        {"text": "中文倾斜", "bold": False, "italic": True, "script": "normal"},
        {"text": "中文加粗倾斜", "bold": True, "italic": True, "script": "normal"},
    ]
    emphasis_tex = latex_renderer.rich_runs(emphasis_runs)
    emphasis_tex = emphasis_tex.replace(r"\nobreak{}", "")
    assert r"\textbf{中文加粗}" in emphasis_tex
    assert r"\textit{中文倾斜}" in emphasis_tex
    assert r"\textbf{\textit{中文加粗倾斜}}" in emphasis_tex
    font_dir = output_root / "synthetic-fonts"
    font_tex = latex_renderer.fonts_tex(
        {
            "SimSun": font_dir / "simsun.ttf",
            "SimHei": font_dir / "simhei.ttf",
            "Times New Roman": font_dir / "times.ttf",
            "Times New Roman Bold": font_dir / "timesbd.ttf",
            "Times New Roman Italic": font_dir / "timesi.ttf",
            "Times New Roman Bold Italic": font_dir / "timesbi.ttf",
        }
    )
    assert "AutoFakeBold=3,AutoFakeSlant=0.2" in font_tex
    assert font_tex.count("AutoFakeSlant=0.2") >= 4
    assert layout["typography"]["title"]["size_pt"] == 18.0
    assert layout["typography"]["label"]["size_pt"] == 12.0
    assert layout["typography"]["body"]["size_pt"] == 10.5
    assert layout["typography"]["title"]["bold"] is True
    assert layout["typography"]["label"]["bold"] is False
    assert layout["table"]["border_pt"] == 0.48
    assert layout["table"]["title_single_line_vertical_padding_mm"] == 0.0
    assert layout["table"]["title_vertical_padding_mm"] == 1.5
    assert layout["table"]["flow_vertical_padding_mm"] == 0.8
    assert layout["table"]["section_title_content_gap_mm"] == 0.0
    assert layout["signature_regions"]["teacher_opinion_region_height_mm"] == 50.0
    assert layout["signature_regions"]["opinion_transition_gap_mm"] == 1.0
    assert layout["paragraphs"]["list_max_depth"] == 4
    assert layout["paragraphs"]["unordered_list_markers"] == ["•", "◦", "▪", "▫"]
    assert layout["paragraphs"]["list_marker_gap_em"] == 0.25
    assert layout["paragraphs"]["list_level_indent_em"] == 1.0
    assert layout["paragraphs"]["list_hanging_indent_em"] == 0.0

    latex_template = (proposal_dir / "latex" / "main.tex").read_text(encoding="utf-8")
    assert r"\input{proposal-typography.tex}" in latex_template
    assert r"\input{sztu-process-form.tex}" in latex_template
    assert r"{\heitiBold\bfseries\SZTUTitleSize 深圳技术大学本科毕业论文（设计）\par}" in latex_template
    assert r"\newcommand{\Label}[1]{{\heiti\SZTULabelSize #1}}" in latex_template
    assert r"\newcommand{\SignatureText}[1]{{\songti\SZTUSignatureSize #1}}" in latex_template
    assert r"\noindent 本课题研究方法、手段如下" not in latex_template
    assert r"\noindent 本课题研究步骤如下" not in latex_template
    assert r"\newcommand{\NestedOrderedItem}" in latex_template
    assert r"\newcommand{\ProposalListItem}" in latex_template
    assert r"\newcommand{\ProposalFigure}" in latex_template
    assert r"height=\SZTUImageMaxHeight" in latex_template
    assert r"\ProposalSeparator" not in latex_template
    assert latex_template.count(r"\begin{tcolorbox}[proposalflow]") == 4
    assert latex_template.count(r"\begin{tcolorbox}[proposalfixed") == 2
    assert latex_template.count(r"\SZTUFormTightJoin") == 6
    signature_section_start = latex_template.index(
        r"\begin{tcolorbox}[proposalfixed"
    )
    body_end = latex_template.rfind(
        r"\end{tcolorbox}", 0, signature_section_start
    )
    between_body_and_signature = latex_template[body_end:signature_section_start]
    assert r"\end{adjustwidth}" not in between_body_and_signature
    assert r"\begin{center}" not in between_body_and_signature

    fixture_names = (
        "short",
        "title-two-lines",
        "normal",
        "long",
        "layout_stress",
        "nested-list",
        "structured-content",
        "page-range-maximum",
    )
    page_range_recipes = load_page_range_recipes(
        proposal_dir.parent / "common" / "fixtures" / "page-range-stress-recipes.json"
    )
    page_range_recipe = page_range_recipes["proposal"]
    page_range_fixture = materialize_page_range_fixture(
        proposal_dir / "fixtures" / page_range_recipe["base_fixture"],
        page_range_recipe,
        output_root / "page-range-maximum.json",
    )
    for name in fixture_names:
        fixture = (
            page_range_fixture
            if name == "page-range-maximum"
            else proposal_dir / "fixtures" / f"{name}.json"
        )
        fixture_data = json.loads(fixture.read_text(encoding="utf-8"))
        assert fixture_data["schema_version"] == "0.2"
        assert "methods_and_means" in fixture_data["sections"]
        assert "research_steps" in fixture_data["sections"]
        assert "methods_and_steps" not in fixture_data["sections"]
        normalized = word_renderer.validate_data(fixture_data)
        docx_output = output_root / f"{name}.docx"
        run(
            [
                sys.executable,
                str(proposal_dir / "word" / "render.py"),
                "--data",
                str(fixture),
                "--output",
                str(docx_output),
                "--overwrite",
            ],
            cwd=project_dir,
        )
        document = Document(docx_output)
        assert len(document.tables) == 1
        table = document.tables[0]
        assert len(table.rows) == 9 and len(table.columns) == 7
        if name in {"short", "title-two-lines"}:
            title_row = table.rows[0]
            assert title_row.height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
            assert title_row.height.twips == 454
            for cell in (title_row.cells[0], title_row.cells[1]):
                margins = cell._tc.tcPr.find(qn("w:tcMar"))
                assert margins is not None
                for side in ("top", "bottom"):
                    margin = margins.find(qn(f"w:{side}"))
                    assert margin is not None
                    assert int(margin.get(qn("w:w"))) == 85
        assert "学生签名" in table.rows[7].cells[0].text
        assert "指导教师意见" in table.rows[8].cells[0].text
        assert "学院领导意见" in table.rows[8].cells[0].text
        assert not table.rows[5].cells[0].paragraphs[0].paragraph_format.page_break_before
        for section, row_index, starter in (
            ("research_content", 4, "本课题研究内容如下："),
            ("methods_and_means", 5, "本课题研究方法、手段如下："),
            ("research_steps", 5, "本课题研究步骤如下："),
        ):
            first = fixture_data["sections"][section][0]
            if first == starter:
                assert starter in table.rows[row_index].cells[0].text
            else:
                assert starter not in table.rows[row_index].cells[0].text
        for row_index in (3, 4, 5, 6):
            label = table.rows[row_index].cells[0].paragraphs[0]
            assert label.paragraph_format.space_before.pt == 0
            assert label.paragraph_format.space_after.pt == 0
            assert label.paragraph_format.line_spacing == 1
        method_paragraphs = table.rows[5].cells[0].paragraphs[1:]
        assert method_paragraphs
        if name not in {"nested-list", "structured-content"}:
            assert all(
                paragraph.paragraph_format.first_line_indent is not None
                and paragraph.paragraph_format.first_line_indent.pt
                == layout["typography"]["body"]["size_pt"]
                * layout["paragraphs"]["first_line_indent_em"]
                for paragraph in method_paragraphs
            )
        if name == "long":
            assert "第一阶段完成" in table.rows[5].cells[0].text
            assert "1、第一阶段完成" not in table.rows[5].cells[0].text
        assert table.rows[3].cells[0].paragraphs[1].paragraph_format.line_spacing == 1
        title_runs = table.rows[0].cells[1].paragraphs[0].runs
        assert title_runs and all(run.font.size.pt == 10.5 for run in title_runs)
        if name == "normal":
            assert any(run.text == "2" and run.font.subscript for run in title_runs)
            assert any(run.text == "3" and run.font.subscript for run in title_runs)
            assert any(run.text == "4" and run.font.subscript for run in title_runs)
            significance_runs = [
                run
                for paragraph in table.rows[3].cells[0].paragraphs
                for run in paragraph.runs
            ]
            assert any(run.text == "[1]" and run.font.superscript for run in significance_runs)
            methods_text = table.rows[5].cells[0].text
            assert "采用模板蒸馏方法" in methods_text
            assert "完成官方文件的结构分析" in methods_text
            assert "1、采用模板蒸馏方法" not in methods_text
            assert "2、采用确定性程序" not in methods_text
            assert "1、完成官方文件的结构分析" not in methods_text
        for row_index, cell_index in ((0, 0), (1, 0), (1, 2), (1, 5), (2, 0), (2, 2)):
            cell = table.rows[row_index].cells[cell_index]
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
            assert all(
                paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
                for paragraph in cell.paragraphs
            )
        for row_index, cell_index in ((0, 1), (1, 6), (2, 1)):
            cell = table.rows[row_index].cells[cell_index]
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
            assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
        for row_index, cell_index in ((1, 1), (1, 3), (2, 4)):
            cell = table.rows[row_index].cells[cell_index]
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
            assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
        if name == "normal":
            nested_paragraphs = [
                paragraph
                for paragraph in method_paragraphs
                if paragraph.text.startswith(("（1）", "（2）", "（3）"))
            ]
            assert [paragraph.text[:3] for paragraph in nested_paragraphs] == [
                "（1）",
                "（2）",
                "（3）",
            ]
            assert all(
                paragraph.paragraph_format.left_indent is not None
                and paragraph.paragraph_format.left_indent.pt
                == layout["typography"]["body"]["size_pt"]
                * layout["paragraphs"]["nested_list_left_indent_em"]
                for paragraph in nested_paragraphs
            )
            continuation = next(
                paragraph
                for paragraph in method_paragraphs
                if paragraph.text.startswith("通过统一数据")
            )
            assert continuation.paragraph_format.left_indent is None
            reference_runs = [
                run
                for paragraph in table.rows[6].cells[0].paragraphs
                for run in paragraph.runs
            ]
            assert any(run.text == "3" and run.font.subscript for run in reference_runs)
            assert any(run.text == "4" and run.font.subscript for run in reference_runs)

        if name == "nested-list":
            research_paragraphs = table.rows[4].cells[0].paragraphs[1:]
            by_prefix = {
                prefix: next(p for p in research_paragraphs if p.text.startswith(prefix))
                for prefix in ("•", "（1）", "▪", "A.", "列表结束后")
            }
            latin_font, cjk_font = word_renderer._list_font_files()
            for depth, prefix in enumerate(("•", "（1）", "▪", "A."), start=1):
                paragraph = by_prefix[prefix]
                assert paragraph.paragraph_format.left_indent is not None
                marker_width = marker_advance_pt(
                    prefix, size_pt=10.5, latin_font=latin_font, cjk_font=cjk_font
                )
                marker_start = 21.0 + (depth - 1) * 10.5
                assert abs(paragraph.paragraph_format.left_indent.pt - (marker_start + marker_width)) <= 0.05
                assert abs(paragraph.paragraph_format.first_line_indent.pt + marker_width) <= 0.05
            trailing = by_prefix["列表结束后"]
            assert trailing.paragraph_format.left_indent is None
            assert trailing.paragraph_format.first_line_indent.pt == 21.0
            assert "1、先说明方法选择依据" not in table.rows[5].cells[0].text
            assert "1.\u2009建立结构化测试数据" in table.rows[5].cells[0].text
            assert "2.\u2009比较 Word 与 LaTeX 输出" in table.rows[5].cells[0].text
            method_intro = next(
                p for p in method_paragraphs if p.text.startswith("先说明方法选择依据")
            )
            assert method_intro.paragraph_format.first_line_indent.pt == 21.0
            nested_runs = [run for paragraph in research_paragraphs for run in paragraph.runs]
            assert any(run.text == "2" and run.font.subscript for run in nested_runs)
            assert normalized["sections"]["research_content"][1]["type"] == "unordered_list"

        if name == "structured-content":
            assert normalized["sections"]["significance_and_status"][0][1]["type"] == "inline_equation"
            assert table.rows[3].cells[0]._tc.xpath(".//m:oMath")
            research = normalized["sections"]["research_content"]
            embedded_block = next(item for item in research if item["type"] == "data_table")
            group = next(item for item in research if item["type"] == "figure_group")
            image = next(item for item in research if item["type"] == "image")
            equation = next(item for item in research if item["type"] == "equation")
            assert len(embedded_block["columns"]) == 4
            assert len(embedded_block["rows"]) == 3
            assert group["figure_label"] == "图 1-1"
            assert image["figure_label"] == "图 1-2"
            assert [item["subfigure_label"] for item in group["items"]] == ["（a）", "（b）"]
            assert equation["expression"]["type"] == "row"
            assert equation["equation_label"] == "(1-1)"
            assert "{{eq:" not in research[0]["runs"][0]["text"]
            assert "(1-1)" in research[0]["runs"][0]["text"]
            assert research[0]["runs"][1]["type"] == "inline_equation"
            research_cell = table.rows[4].cells[0]
            embedded = research_cell.tables[0]
            equation_table = research_cell.tables[1]
            group_table = research_cell.tables[2]
            assert len(embedded.rows) == 4 and len(embedded.columns) == 4
            assert embedded.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
            assert all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None for row in embedded.rows)
            assert len(group_table.rows) == 1 and len(group_table.columns) == 2
            assert group_table.rows[0]._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            assert [item.get("descr") for item in document.element.xpath(".//wp:docPr")] == [
                "虚构的结构化输入流程图",
                "虚构的双路输出流程图",
                "虚构的开题报告单图回归示意图",
            ]
            research_cell_paragraphs = research_cell.paragraphs
            assert any(paragraph._p.xpath("./m:oMath/m:f") for paragraph in research_cell_paragraphs)
            image_paragraph = next(
                paragraph for paragraph in research_cell_paragraphs
                if paragraph.text.strip() == "图 1-2 开题报告单图回归示意"
            )
            assert image_paragraph._p.xpath(".//w:drawing")
            assert image_paragraph._p.xpath(".//w:br")
            assert image_paragraph.paragraph_format.keep_together
            for math_tag in ("m:oMathPara", "m:f", "m:sSub", "m:sSup", "m:rad"):
                assert f"<{math_tag}" in document.element.xml
            assert "效率等于输出浓度" in document.element.xml
            assert equation_table.rows[0].cells[2].text == "(1-1)"
            assert equation_table.rows[0]._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            assert all(
                border.get(qn("w:val")) == "nil"
                for border in equation_table._tbl.tblPr.find(qn("w:tblBorders"))
            )
            assert equation_table.rows[0].cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

        if args.skip_pdf:
            if name == "nested-list":
                data_tex = latex_renderer.data_tex(
                    normalized,
                    data_dir=fixture.parent,
                    assets_dir=output_root / "skip-assets",
                )
                for depth, marker in ((1, "•"), (2, "（1）"), (3, "▪"), (4, "A.")):
                    assert rf"\ProposalListItem{{{depth}}}{{{marker}}}" in data_tex
                assert r"\ProposalParagraph{列表结束后" in data_tex
                assert r"H\textsubscript{2}O" in data_tex
                assert r"\ProposalListItem{1}{1.}{建立结构化测试数据}" in data_tex
                assert r"\ProposalListItem{1}{2.}{比较 Word 与 LaTeX 输出}" in data_tex
            continue

        latex_dir = output_root / f"latex-{name}"
        latex_compile_output = run(
            [
                sys.executable,
                str(proposal_dir / "latex" / "render.py"),
                "--data",
                str(fixture),
                "--output-dir",
                str(latex_dir),
                "--overwrite",
                "--compile",
            ],
            cwd=project_dir,
        )
        assert "Overfull" not in latex_compile_output
        assert "Missing character" not in latex_compile_output
        data_tex = (latex_dir / "proposal-data.tex").read_text(encoding="utf-8")
        fonts_tex = (latex_dir / "proposal-fonts.tex").read_text(encoding="utf-8")
        typography_tex = (latex_dir / "proposal-typography.tex").read_text(
            encoding="utf-8"
        )
        process_form_tex = (latex_dir / "sztu-process-form.tex").read_text(
            encoding="utf-8"
        )
        assert "ProposalAdaptiveLayout" not in data_tex
        assert r"\long\def\MethodsAndMeans" in data_tex
        assert r"\long\def\ResearchSteps" in data_tex
        assert "AutoFakeBold=3" in fonts_tex
        assert r"\newcommand{\SZTUTitleSize}{\zihao{-2}}" in typography_tex
        assert r"\newcommand{\SZTULabelSize}{\zihao{-4}}" in typography_tex
        assert r"\newcommand{\SZTUBodySize}{\zihao{5}}" in typography_tex
        assert r"\newcommand{\SZTUFormRuleWidth}{0.48pt}" in typography_tex
        assert r"\newcommand{\SZTUFlowVerticalPadding}{0.8mm}" in typography_tex
        assert r"\newcommand{\SZTUTitleSingleLineVerticalPadding}{0mm}" in typography_tex
        assert r"\newcommand{\SZTUTitleVerticalPadding}{1.5mm}" in typography_tex
        assert r"\newcommand{\SZTUSectionTitleContentGap}{0mm}" in typography_tex
        assert "sztuformflow/.style" in process_form_tex
        if name == "normal":
            assert r"\long\def\ProposalTitle" in data_tex
            assert r"H\textsubscript{2}O\textsubscript{2}" in data_tex
            assert r"g-C\textsubscript{3}N\textsubscript{4}" in data_tex
            assert r"\textsubscript{3}" in data_tex
            assert r"\textsubscript{4}" in data_tex
            assert r"\textsuperscript{[1]}" in data_tex
            for marker in ("（1）", "（2）", "（3）"):
                assert rf"\NestedOrderedItem{{{marker}}}" in data_tex
            assert "采用模板蒸馏方法" in data_tex
            assert "完成官方文件的结构分析" in data_tex
            assert "1、采用模板蒸馏方法" not in data_tex
            assert "2、采用确定性程序" not in data_tex
            assert "1、完成官方文件的结构分析" not in data_tex
        if name == "long":
            assert "第一阶段完成" in data_tex
            assert "1、第一阶段完成" not in data_tex
        if name == "nested-list":
            for depth, marker in ((1, "•"), (2, "（1）"), (3, "▪"), (4, "A.")):
                assert rf"\ProposalListItem{{{depth}}}{{{marker}}}" in data_tex
            assert r"\ProposalParagraph{列表结束后" in data_tex
            assert r"H\textsubscript{2}O" in data_tex
            assert r"\ProposalListItem{1}{1.}{建立结构化测试数据}" in data_tex
            assert r"\ProposalListItem{1}{2.}{比较 Word 与 LaTeX 输出}" in data_tex
        if name == "structured-content":
            assert r"过程文档中的局部比值 $\frac{a}{b}$" in data_tex
            assert r"\begin{tblr}" in data_tex
            assert r"\begin{minipage}[t]{0.4891\linewidth}" in data_tex
            assert "图 1-1" in data_tex
            assert "图 1-2" in data_tex
            assert r"\ProposalFigure{assets/research-workflow-" in data_tex
            assert r"}{72mm}{图 1-2 开题报告单图回归示意}" in data_tex
            assert "（a）" in data_tex and "（b）" in data_tex
            assert r"\frac{" in data_tex and r"_{out}" in data_tex
            assert r"$\frac{a}{b}$" in data_tex
            assert r"\sqrt{{x}^{2}}" in data_tex
            assert r"\input" not in data_tex
            assert "(1-1)" in data_tex and "{{eq:" not in data_tex

        # PDF metadata CLI labels are localized on Windows.
        with pdfplumber.open(latex_dir / "main.pdf") as counted_pdf:
            pages = len(counted_pdf.pages)
        if name == "page-range-maximum":
            minimum, maximum = page_range_recipe["expected_pdf_pages"]
            assert minimum <= pages <= maximum
        else:
            assert 1 <= pages <= 8
        fonts = run(["pdffonts", str(latex_dir / "main.pdf")], cwd=project_dir)
        assert "SimSun" in fonts
        assert "SimHei" in fonts
        assert "TimesNewRoman" in fonts
        with pdfplumber.open(latex_dir / "main.pdf") as pdf:
            page_chars = [list(page.chars) for page in pdf.pages]
            page_words = [page.extract_words() for page in pdf.pages]
            page_text_lines = [page.extract_text_lines(return_chars=True) for page in pdf.pages]
            chars = [
                char
                for current_page_chars in page_chars
                for char in current_page_chars
                if char.get("text", "").strip()
            ]
            first_page_chars = page_chars[0]
            first_page_lines = list(pdf.pages[0].lines)
            first_page_edges = list(pdf.pages[0].edges)
            page_edges = [list(page.edges) for page in pdf.pages]
            page_dimensions = [(float(page.width), float(page.height)) for page in pdf.pages]
        if name == "nested-list":
            def word_x(prefix: str) -> float:
                return float(next(
                    word["x0"] for word in page_words[0]
                    if word["text"].startswith(prefix)
                ))

            marker_x = [
                word_x("•一级无序：建立"),
                word_x("（1）二级有序："),
                word_x("▪三级无序："),
                word_x("A.四级有序："),
            ]
            assert all(
                3.0 <= following - current <= 20.0
                for current, following in zip(marker_x, marker_x[1:])
            ), marker_x
            assert 29.0 <= marker_x[-1] - marker_x[0] <= 34.0
            assert abs(word_x("•一级对齐测试") - marker_x[0]) <= 1.0
            assert abs(word_x("列表结束后") - marker_x[0]) <= 1.0
            assert 8.0 <= word_x("◦检查项目符号") - word_x("1.建立结构化测试数据") <= 15.0
            lines = page_text_lines[0]
            for prefix, marker_chars in (
                ("•一级对齐测试", 1),
                ("（1）二级有序：", 3),
                ("◦二级对齐测试", 1),
            ):
                first_index = next(
                    index for index, line in enumerate(lines)
                    if line["text"].startswith(prefix)
                )
                first_line = lines[first_index]
                continuation = next(
                    line for line in lines[first_index + 1 :]
                    if line["top"] - first_line["top"] >= 10.0
                )
                assert 12.0 <= continuation["top"] - first_line["top"] <= 20.0
                first_body_x = float(first_line["chars"][marker_chars]["x0"])
                continuation_x = float(continuation["chars"][0]["x0"])
                assert abs(first_body_x - continuation_x) <= 0.75, (
                    prefix, first_body_x, continuation_x
                )
        simsun_sizes = {
            round(float(char["size"]), 2)
            for char in chars
            if "SimSun" in char["fontname"]
        }
        simhei_sizes = {
            round(float(char["size"]), 2)
            for char in chars
            if "SimHei" in char["fontname"]
        }
        assert all(
            any(
                name in char["fontname"]
                for name in (
                    "SimSun",
                    "SimHei",
                    "TimesNewRoman",
                    "CMMI",
                    "CMR",
                    "CMSY",
                )
            )
            for char in chars
        )
        assert 10.50 in simsun_sizes and 12.00 in simsun_sizes
        assert 12.00 in simhei_sizes and 18.00 in simhei_sizes
        horizontal_candidates = [
            line
            for line in first_page_edges
            if abs(float(line["bottom"]) - float(line["top"])) < 0.1
            and float(line["top"]) < 400
        ]
        horizontal = [
            top
            for top in clustered(
                [round(float(line["top"]), 2) for line in horizontal_candidates]
            )
            if (
                max(
                    float(line["x1"])
                    for line in horizontal_candidates
                    if abs(float(line["top"]) - top) <= 1.0
                )
                - min(
                    float(line["x0"])
                    for line in horizontal_candidates
                    if abs(float(line["top"]) - top) <= 1.0
                )
                > 400
            )
        ]
        assert len(horizontal) >= 4
        rows = horizontal[:4]
        if name in {"short", "title-two-lines"}:
            title_height_mm = (rows[1] - rows[0]) * 25.4 / 72
            title_chars = cell_chars(
                first_page_chars,
                (84.6, rows[0], 504.7, rows[1]),
                font="SimSun",
                size=10.50,
            )
            title_chars = [char for char in title_chars if float(char["x0"]) >= 169]
            title_line_tops = clustered([round(float(char["top"]), 1) for char in title_chars])
            if name == "short":
                assert 8.0 <= title_height_mm <= 8.6, title_height_mm
                assert len(title_line_tops) == 1, title_line_tops
            else:
                assert 11.5 <= title_height_mm <= 13.2, title_height_mm
                assert len(title_line_tops) == 2, title_line_tops
                assert min(float(char["top"]) for char in title_chars) - rows[0] >= 3.5
                assert rows[1] - max(float(char["bottom"]) for char in title_chars) >= 3.5
        vertical = sorted(
            {
                round(float(line["x0"]), 2)
                for line in first_page_edges
                if abs(float(line["x1"]) - float(line["x0"])) < 0.1
                and float(line["top"]) < rows[-1] - 0.1
                and float(line["bottom"]) > rows[0] + 0.1
            }
        )
        assert len(vertical) == 8
        cols = vertical

        label_cells = (
            (cols[0], rows[0], cols[1], rows[1]),
            (cols[0], rows[1], cols[1], rows[2]),
            (cols[2], rows[1], cols[3], rows[2]),
            (cols[5], rows[1], cols[6], rows[2]),
            (cols[0], rows[2], cols[1], rows[3]),
            (cols[2], rows[2], cols[4], rows[3]),
        )
        for bounds in label_cells:
            assert_centered(
                cell_chars(first_page_chars, bounds, font="SimHei", size=12.00),
                bounds,
            )
        centered_data_cells = (
            ((cols[1], rows[1], cols[2], rows[2]), "SimSun"),
            ((cols[3], rows[1], cols[5], rows[2]), "TimesNewRoman"),
            ((cols[4], rows[2], cols[7], rows[3]), "SimSun"),
        )
        for bounds, font in centered_data_cells:
            assert_centered(
                cell_chars(first_page_chars, bounds, font=font, size=10.50),
                bounds,
            )
        left_data_cells = (
            (cols[1], rows[0], cols[7], rows[1]),
            (cols[6], rows[1], cols[7], rows[2]),
            (cols[1], rows[2], cols[2], rows[3]),
        )
        for bounds in left_data_cells:
            selected = cell_chars(first_page_chars, bounds, font="SimSun", size=10.50)
            assert_vertically_centered(selected, bounds)
            # A left-aligned cell must retain the 2 mm internal inset instead
            # of placing glyphs directly against the border.
            assert min(float(char["x0"]) for char in selected) - bounds[0] >= 4.5

        if name == "short":
            # Short majors are centered; long majors are covered by the
            # left-aligned multi-line layout_stress fixture below.
            assert_centered(
                cell_chars(
                    first_page_chars,
                    left_data_cells[1],
                    font="SimSun",
                    size=10.50,
                ),
                left_data_cells[1],
            )

        if name == "layout_stress":
            assert abs(rows[0] - 134.64) <= 1.0
            assert abs(cols[0] - 84.60) <= 1.0
            assert abs(cols[-1] - 504.70) <= 1.0
            metadata_vertical_bottom = max(
                float(line["bottom"])
                for line in first_page_lines
                if abs(float(line["x1"]) - float(line["x0"])) < 0.1
                and float(line["top"]) < rows[-1]
            )
            # The shared tight-join skin overlaps exactly one rule width, so
            # the metadata verticals may end on the inner edge of that rule.
            assert abs(metadata_vertical_bottom - rows[-1]) <= 0.6
            assert all(
                abs(float(line["linewidth"]) - 0.48) <= 0.01
                for line in first_page_lines
                if float(line["top"]) < rows[-1] + 0.1
            )

            for bounds in left_data_cells:
                selected = cell_chars(first_page_chars, bounds, font="SimSun", size=10.50)
                starts = []
                for top in sorted({round(float(char["top"]), 1) for char in selected}):
                    starts.append(
                        min(
                            float(char["x0"])
                            for char in selected
                            if round(float(char["top"]), 1) == top
                        )
                    )
                assert len(starts) >= 2
                assert max(starts) - min(starts) <= 0.8
                assert min(starts) - bounds[0] >= 4.5
                assert min(starts) - bounds[0] <= 9.0

            extracted = run(
                ["pdftotext", "-layout", str(latex_dir / "main.pdf"), "-"],
                cwd=project_dir,
            )
            assert "ENDREF-A" in extracted and "ENDREF-B" in extracted
            reference_lines = extracted.split("参考文献：", 1)[1].split("学生签名", 1)[0]
            assert len([line for line in reference_lines.splitlines() if line.strip()]) >= 8
        if name == "long":
            raster_dir = latex_dir / "raster-180"
            raster_dir.mkdir(parents=True, exist_ok=True)
            run(
                [
                    "pdftocairo",
                    "-png",
                    "-r",
                    "180",
                    str(latex_dir / "main.pdf"),
                    str(raster_dir / "page"),
                ],
                cwd=project_dir,
            )
            break_rule_checks = 0
            separator_checks = 0
            for page_index, (edges, dimensions) in enumerate(
                zip(page_edges, page_dimensions, strict=True), start=1
            ):
                page_width, page_height = dimensions
                with Image.open(raster_dir / f"page-{page_index}.png") as image:
                    boundary_edges = [
                        edge
                        for edge in edges
                        if abs(float(edge["bottom"]) - float(edge["top"])) < 0.1
                        and float(edge["x1"]) - float(edge["x0"]) > 400
                        and (
                            float(edge["top"]) < 100
                            or float(edge["top"]) > page_height - 100
                        )
                    ]
                    for top in clustered(
                        [float(edge["top"]) for edge in boundary_edges]
                    ):
                        edge = max(
                            (
                                candidate
                                for candidate in boundary_edges
                                if abs(float(candidate["top"]) - top) <= 1.0
                            ),
                            key=lambda candidate: float(candidate["x1"])
                            - float(candidate["x0"]),
                        )
                        assert_rule_has_raster_ink(
                            image,
                            page_width=page_width,
                            page_height=page_height,
                            x0=float(edge["x0"]),
                            x1=float(edge["x1"]),
                            top=float(edge["top"]),
                        )
                        break_rule_checks += 1

                    for edge in edges:
                        top = float(edge["top"])
                        if not (
                            abs(float(edge["bottom"]) - top) < 0.1
                            and float(edge["x1"]) - float(edge["x0"]) > 400
                            and 0.35 <= float(edge.get("linewidth", 0)) <= 0.50
                            and 100 < top < page_height - 100
                        ):
                            continue
                        crossing_sides = [
                            float(side["x0"])
                            for side in edges
                            if abs(float(side["x1"]) - float(side["x0"])) < 0.1
                            and float(side["top"]) <= top
                            and float(side["bottom"]) >= top
                            and float(side["bottom"]) - float(side["top"]) > 20
                        ]
                        if len(crossing_sides) < 2:
                            continue
                        assert float(edge["x0"]) - min(crossing_sides) <= 0.6
                        assert max(crossing_sides) - float(edge["x1"]) <= 0.6
                        assert_rule_has_raster_ink(
                            image,
                            page_width=page_width,
                            page_height=page_height,
                            x0=min(crossing_sides),
                            x1=max(crossing_sides),
                            top=top,
                        )
                        separator_checks += 1
            assert break_rule_checks >= 2
            assert separator_checks >= 2
        if name == "normal":
            all_words = [word for words in page_words for word in words]
            for expected_label in (
                "本选题的意义及国内外发展状况：",
                "研究内容：",
                "研究方法、手段及步骤：",
                "参考文献：",
            ):
                label_page_index, label = next(
                    (page_index, word)
                    for page_index, words in enumerate(page_words)
                    for word in words
                    if expected_label in word["text"]
                )
                preceding_rules = [
                    edge
                    for edge in page_edges[label_page_index]
                    if abs(float(edge["bottom"]) - float(edge["top"])) < 0.1
                    and float(edge["x1"]) - float(edge["x0"]) > 400
                    and 0.35 <= float(edge.get("linewidth", 0)) <= 0.50
                    and float(edge["top"]) <= float(label["top"])
                ]
                assert preceding_rules
                preceding_rule = max(preceding_rules, key=lambda edge: float(edge["top"]))
                title_top_gap_pt = float(label["top"]) - float(preceding_rule["top"])
                assert 1.0 <= title_top_gap_pt <= 4.5
            methods_label = next(
                word for word in all_words if word["text"] == "研究方法、手段及步骤："
            )
            methods_intro = next(
                word for word in all_words if word["text"] == "本课题研究方法、手段如下："
            )
            steps_intro = next(
                word for word in all_words if word["text"] == "本课题研究步骤如下："
            )
            first_method = next(
                word for word in all_words
                if word["text"].startswith("采用模板蒸馏方法提取官方")
            )
            title_to_intro_gap = float(methods_intro["top"]) - float(methods_label["bottom"])
            intro_to_body_gap = float(first_method["top"]) - float(methods_intro["bottom"])
            assert abs(title_to_intro_gap - intro_to_body_gap) <= 2.0
            continuation_word = next(
                word for word in all_words if word["text"].startswith("通过统一数据")
            )
            nested_markers = [
                word
                for word in all_words
                if word["text"] in {"（1）", "（2）", "（3）"}
            ]
            assert 19.0 <= float(methods_intro["x0"]) - float(methods_label["x0"]) <= 23.0
            assert abs(float(steps_intro["x0"]) - float(methods_intro["x0"])) <= 0.3
            assert abs(float(continuation_word["x0"]) - float(methods_intro["x0"])) <= 0.3
            assert len(nested_markers) == 3
            assert all(
                3.0 <= float(marker["x0"]) - float(methods_intro["x0"]) <= 7.0
                for marker in nested_markers
            )

            signature_labels = ("学生签名", "指导教师意见")
            for label_text in signature_labels:
                label_page_index, label = next(
                    (index, word)
                    for index, words in enumerate(page_words)
                    for word in words
                    if word["text"] == f"{label_text}："
                )
                preceding_rules = [
                    edge
                    for edge in page_edges[label_page_index]
                    if abs(float(edge["bottom"]) - float(edge["top"])) < 0.1
                    and float(edge["x1"]) - float(edge["x0"]) > 400
                    and 0.35 <= float(edge.get("linewidth", 0)) <= 0.50
                    and float(edge["top"]) <= float(label["top"])
                ]
                assert preceding_rules
                preceding_rule = max(
                    preceding_rules, key=lambda edge: float(edge["top"])
                )
                label_top_gap_pt = float(label["top"]) - float(preceding_rule["top"])
                assert 1.0 <= label_top_gap_pt <= 4.5
                assert abs(float(preceding_rule["x0"]) - 84.60) <= 1.0
                assert abs(float(preceding_rule["x1"]) - 504.70) <= 1.0

            teacher_opinion = next(
                word for word in all_words if word["text"] == "指导教师意见："
            )
            college_opinion = next(
                word for word in all_words if word["text"] == "学院领导意见："
            )
            opinion_label_gap_pt = float(college_opinion["top"]) - float(
                teacher_opinion["top"]
            )
            assert 138.0 <= opinion_label_gap_pt <= 146.0

            signature_runs = [
                char
                for char in chars
                if char["text"] == "签"
                and "SimSun" in char["fontname"]
                and round(float(char["size"]), 2) == 12.00
            ]
            assert len(signature_runs) == 2
            assert all(357 <= float(char["x0"]) <= 362 for char in signature_runs)
            signature_tops = sorted(float(char["top"]) for char in signature_runs)
            assert 120.0 <= signature_tops[0] - float(teacher_opinion["top"]) <= 132.0
            assert 72.0 <= signature_tops[1] - float(college_opinion["top"]) <= 86.0
            date_year_runs = [
                char
                for char in chars
                if char["text"] == "年"
                and "SimSun" in char["fontname"]
                and round(float(char["size"]), 2) == 12.00
            ]
            assert len(date_year_runs) == 2
            assert all(403 <= float(char["x0"]) <= 408 for char in date_year_runs)
            date_tops = sorted(float(char["top"]) for char in date_year_runs)
            assert date_tops[-1] > signature_tops[-1]
            assert 12.0 <= date_tops[-1] - signature_tops[-1] <= 20.0
        run(
            [
                sys.executable,
                str(project_dir / "scripts" / "validate_cjk_render.py"),
                str(latex_dir / "main.pdf"),
            ],
            cwd=project_dir,
        )

    overrides = json.loads((proposal_dir / "fixtures" / "intro-variants.json").read_text(encoding="utf-8"))
    variant = json.loads((proposal_dir / "fixtures" / "normal.json").read_text(encoding="utf-8"))
    variant["sections"].update(overrides)
    variant_path = output_root / "intro-variants.json"
    variant_path.write_text(json.dumps(variant, ensure_ascii=False), encoding="utf-8")
    variant_docx = output_root / "intro-variants.docx"
    word_renderer.render(proposal_dir / "word" / "official-template.docx", variant_path, variant_docx, overwrite=True)
    variant_table = Document(variant_docx).tables[0]
    assert "研究范围和目标如下：" in variant_table.rows[4].cells[0].text
    assert "实施安排如下：" in variant_table.rows[5].cells[0].text
    assert "本课题研究方法、手段如下：" not in variant_table.rows[5].cells[0].text
    variant_tex = latex_renderer.data_tex(word_renderer.validate_data(variant), data_dir=output_root, assets_dir=output_root)
    assert "研究范围和目标如下：" in variant_tex
    assert "实施安排如下：" in variant_tex
    assert "本课题研究方法、手段如下：" not in variant_tex

    legacy = json.loads((proposal_dir / "fixtures" / "normal.json").read_text(encoding="utf-8"))
    legacy["sections"]["methods_and_means"].pop(0)
    legacy["sections"]["methods_and_means_intro"] = "本课题研究方法、手段如下："
    legacy_path = output_root / "legacy-intro.json"
    legacy_path.write_text(json.dumps(legacy, ensure_ascii=False), encoding="utf-8")
    legacy_docx = output_root / "legacy-intro.docx"
    word_renderer.render(proposal_dir / "word" / "official-template.docx", legacy_path, legacy_docx, overwrite=True)
    assert Document(legacy_docx).tables[0].rows[5].cells[0].text.count("本课题研究方法、手段如下：") == 1

    print("proposal regression checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
